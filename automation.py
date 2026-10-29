"""Read a selected source sheet and fill a fresh copy of a spreadsheet template."""

from __future__ import annotations

import csv
import re
import unicodedata
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from email_validation import validated_email


FIELD_MAPPING = {
    "debtor_email": "email",
    "OLD I.C": "indexoneuid",
    "Account no.": "primary_key",
    "financing/Card no.": "cardnumber",
    "Name": "FullName",
    "endo date": "endodate",
    "level": "level",
    "product": "product",
    "OB (Masterlist)": "OB",
    "phone_1": "phone",
    "Christmas Amnesty": "ChristmasAmnesty",
}
OPTIONAL_FIELD_MAPPING = {"Area": "AREA"}
SOURCE_ALIASES = {
    "debtor_email": ("email", "email address", "e-mail address", "email 1", "debtor email address", "email id"),
    "Account no.": ("account", "account number", "acct no", "acct number"),
    "financing/Card no.": ("card_no",),
    "endo date": ("assign_date",),
    "level": ("placement",),
    "product": ("Product Type",),
    "OB (Masterlist)": ("OB",),
    "phone_1": ("phone number", "phone no", "mobile", "mobile number", "mobile no", "contact number", "contact no", "cellphone", "cellphone number"),
}
MAPPING_VERSION = 6
WORKBOOK_READER_VERSION = 1
OB_SORT_VERSION = 1
TARGET_HEADERS = tuple(FIELD_MAPPING.values())
DERIVED_TEMPLATE_HEADERS = ("doclass",)
GEO_DOCLASS_BY_AREA = {
    "mindanao": "SBC GEO Mindanao",
    "r3": "SBC GEO R3",
    "region3": "SBC GEO R3",
}
# Keep the full campaign layout separate from the fields required for mapping.
DEFAULT_TEMPLATE_HEADERS = (
    "email", "primary_key", "indexoneuid", "cardnumber", "FullName", "phone", "OB",
    "Monthly", "Terms", "doclass", "Level", "Product", "EndoDate", "AREA",
    "ChristmasAmnesty", "NewEndo",
)
CSV_OUTPUT_ENCODINGS = {
    "CSV (Comma delimited)": "cp1252",
    "CSV UTF-8 (Comma delimited)": "utf-8-sig",
}
IDENTIFIER_FIELDS = frozenset({"indexoneuid", "primary_key", "cardnumber", "phone"})
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class AutomationError(ValueError):
    """An actionable input problem that can be displayed in the app."""


@dataclass(frozen=True)
class SheetInfo:
    name: str
    header_row: int
    matched_fields: int
    has_area: bool = False
    available_fields: tuple[str, ...] = ()
    headers: tuple[str, ...] = ()


@dataclass
class MappedData:
    records: list[dict[str, Any]]
    source_rows: list[int]
    blank_rows_skipped: int = 0
    no_contact_rows_skipped: int = 0
    invalid_email_rows: int = 0
    valid_email_rows: int = 0
    missing_email_rows: int = 0


def normalize_header(value: Any) -> str:
    """Match headers despite capitalization, extra whitespace, and punctuation."""
    text = unicodedata.normalize("NFKC", str(value or "")).casefold()
    return "".join(character for character in text if character.isalnum())


def ob_sort_key(value: Any) -> tuple[bool, Decimal]:
    """Compare OB amounts numerically; reverse sorting puts missing amounts last."""
    if value is not None and not isinstance(value, bool):
        text = re.sub(r"^(?:PHP|₱)\s*", "", str(value).strip(), flags=re.IGNORECASE)
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1].strip()
        try:
            amount = Decimal(text.replace(",", ""))
            if amount.is_finite():
                return True, amount
        except InvalidOperation:
            pass
    return False, Decimal(0)


def geo_doclass(area: Any) -> str | None:
    """Derive the GEO campaign class from the account's source Area."""
    return GEO_DOCLASS_BY_AREA.get(normalize_header(area))


def _source_aliases(source: str) -> set[str]:
    target = {**FIELD_MAPPING, **OPTIONAL_FIELD_MAPPING}.get(source, source)
    return {normalize_header(label) for label in (source, target, *SOURCE_ALIASES.get(source, ()))}


def source_column_positions(headers: list[Any] | tuple[str, ...]) -> dict[str, int]:
    """Find unambiguous defaults for the source-column mapping controls."""
    labels = [normalize_header(value) for value in headers]
    positions = {}
    for source, target in {**FIELD_MAPPING, **OPTIONAL_FIELD_MAPPING}.items():
        matches = [index for index, label in enumerate(labels) if label in _source_aliases(source)]
        if len(matches) == 1:
            positions[target] = matches[0]
    return positions


def _load(data: bytes, *, read_only: bool = True, data_only: bool = False):
    if not data:
        raise AutomationError("The workbook is empty. Upload an .xlsx file.")
    try:
        workbook = load_workbook(BytesIO(data), read_only=read_only, data_only=data_only)
        if read_only:
            for sheet in workbook.worksheets:
                sheet.reset_dimensions()
        return workbook
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise AutomationError(
            "This file could not be read as an Excel workbook. "
            "Open it in Excel and save an unencrypted .xlsx copy."
        ) from exc


def _check_header_row(row: int) -> None:
    if not isinstance(row, int) or row < 1 or row >= 1_048_576:
        raise AutomationError("The header row must be between 1 and 1,048,575.")


def _header_groups(template: bool) -> list[set[str]]:
    if template:
        return [{normalize_header(header)} for header in TARGET_HEADERS]
    return [_source_aliases(source) for source in FIELD_MAPPING]


def _detect_header(sheet, *, template: bool = False) -> tuple[int, int]:
    groups = _header_groups(template)
    best_row, best_score = 1, 0
    for row_number, row in enumerate(
        sheet.iter_rows(max_row=min(sheet.max_row or 50, 50), values_only=True), 1
    ):
        labels = {normalize_header(value) for value in row if value is not None}
        score = sum(bool(labels & group) for group in groups)
        if score > best_score:
            best_row, best_score = row_number, score
        if score == len(groups):
            break
    return best_row, best_score


def _sheet_info(sheet, *, template: bool = False, header_row: int | None = None) -> SheetInfo:
    if header_row is None:
        header_row, matched_fields = _detect_header(sheet, template=template)
    else:
        labels = {normalize_header(value) for value in _headers(sheet, header_row)}
        matched_fields = sum(bool(labels & group) for group in _header_groups(template))
    headers = tuple("" if value is None else str(value) for value in _headers(sheet, header_row))
    available = tuple(source_column_positions(headers))
    return SheetInfo(sheet.title, header_row, matched_fields, "AREA" in available, available, headers)


def inspect_source_sheet(data: bytes, sheet_name: str, header_row: int) -> SheetInfo:
    """Check the available source fields after a manual header-row change."""
    workbook = _load(data)
    try:
        return _sheet_info(_worksheet(workbook, sheet_name), header_row=header_row)
    finally:
        workbook.close()


def preview_source_rows(data: bytes, sheet_name: str, header_row: int) -> list[tuple[Any, ...]]:
    """Show a few saved source values before mapping or validating contacts."""
    _check_header_row(header_row)
    workbook = _load(data, data_only=True)
    try:
        return list(_worksheet(workbook, sheet_name).iter_rows(
            min_row=header_row + 1, max_row=min(header_row + 5, 1_048_576), values_only=True,
        ))
    finally:
        workbook.close()


def inspect_workbook(data: bytes, *, template: bool = False) -> list[SheetInfo]:
    """Find likely header rows without loading all the worksheet data into memory."""
    workbook = _load(data)
    try:
        return [_sheet_info(sheet, template=template) for sheet in workbook.worksheets]
    finally:
        workbook.close()


def default_source_sheet(sheets: list[SheetInfo]) -> str:
    if not sheets:
        raise AutomationError("The source workbook has no worksheets.")
    usable = [sheet for sheet in sheets
              if {"email", "phone"}.intersection(sheet.available_fields)
              and {"primary_key", "cardnumber"}.intersection(sheet.available_fields)]
    candidates = usable or sheets
    return next(
        (sheet.name for sheet in candidates if normalize_header(sheet.name) == "negative"),
        max(candidates, key=lambda sheet: sheet.matched_fields).name,
    )


def _worksheet(workbook, name: str):
    if name not in workbook.sheetnames:
        raise AutomationError(f'The worksheet "{name}" was not found.')
    return workbook[name]


def _headers(sheet, row: int) -> list[Any]:
    _check_header_row(row)
    return list(next(sheet.iter_rows(min_row=row, max_row=row, values_only=True), ()))


def _column_lookup(
    headers: list[Any], *, template: bool = False, allow_missing: bool = False,
    source_columns: dict[str, int | None] | None = None,
) -> dict[str, int]:
    """Resolve required fields by column position; reject ambiguous headers."""
    normalized = [normalize_header(value) for value in headers]
    positions = {}
    missing = []
    fields = {**FIELD_MAPPING, **OPTIONAL_FIELD_MAPPING}
    if template:
        fields.update((header, header) for header in DERIVED_TEMPLATE_HEADERS)
    elif source_columns is not None and set(source_columns) - set(fields.values()):
        raise AutomationError("The source-column mapping contains an unknown template field.")
    for source, target in fields.items():
        if not template and source_columns is not None and target in source_columns:
            column = source_columns[target]
            if column is None:
                if source in FIELD_MAPPING:
                    missing.append(source)
                continue
            if type(column) is not int or not 0 <= column < len(headers):
                raise AutomationError(f'Choose an existing source column for "{target}".')
            positions[target] = column
            continue
        aliases = {normalize_header(target)} if template else _source_aliases(source)
        matches = [index for index, value in enumerate(normalized) if value in aliases]
        label = target if template else source
        if not matches and source in FIELD_MAPPING:
            missing.append(label)
        elif len(matches) > 1:
            raise AutomationError(
                f'Duplicate or ambiguous header "{label}". '
                "Give that column a unique header before processing."
            )
        elif matches:
            positions[target] = matches[0]
    if missing and (template or not allow_missing):
        kind = "Template" if template else "Source"
        raise AutomationError(
            f"{kind} headers are missing: {', '.join(missing)}. "
            "Check the selected sheet and header row."
        )
    return positions


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def has_contact_value(value: Any) -> bool:
    return not _is_blank(value)


def has_contact_details(record: dict[str, Any]) -> bool:
    """Keep accounts with a valid email address or a nonempty phone field."""
    return validated_email(record.get("email")) is not None or has_contact_value(record.get("phone"))


def _identifier(value: Any, number_format: str) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        text = str(int(value)) if float(value).is_integer() else str(value)
        # Excel may store 123 with a display mask of 000000. Retain 000123.
        mask = number_format.split(";")[0].strip()
        if re.fullmatch(r"0+", mask) and float(value).is_integer() and value >= 0:
            text = text.zfill(len(mask))
        return text
    return str(value).strip()


def read_source(
    data: bytes, sheet_name: str, header_row: int, *, allow_missing: bool = False,
    source_columns: dict[str, int | None] | None = None,
) -> MappedData:
    """Read all mapped values, including hidden/filtered rows, in source order."""
    _check_header_row(header_row)
    formulas = _load(data)
    cached = None
    try:
        source = _worksheet(formulas, sheet_name)
        positions = _column_lookup(_headers(source, header_row), allow_missing=allow_missing, source_columns=source_columns)
        contact_columns = {positions[target] for target in ("email", "phone") if target in positions}
        account_columns = {positions[target] for target in ("primary_key", "cardnumber", "indexoneuid") if target in positions}
        if contact_columns & account_columns:
            raise AutomationError("Choose separate source columns for account identifiers and contact details.")
        if allow_missing:
            if not {"email", "phone"}.intersection(positions):
                raise AutomationError("No email or phone column was found. Check the source sheet and header row.")
            if not {"primary_key", "cardnumber"}.intersection(positions):
                raise AutomationError("No account-number column was found. Use Account no., primary_key, or cardnumber as its header.")
        cached = _load(data, data_only=True)
        values = _worksheet(cached, sheet_name)
        result = MappedData([], [])
        max_column = max(positions.values()) + 1
        source_iterator = source.iter_rows(min_row=header_row + 1, max_col=max_column)
        value_iterator = values.iter_rows(min_row=header_row + 1, max_col=max_column)
        for row_number, (source_row, value_row) in enumerate(
            zip(source_iterator, value_iterator), header_row + 1
        ):
            record = dict.fromkeys(TARGET_HEADERS)
            invalid_email = False
            for target, column in positions.items():
                original, cell = source_row[column], value_row[column]
                value = cell.value
                if original.data_type == "f" and value is None:
                    raise AutomationError(
                        f'{sheet_name}, row {row_number}, {target}: a formula has no saved value. '
                        "Open the source in Excel, recalculate and save it, then upload it again."
                    )
                if cell.data_type == "e":
                    raise AutomationError(
                        f'{sheet_name}, row {row_number}, {target}: Excel error {value}. '
                        "Fix that source cell before processing."
                    )
                if target in IDENTIFIER_FIELDS:
                    value = _identifier(value, cell.number_format)
                elif target == "email":
                    email = validated_email(value)
                    if not _is_blank(value) and email is None:
                        invalid_email = True
                        result.invalid_email_rows += 1
                    value = email
                elif _is_blank(value):
                    value = None
                if target == "ChristmasAmnesty":
                    value = "QUALIFIED" if normalize_header(value) == "qualified" else None
                record[target] = value
            if not invalid_email and all(_is_blank(value) for value in record.values()):
                result.blank_rows_skipped += 1
                continue
            if record.get("email") is not None:
                result.valid_email_rows += 1
            elif not invalid_email:
                result.missing_email_rows += 1
            if not has_contact_details(record):
                result.no_contact_rows_skipped += 1
                continue
            result.records.append(record)
            result.source_rows.append(row_number)
        return result
    finally:
        formulas.close()
        if cached is not None:
            cached.close()


def create_default_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Template"
    sheet.append(DEFAULT_TEMPLATE_HEADERS)
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="167D78")
        sheet.column_dimensions[cell.column_letter].width = 21
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _write_excel_value(cell, value: Any, target: str | None = None) -> None:
    cell.value = value
    # Values supplied by the source are data, even when they begin with '='.
    if isinstance(value, str):
        cell.data_type = "s"
    if target in IDENTIFIER_FIELDS:
        cell.number_format = "@"
    elif isinstance(value, (date, datetime)) and cell.number_format == "General":
        cell.number_format = "yyyy-mm-dd"


def fill_excel_template(
    template: bytes,
    records: list[dict[str, Any]],
    sheet_name: str,
    header_row: int,
) -> bytes:
    """Keep template headers, column order, styles, and other sheets in a new file.

    The chosen sheet's data area is replaced. Non-mapped template columns are
    left blank; this prevents old example data from leaking into a fresh output.
    """
    _check_header_row(header_row)
    if header_row + len(records) > 1_048_576:
        raise AutomationError("The output exceeds Excel's worksheet row limit.")
    workbook = _load(template, read_only=False)
    try:
        sheet = _worksheet(workbook, sheet_name)
        headers = _headers(sheet, header_row)
        positions = _column_lookup(headers, template=True)
        if any(area.max_row > header_row for area in sheet.merged_cells.ranges):
            raise AutomationError(
                "The template has merged cells in its data area. "
                "Unmerge them before filling the template."
            )
        first_data_row = header_row + 1
        styles = [copy(sheet.cell(first_data_row, column)._style)
                  for column in range(1, len(headers) + 1)]
        row_height = sheet.row_dimensions[first_data_row].height
        for row in sheet.iter_rows(min_row=first_data_row):
            for cell in row:
                cell.value = None
                cell.hyperlink = None
                cell.comment = None
        for offset, record in enumerate(records):
            row_number = first_data_row + offset
            if row_height is not None:
                sheet.row_dimensions[row_number].height = row_height
            # Imported records should remain visible even if the template was filtered.
            sheet.row_dimensions[row_number].hidden = False
            for column, style in enumerate(styles, 1):
                sheet.cell(row_number, column)._style = copy(style)
            for target, column in positions.items():
                value = geo_doclass(record.get("AREA")) if target == "doclass" else record.get(target)
                _write_excel_value(sheet.cell(row_number, column + 1), value, target)
        last_row = max(header_row + len(records), first_data_row)
        if sheet.auto_filter.ref:
            sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
            sheet.auto_filter.filterColumn = []
            sheet.auto_filter.sortState = None
        for table in sheet.tables.values():
            min_column, min_row, max_column, _ = range_boundaries(table.ref)
            if min_row != header_row or table.totalsRowCount:
                raise AutomationError(
                    "The template contains a table outside its header/data area or a totals row. "
                    "Use a template with a single header row and no totals row."
                )
            table.ref = (
                f"{get_column_letter(min_column)}{min_row}:"
                f"{get_column_letter(max_column)}{last_row}"
            )
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
                table.autoFilter.filterColumn = []
                table.autoFilter.sortState = None
            for column in table.tableColumns:
                column.calculatedColumnFormula = None
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


def _csv_contents(data: bytes) -> tuple[list[str], str]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252")
    try:
        delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    try:
        headers = next(csv.reader(StringIO(text), delimiter=delimiter), [])
    except csv.Error as exc:
        raise AutomationError("The CSV template could not be read.") from exc
    _column_lookup(headers, template=True)
    return headers, delimiter


def csv_template_headers(data: bytes) -> list[str]:
    return _csv_contents(data)[0]


def _encode_csv(text: str, encoding: str) -> bytes:
    if encoding not in CSV_OUTPUT_ENCODINGS.values():
        raise AutomationError("Choose a supported CSV file type.")
    try:
        return text.encode(encoding)
    except UnicodeEncodeError as exc:
        raise AutomationError(
            "Some characters cannot be saved in CSV (Comma delimited). "
            "Select CSV UTF-8 (Comma delimited) to keep the original text."
        ) from exc


def write_text_csv_sheet(workbook: Workbook, data: bytes, title: str):
    """Write literal text cells so Excel cannot round or evaluate exported values."""
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    encoding = "utf-8-sig" if data.startswith(b"\xef\xbb\xbf") else "cp1252"
    for row_index, row in enumerate(csv.reader(StringIO(data.decode(encoding), newline=""))):
        cells = []
        for column, value in enumerate(row, 1):
            if row_index == 0:
                sheet.column_dimensions[get_column_letter(column)].width = 24
            cell = WriteOnlyCell(sheet, value=value)
            cell.data_type = "s"
            cell.number_format = "@"
            cells.append(cell)
        sheet.append(cells)
    return sheet


def _csv_header(headers: list[Any] | tuple[str, ...], encoding: str = "utf-8-sig") -> bytes:
    output = StringIO(newline="")
    csv.writer(output).writerow(headers)
    return _encode_csv(output.getvalue(), encoding)


def create_default_csv_template(*, encoding: str = "cp1252") -> bytes:
    return _csv_header(DEFAULT_TEMPLATE_HEADERS, encoding)


def excel_template_to_csv(template: bytes, sheet_name: str, header_row: int) -> bytes:
    """Use only the selected Excel header row as the CSV column layout."""
    workbook = _load(template)
    try:
        headers = _headers(_worksheet(workbook, sheet_name), header_row)
        _column_lookup(headers, template=True)
        return _csv_header(headers)
    finally:
        workbook.close()


def fill_csv_template(
    template: bytes, records: list[dict[str, Any]], *, encoding: str = "cp1252",
) -> bytes:
    """Always write comma-delimited CSV, regardless of the template delimiter."""
    headers, _ = _csv_contents(template)
    positions = _column_lookup(headers, template=True)
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=",", lineterminator="\r\n")
    writer.writerow(headers)
    for record in records:
        row = [""] * len(headers)
        for target, position in positions.items():
            value = geo_doclass(record.get("AREA")) if target == "doclass" else record.get(target)
            if isinstance(value, (date, datetime)):
                value = value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
            row[position] = "" if value is None else value
        writer.writerow(row)
    return _encode_csv(output.getvalue(), encoding)


def output_filename(source_name: str, sheet_name: str, extension: str = ".csv") -> str:
    safe_sheet = re.sub(r'[^\w .-]+', "_", sheet_name).strip(" .") or "sheet"
    safe_source = re.sub(r'[^\w .-]+', "_", Path(source_name).stem).strip(" .") or "source"
    return f"{safe_source}_{safe_sheet}_filled{extension}"
