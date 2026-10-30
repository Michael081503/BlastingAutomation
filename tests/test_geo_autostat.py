from __future__ import annotations

import csv
import unittest
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from autostat import build_autostat_channel_excel, build_autostat_excel, build_autostat_excel_zip, build_autostat_exports
from geo_autostat import mapped_contact_source


ROOT = Path(__file__).resolve().parents[1]


class Upload(BytesIO):
    def __init__(self, data, name="source.xlsx"):
        super().__init__(data)
        self.name = name


def workbook_data(first_account="000000000000000001.0", first_card="12345678901234567.0"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NEGATIVE"
    headers = ["primary_key", "cardnumber", "indexoneuid", "email", "phone", "OB", "Product", "Level", "ChristmasAmnesty"]
    sheet.append(headers)
    for row in (
        [first_account, first_card, "U1", "first@example.com", "09000000001", 9, "CARDS", 1, None],
        ["000000000000000002", "000123456789012345", "U2", "second@example.com", "09000000002", 1000, "CARDS", 6, None],
        ["000000000000000003", "000123456789012345", "U3", "third@example.com", "09000000003", 500, "CARDS", 1, None],
        [first_account.removesuffix(".0"), first_card.removesuffix(".0"), "U1", "ignored@example.com", "09000000004", 99999, "CARDS", 1, None],
    ):
        sheet.append(row)
    sheet = workbook.create_sheet("GEO")
    sheet.append([*headers, "Area"])
    sheet.append(["000000000000000019", "CARD-GEO", "GEO-1", "geo@example.com", "09000000005", 100, "BEL", 1, "QUALIFIED", "Mindanao"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def log_rows(file):
    return list(csv.DictReader(StringIO(file.data.decode("cp1252"))))


class AutomaticGeoAutostatTests(unittest.TestCase):
    def test_separate_email_viber_and_zip_downloads_follow_template_on_geo_and_reco(self):
        upload = Upload(workbook_data())
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: upload if kw.get("key") == "source_upload" else None), \
                patch("geo_autostat.build_autostat_channel_excel", wraps=build_autostat_channel_excel) as channel_excel, \
                patch("geo_autostat.build_autostat_excel_zip", wraps=build_autostat_excel_zip) as excel_zip, \
                patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            channel_excel.assert_not_called()
            excel_zip.assert_not_called()
            for sheet_name, accounts, phones in (
                ("NEGATIVE", ["000000000000000002", "000000000000000003", "000000000000000001"],
                 ["9000000002", "9000000003", "9000000001"]),
                ("GEO", ["000000000000000019"], ["9000000005"]),
            ):
                with self.subTest(sheet=sheet_name):
                    if sheet_name == "GEO":
                        download.reset_mock()
                        next(widget for widget in app.selectbox if widget.label == "Source sheet").select(sheet_name).run(timeout=20)
                    self.assertEqual(len(app.error), 0)
                    calls = {call.kwargs.get("key"): call for call in download.call_args_list}
                    email = calls["geo_autostat_email_excel"]
                    viber = calls["geo_autostat_viber_excel"]
                    self.assertEqual(email.args[2], "AUTOSTAT_Email.xlsx")
                    self.assertEqual(viber.args[2], "AUTOSTAT_Viber.xlsx")
                    self.assertEqual(viber.args[0], "Download AUTOSTAT_Viber.xlsx")
                    self.assertTrue(callable(viber.args[1]))
                    workbook = load_workbook(BytesIO(viber.args[1]()))
                    self.assertEqual(workbook.sheetnames, ["Viber"])
                    output = workbook.active
                    self.assertEqual(tuple(cell.value for cell in output[1]),
                                     ("Time", "Account No.", "Status", "Remark", "Remark By"))
                    self.assertEqual([cell.value for cell in output["B"][1:]], accounts)
                    self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in output["B"][1:]))
                    self.assertTrue(all(cell.value == "SMS BLAST SENT - WAITING FOR REPLY" for cell in output["C"][1:]))
                    self.assertEqual([cell.value for cell in output["D"][1:]], [
                        f"SPMA | 08 With SMS / email / DL without response - {phone} - VIBER SENT - BLASTING" for phone in phones
                    ])
                    self.assertTrue(all(cell.value == "MBURGOS" for cell in output["E"][1:]))
                    workbook.close()
                    workbook = load_workbook(BytesIO(email.args[1]()))
                    self.assertEqual(workbook.sheetnames, ["Email"])
                    self.assertEqual(workbook.active.max_column, 11)
                    workbook.close()
                    with ZipFile(BytesIO(calls["geo_autostat_zip"].args[1]())) as archive:
                        self.assertEqual(archive.namelist(), ["AUTOSTAT_Email.xlsx", "AUTOSTAT_Viber.xlsx"])
                        workbook = load_workbook(BytesIO(archive.read("AUTOSTAT_Viber.xlsx")))
                        self.assertEqual([cell.value for cell in workbook.active["B"][1:]], accounts)
                        workbook.close()

    def test_viber_empty_channels_explain_missing_or_invalid_phones_and_phone_only_files_work(self):
        for headers, record, expected_viber_rows, warning in (
            (["primary_key", "email"], ["0001", "one@example.com"], 0, "No phone column is selected"),
            (["primary_key", "email", "phone"], ["0001", "one@example.com", "123"], 0, "No valid mobile numbers"),
            (["primary_key", "phone"], ["0001", "09171234567"], 1, None),
        ):
            with self.subTest(headers=headers):
                workbook = Workbook()
                workbook.active.append(headers)
                workbook.active.append(record)
                data = BytesIO()
                workbook.save(data)
                workbook.close()
                upload = Upload(data.getvalue())
                with patch("streamlit.file_uploader", side_effect=lambda label, **kw: upload if kw.get("key") == "source_upload" else None), \
                        patch("streamlit.download_button") as download:
                    app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
                    self.assertEqual(len(app.exception), 0)
                    self.assertEqual(len(app.error), 0)
                    if warning:
                        self.assertTrue(any(warning in item.value for item in app.warning))
                    call = next(call for call in download.call_args_list if call.kwargs.get("key") == "geo_autostat_viber_excel")
                    saved = load_workbook(BytesIO(call.args[1]()))
                    self.assertEqual(saved.active.max_row, expected_viber_rows + 1)
                    self.assertEqual(saved.active.max_column, 5)
                    if expected_viber_rows:
                        self.assertEqual(saved.active["B2"].value, "0001")
                        self.assertTrue(saved.active["D2"].value.endswith("9171234567 - VIBER SENT - BLASTING"))
                    saved.close()

    def test_upload_immediately_prepares_download_and_reruns_reuse_it(self):
        uploads = [Upload(workbook_data())]
        uploader = lambda label, **kw: uploads[0] if kw.get("key") == "source_upload" else None
        with patch("streamlit.file_uploader", side_effect=uploader), \
                patch("geo_autostat.build_autostat_exports", wraps=build_autostat_exports) as prepare, \
                patch("geo_autostat.build_autostat_excel", wraps=build_autostat_excel) as excel, \
                patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertNotIn("result", app.session_state)  # Campaign Prepare has not been clicked.
            result = app.session_state["geo_autostat_result"]
            self.assertEqual(prepare.call_count, 1)
            excel.assert_not_called()
            self.assertEqual([file.count for file in result["bundle"].files], [3, 3])
            self.assertEqual([row["EMAIL/ NUMBER"] for row in log_rows(result["bundle"].files[0])],
                             ["second@example.com", "third@example.com", "first@example.com"])
            call = next(call for call in download.call_args_list if call.kwargs.get("key") == "geo_autostat_excel")
            self.assertEqual(call.args[2], "AUTOSTAT.xlsx")
            self.assertTrue(callable(call.args[1]))
            workbook = load_workbook(BytesIO(call.args[1]()))
            self.assertEqual(workbook.sheetnames, ["Email", "Viber"])
            for sheet in workbook:
                self.assertEqual([cell.value for cell in sheet["B"][1:]],
                                 ["000000000000000002", "000000000000000003", "000000000000000001"])
                self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in sheet["B"][1:]))
            workbook.close()
            app.run(timeout=20)
            self.assertEqual(prepare.call_count, 1)
            self.assertEqual(app.session_state["geo_autostat_result"]["timestamp"], result["timestamp"])
            next(widget for widget in app.selectbox if widget.label == "Source sheet").select("GEO").run(timeout=20)
            self.assertEqual(prepare.call_count, 2)
            self.assertEqual([row["Account No."] for row in log_rows(app.session_state["geo_autostat_result"]["bundle"].files[0])],
                             ["000000000000000019"])
            uploads[0] = Upload(workbook_data("000000000000000025.0"))  # Same filename, different contents.
            app.run(timeout=20)
            self.assertEqual(prepare.call_count, 3)
            self.assertEqual(log_rows(app.session_state["geo_autostat_result"]["bundle"].files[0])[-1]["Account No."], "000000000000000025")
            uploads[0] = None
            download.reset_mock()
            app.run(timeout=20)
            self.assertNotIn("geo_autostat_result", app.session_state)
            self.assertFalse(any(call.kwargs.get("key") == "geo_autostat_excel" for call in download.call_args_list))

    def test_manual_mapping_controls_identity_and_ob_instead_of_unselected_alias_columns(self):
        workbook = Workbook()
        workbook.active.title = "Custom"
        workbook.active.append(["Ledger", "Inbox", "Handset", "Chosen ID", "Chosen OB", "OB", "Old I.C."])
        for row in (
            ["001", "first@example.com", "09000000001", "U1", 100, 1, "old-a"],
            ["001", "second@example.com", "09000000002", "U2", 9, 500, "old-b"],
            ["001", "ignored@example.com", "09000000003", "U1", 999, 999, "old-c"],
        ):
            workbook.active.append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        source = mapped_contact_source("custom.xlsx", output.getvalue(), "Custom", 1, 7,
            {"primary_key": 0, "cardnumber": None, "email": 1, "phone": 2, "indexoneuid": 3, "OB": 4})
        bundle = build_autostat_exports([source], timestamp=None)
        self.assertEqual([row["EMAIL/ NUMBER"] for row in log_rows(bundle.files[0])], ["first@example.com", "second@example.com"])
        self.assertEqual(bundle.summary[0]["Duplicate Email rows removed"], 1)

    def test_missing_primary_keys_are_excluded_without_substituting_valid_cardnumbers(self):
        upload = Upload(workbook_data(first_account=""))
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: upload if kw.get("key") == "source_upload" else None):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            bundle = app.session_state["geo_autostat_result"]["bundle"]
            self.assertEqual([file.count for file in bundle.files], [2, 2])
            self.assertEqual(bundle.summary[0]["Missing account"], 2)
            for file in bundle.files:
                self.assertEqual([row["Account No."] for row in log_rows(file)],
                                 ["000000000000000002", "000000000000000003"])

    def test_log_settings_rebuild_automatically_and_scientific_accounts_clear_downloads(self):
        uploads = [Upload(workbook_data())]
        uploader = lambda label, **kw: uploads[0] if kw.get("key") == "source_upload" else None
        with patch("streamlit.file_uploader", side_effect=uploader), patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            app.checkbox(key="geo_autostat_blank_time").set_value(True).run(timeout=20)
            app.text_input(key="geo_autostat_remark_by").set_value("USER2").run(timeout=20)
            result = app.session_state["geo_autostat_result"]
            self.assertIsNone(result["timestamp"])
            self.assertTrue(all(row["Time"] == "" and row["Remark By"] == "USER2" for file in result["bundle"].files for row in log_rows(file)))
            uploads[0] = Upload(workbook_data("5.15604E+15"))
            download.reset_mock()
            app.run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertIn("scientific notation", app.session_state["geo_autostat_result"]["error"])
            self.assertNotIn("bundle", app.session_state["geo_autostat_result"])
            self.assertFalse(any(call.kwargs.get("key") == "geo_autostat_excel" for call in download.call_args_list))

    def test_excluded_uploads_do_not_leave_previous_autostat_available(self):
        uploads = [Upload(workbook_data())]
        uploader = lambda label, **kw: uploads[0] if kw.get("key") == "source_upload" else None
        with patch("streamlit.file_uploader", side_effect=uploader), patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            self.assertIn("bundle", app.session_state["geo_autostat_result"])
            for filename in ("SBC INSURANCE.xlsx", "NEGATIVE B2.xlsx", "NEGATIVE B4.xlsx"):
                uploads[0] = Upload(workbook_data(), filename)
                download.reset_mock()
                app.run(timeout=20)
                self.assertEqual(len(app.exception), 0)
                self.assertNotIn("geo_autostat_result", app.session_state)
                self.assertFalse(any(call.kwargs.get("key") == "geo_autostat_excel" for call in download.call_args_list))


if __name__ == "__main__":
    unittest.main()
