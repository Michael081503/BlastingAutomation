from __future__ import annotations

import csv
import unittest
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from automation import AutomationError
from autostat import SourceSelection, build_autostat_excel, build_autostat_excel_zip, inspect_contact_source
from b2_autostat import B2ActivitySource, b2_metadata_columns, build_b2_autostat_exports, build_b4_autostat_exports
from folder_output import MANILA_TIMEZONE


ROOT = Path(__file__).resolve().parents[1]
NOW = datetime(2026, 10, 30, 14, 0, 0, tzinfo=MANILA_TIMEZONE)
DATA = (
    b"cardnumber,primary_key,indexoneuid,FullName,email,phone,Collector,Terms,OB\n"
    b"WRONG-CARD-1,000000000000000001,ID1,First Person,one@example.com,+639171234567,OLD,AMRANCELI,9\n"
    b"WRONG-CARD-2,000000000000000002,ID2,Second Person,two@example.com,09181234567,OLD,KPALIMA,1000\n"
    b"WRONG-CARD-3,000000000000000003,ID3,Invalid contacts,bad..mail@example.com,123,OLD,OTHER,9000\n"
)


def source(data=DATA, name="NEGATIVE B4 - CARDS.csv"):
    sheet = inspect_contact_source(data, Path(name).suffix)[0]
    return B2ActivitySource(SourceSelection(name, data, sheet), *b2_metadata_columns(sheet.headers))


def rows(file):
    return list(csv.DictReader(StringIO(file.data.decode("utf-8-sig" if file.data.startswith(b"\xef\xbb\xbf") else "cp1252"))))


class Upload(BytesIO):
    def __init__(self, data=DATA, name="NEGATIVE B4 - CARDS.csv"):
        super().__init__(data)
        self.name = name


class B4AutostatTests(unittest.TestCase):
    def test_same_b2_template_values_with_b4_filenames_primary_keys_and_terms(self):
        b2 = build_b2_autostat_exports([source()], timestamp=NOW)
        b4 = build_b4_autostat_exports([source()], timestamp=NOW)
        self.assertEqual([file.filename for file in b4.files],
                         ["AUTOSTAT_B4_Email.csv", "AUTOSTAT_B4_SMS.csv", "AUTOSTAT_B4_Viber.csv"])
        self.assertEqual([file.headers for file in b4.files], [file.headers for file in b2.files])
        self.assertEqual([file.data for file in b4.files], [file.data for file in b2.files])
        self.assertEqual([file.count for file in b4.files], [2, 2, 2])
        for file in b4.files:
            header = "Account No." if file.filename.endswith("Email.csv") else "Account number"
            self.assertEqual([row[header] for row in rows(file)], ["000000000000000002", "000000000000000001"])
            self.assertTrue(all(row["REMARKS DATE"] == "10/30/2026 14:00:00" for row in rows(file)))
        self.assertEqual([row["Collector"] for row in rows(b4.files[0])], ["KPALIMA", "AMRANCELI"])
        self.assertEqual(rows(b4.files[0])[0]["REMARKS"], "two@example.com | EMAIL SENT | KPALIMA")
        self.assertEqual(rows(b4.files[1])[0]["Mobile Number"], "09181234567")

    def test_excel_and_zip_names_and_account_text_follow_b4_channels(self):
        bundle = build_b4_autostat_exports([source()], timestamp=NOW)
        workbook = load_workbook(BytesIO(build_autostat_excel(bundle)))
        self.assertEqual(workbook.sheetnames, ["B4_Email", "B4_SMS", "B4_Viber"])
        for sheet in workbook:
            column = "B" if sheet.title.endswith("Email") else "A"
            self.assertEqual([cell.value for cell in sheet[column][1:]], ["000000000000000002", "000000000000000001"])
            self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in sheet[column][1:]))
            self.assertEqual(sheet["F2"].value, "10/30/2026 14:00:00")
        workbook.close()
        with ZipFile(BytesIO(build_autostat_excel_zip(bundle))) as archive:
            self.assertEqual(archive.namelist(), ["AUTOSTAT_B4_Email.xlsx", "AUTOSTAT_B4_SMS.xlsx", "AUTOSTAT_B4_Viber.xlsx"])

    def test_full_endorsement_is_not_accepted_as_a_b4_campaign_file(self):
        data = b"Account Key,Contract Number,Bucket Code,Email (01),Phone (01)\nKEY,CARD,4,one@example.com,09171234567\n"
        with self.assertRaisesRegex(AutomationError, "B4 CARDS/BEL output files containing primary_key"):
            build_b4_autostat_exports([source(data, "endorsement.csv")], timestamp=None)


class B4AutostatPageTests(unittest.TestCase):
    def test_negative_page_prepares_b4_downloads_automatically_with_compact_options(self):
        uploads = {"b4_autostat_uploads": [Upload()]}
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: uploads.get(kw["key"])), \
                patch("autostat_page.build_b4_autostat_exports", wraps=build_b4_autostat_exports) as prepare, \
                patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(ROOT / "pages" / "NEGATIVE_B2.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(prepare.call_count, 1)
            self.assertEqual(len(app.metric), 0)
            self.assertFalse(any(widget.key == "b4_autostat_prepare" for widget in app.button))
            self.assertTrue(all(not panel.proto.expanded for panel in app.expander))
            bundle = app.session_state["b4_autostat_result"]["bundle"]
            self.assertEqual([file.count for file in bundle.files], [2, 2, 2])
            calls = {call.kwargs.get("key"): call for call in download.call_args_list}
            self.assertEqual(calls["b4_autostat_excel"].args[2], "AUTOSTAT_B4.xlsx")
            self.assertEqual(calls["b4_autostat_download_0"].args[2], "AUTOSTAT_B4_Email.xlsx")
            with ZipFile(BytesIO(calls["b4_autostat_zip"].args[1]())) as archive:
                self.assertEqual(archive.namelist(), ["AUTOSTAT_B4_Email.xlsx", "AUTOSTAT_B4_SMS.xlsx", "AUTOSTAT_B4_Viber.xlsx"])
            app.run(timeout=20)
            self.assertEqual(prepare.call_count, 1)
            app.selectbox(key="b4_autostat_output_format").select("CSV (.csv)").run(timeout=20)
            self.assertEqual(prepare.call_count, 2)
            self.assertEqual(download.call_args_list[-1].kwargs.get("key"), "b4_autostat_download_2")
            self.assertEqual(download.call_args_list[-1].args[2], "AUTOSTAT_B4_Viber.csv")

    def test_standalone_b4_mode_is_separate_from_b2_and_clears_all_uploads(self):
        uploads = {"autostat_b4_uploads": [Upload()]}
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: uploads.get(kw["key"])):
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            app.radio(key="autostat_upload_type").set_value("B4").run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertIn("autostat_b4_result", app.session_state)
            self.assertNotIn("autostat_b2_result", app.session_state)
            self.assertTrue(all(file.filename.startswith("AUTOSTAT_B4_") for file in app.session_state["autostat_b4_result"]["bundle"].files))
            app.button(key="clear_all_files").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertNotIn("autostat_b4_result", app.session_state)
            self.assertEqual(len(app.get("download_button")), 0)

    def test_b2_and_b4_panels_do_not_mix_and_b4_replacement_and_clear_refresh_logs(self):
        uploads = {"geo_b2_autostat_uploads": [Upload(DATA.replace(b"000000000000000001", b"B2-ACCOUNT"), "B2.csv")],
                   "b4_autostat_uploads": [Upload()]}
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: uploads.get(kw["key"])):
            app = AppTest.from_file(str(ROOT / "pages" / "NEGATIVE_B2.py")).run(timeout=20)
            b2 = app.session_state["geo_b2_autostat_result"]
            self.assertEqual([file.count for file in app.session_state["b4_autostat_result"]["bundle"].files], [2, 2, 2])
            uploads["b4_autostat_uploads"][0] = Upload(DATA.replace(b"000000000000000001", b"NEW-B4-ACCOUNT"))
            app.run(timeout=20)
            self.assertIn("NEW-B4-ACCOUNT", {row["Account No."] for row in rows(app.session_state["b4_autostat_result"]["bundle"].files[0])})
            self.assertEqual(app.session_state["geo_b2_autostat_result"]["timestamp"], b2["timestamp"])
            app.text_input(key="b4_autostat_remark_by").set_value("B4USER").run(timeout=20)
            app.button(key="clear_all_files").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertNotIn("b4_autostat_result", app.session_state)
            self.assertNotIn("geo_b2_autostat_result", app.session_state)
            self.assertEqual(len(app.get("download_button")), 0)
            uploads["b4_autostat_uploads_clear_1"] = [Upload()]
            app.run(timeout=20)
            self.assertEqual(app.text_input(key="b4_autostat_remark_by").value, "B4USER")
            self.assertIn("b4_autostat_result", app.session_state)


if __name__ == "__main__":
    unittest.main()
