from __future__ import annotations

import subprocess
import sys
import textwrap
import unittest
from pathlib import Path


PROJECT_DIRECTORY = Path(__file__).resolve().parents[1]


class StartupTests(unittest.TestCase):
    def run_isolated(self, script: str) -> None:
        result = subprocess.run(
            [sys.executable, "-c", textwrap.dedent(script)],
            cwd=PROJECT_DIRECTORY, capture_output=True, text=True, timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_normal_reruns_keep_module_classes_and_settings(self):
        self.run_isolated("""
            from app_startup import refresh_application_modules
            assert refresh_application_modules() == set()
            import automation, export_config, exports
            error_class = automation.AutomationError
            file_class = exports.ExportFile
            filenames = export_config.OUTPUT_FILENAMES
            assert refresh_application_modules() == set()
            import app
            assert automation.AutomationError is error_class
            assert app.AutomationError is error_class
            assert exports.ExportFile is file_class
            assert export_config.OUTPUT_FILENAMES is filenames
        """)

    def test_cached_autostat_refreshes_to_cardnumber_mapping(self):
        self.run_isolated("""
            import autostat, autostat_page
            del autostat.PREFERRED_ACCOUNT_HEADER
            autostat.AUTOSTAT_VERSION = 1
            autostat.ACCOUNT_ALIASES = ('primary_key', 'Account No.')
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert 'autostat' in refreshed
            assert 'autostat_page' in refreshed
            data = b'primary_key,cardnumber,email\\nOTHER,000123,first@example.com\\n'
            sheet = autostat.inspect_contact_source(data, '.csv')[0]
            contacts = autostat.extract_contacts(autostat.SourceSelection('contacts.csv', data, sheet))
            assert contacts.emails == [('000123', 'first@example.com')]
            assert autostat_page.AUTOSTAT_VERSION == autostat.AUTOSTAT_VERSION
            assert refresh_application_modules() == set()
        """)

    def test_cached_modules_refresh_for_mc_product_normalization_before_grouping(self):
        self.run_isolated("""
            import automation, exports
            del automation.PRODUCT_NORMALIZATION_VERSION
            automation.normalize_product = lambda value: value
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'automation', 'exports'} <= refreshed
            assert automation.normalize_product(' mc ') == 'CARDS'
            assert exports.normalize_product is automation.normalize_product
            records = [{'primary_key': '0001', 'email': 'one@example.com', 'phone': '09171234567',
                        'product': 'MC', 'level': 1, 'AREA': 'Mindanao'}]
            template = automation.create_default_csv_template()
            reco = exports.build_negative_exports(template, records, '', 1, '.csv')
            assert reco.counts['blank_cards_level_1'] == 1
            geo = exports.build_geo_exports(template, records, '', 1, '.csv')
            assert geo.counts['mindanao_blank_cards'] == 1
            assert all(file.preview[0]['product'] == 'CARDS' for bundle in (reco, geo) for file in bundle.files)
            assert automation.MAPPING_VERSION == 8
            assert refresh_application_modules() == set()
        """)

    def test_cached_geo_loan_filenames_and_labels_refresh_together(self):
        self.run_isolated("""
            import export_config, exports
            del export_config.GEO_LOANS_FILENAMES_VERSION
            del exports.GEO_LOANS_VERSION
            export_config.GEO_OUTPUT_FILENAMES['mindanao_blank_other_products'] = 'Wrong CARDS destination'
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'export_config', 'exports'} <= refreshed
            assert export_config.GEO_OUTPUT_FILENAMES['mindanao_blank_other_products'] == 'GEO NEGA LOANS - MINDANAO'
            assert exports.GEO_OUTPUT_FILENAMES is export_config.GEO_OUTPUT_FILENAMES
            assert 'LOANS' in exports.GEO_PRODUCT_GROUP_LABELS['blank_other_products']
            assert exports.GEO_FILENAME_CONFIG_VERSION == 3
            record = {'primary_key': '0001', 'email': 'one@example.com', 'product': 'CSL',
                      'AREA': 'Mindanao', 'ChristmasAmnesty': 'QUALIFIED'}
            groups = exports.separate_geo([record])
            assert groups['mindanao_blank_other_products'] == [record]
            assert groups['mindanao_qualified_review'] == []
            assert refresh_application_modules() == set()
        """)

    def test_cached_modules_refresh_for_ob_ordering(self):
        self.run_isolated("""
            import automation, autostat, autostat_page, exports
            del automation.OB_SORT_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'automation', 'exports', 'autostat', 'autostat_page'} <= refreshed
            data = b'cardnumber,email,OB\\n0001,low@example.com,9\\n0002,high@example.com,1000\\n'
            source = autostat.SourceSelection('source.csv', data, autostat.inspect_contact_source(data, '.csv')[0])
            bundle = autostat_page.build_autostat_exports([source], timestamp=None)
            assert [row['Account No.'] for row in bundle.files[0].preview] == ['0002', '0001']
            records = [
                {'primary_key': 'low', 'email': 'low@example.com', 'OB': '9'},
                {'primary_key': 'high', 'email': 'high@example.com', 'OB': '1,000'},
            ]
            bundle = exports.build_sheet_exports(
                automation.create_default_csv_template(), records, '', 1,
                base_name='RECO', label='RECO', template_extension='.csv',
            )
            assert [row['primary_key'] for row in bundle.files[0].preview] == ['high', 'low']
            assert refresh_application_modules() == set()
        """)

    def test_cached_autostat_refreshes_for_exact_whole_number_account_text(self):
        self.run_isolated("""
            import autostat, autostat_page
            del autostat.ACCOUNT_NUMBER_NORMALIZATION_VERSION
            autostat.AUTOSTAT_VERSION = 10
            autostat.account_number_text = lambda value: value
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'autostat', 'autostat_page'} <= refreshed
            data = b'cardnumber,email\\n5156041234567897.0,first@example.com\\n000123.00,second@example.com\\n'
            source = autostat.SourceSelection('accounts.csv', data, autostat.inspect_contact_source(data, '.csv')[0])
            bundle = autostat_page.build_autostat_exports([source], timestamp=None)
            assert [row['Account No.'] for row in bundle.files[0].preview] == ['5156041234567897', '000123']
            assert autostat_page.AUTOSTAT_VERSION == 13
            assert refresh_application_modules() == set()
        """)

    def test_cached_autostat_refreshes_to_skip_bad_source_rows_in_all_active_workflows(self):
        self.run_isolated("""
            import autostat, geo_autostat, b2_autostat, autostat_page
            del autostat.SOURCE_ROW_SKIP_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'autostat', 'geo_autostat', 'b2_autostat', 'autostat_page'} <= refreshed
            data = b'primary_key,cardnumber,email\\n0001,5.15604E+15,badcard@example.com\\n5.15604E+15,0002,badprimary@example.com\\n0003,0004,valid@example.com\\n'
            source = autostat.SourceSelection('source.csv', data, autostat.inspect_contact_source(data, '.csv')[0])
            bundle = autostat_page.build_autostat_exports([source], timestamp=None)
            assert [row['Account No.'] for row in bundle.files[0].preview] == ['0002', '0004']
            assert bundle.summary[0]['Source error rows skipped'] == 1
            geo_source = geo_autostat.mapped_contact_source('source.csv', data, 'CSV', 1, 3,
                {'primary_key': 0, 'cardnumber': 1, 'email': 2})
            bundle = geo_autostat.build_autostat_exports([geo_source], timestamp=None)
            assert [row['Account No.'] for row in bundle.files[0].preview] == ['0002', '0004']
            bundle = b2_autostat.build_b2_autostat_exports([b2_autostat.B2ActivitySource(source)], timestamp=None)
            assert [row['Account No.'] for row in bundle.files[0].preview] == ['0001', '0003']
            assert bundle.summary[0]['Source error rows skipped'] == 1
            assert autostat.AUTOSTAT_VERSION == 13
            assert autostat_page.build_autostat_exports is autostat.build_autostat_exports
            assert refresh_application_modules() == set()
        """)

    def test_cached_autostat_refreshes_for_bel_split_and_filenames(self):
        self.run_isolated("""
            from datetime import datetime
            import autostat, autostat_page, export_config
            del export_config.AUTOSTAT_BEL_FILENAMES
            del autostat.build_bel_exports
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'export_config', 'autostat', 'autostat_page'} <= refreshed
            data = b'primary_key,email\\n@BEL-1,first@example.com\\n'
            sheet = autostat.inspect_contact_source(data, '.csv')[0]
            bundle = autostat.build_bel_exports(
                [autostat.SourceSelection('contacts.csv', data, sheet)], timestamp=datetime(2026, 10, 28),
            )
            assert [file.filename for file in bundle.files] == ['NEGATIVE B4 - BEL.csv', 'NEGATIVE B4 - CARDS.csv']
            assert autostat_page.build_autostat_exports is autostat.build_autostat_exports
            assert refresh_application_modules() == set()
        """)

    def test_cached_autostat_refreshes_for_account_deduplication(self):
        self.run_isolated("""
            from datetime import datetime
            import autostat, autostat_page
            del autostat.ACCOUNT_DEDUP_VERSION
            del autostat.unique_account_contacts
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'autostat', 'autostat_page'} <= refreshed
            data = b'cardnumber,email,phone\\n0001,first@example.com,09171234567\\n0001,next@example.com,09181234567\\n'
            sheet = autostat.inspect_contact_source(data, '.csv')[0]
            bundle = autostat.build_autostat_exports(
                [autostat.SourceSelection('contacts.csv', data, sheet)], timestamp=datetime(2026, 10, 28),
            )
            assert [file.count for file in bundle.files] == [1, 1]
            assert autostat_page.AUTOSTAT_VERSION == autostat.AUTOSTAT_VERSION
            assert refresh_application_modules() == set()
        """)

    def test_cached_autostat_refreshes_for_text_formatted_excel_download(self):
        self.run_isolated("""
            from io import BytesIO
            from openpyxl import load_workbook
            import autostat, autostat_page
            del autostat.ACCOUNT_TEXT_EXPORT_VERSION
            del autostat.build_autostat_excel
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'autostat', 'autostat_page'} <= refreshed
            data = b'cardnumber,email\\n5156041234567897,first@example.com\\n'
            sheet = autostat.inspect_contact_source(data, '.csv')[0]
            bundle = autostat.build_autostat_exports(
                [autostat.SourceSelection('source.csv', data, sheet)], timestamp=None,
            )
            workbook = load_workbook(BytesIO(autostat_page.build_autostat_excel(bundle)))
            assert workbook['Email']['B2'].value == '5156041234567897'
            assert workbook['Email']['B2'].data_type == 's'
            workbook.close()
            assert refresh_application_modules() == set()
        """)

    def test_cached_autostat_refreshes_to_two_compiled_channel_files(self):
        self.run_isolated("""
            from datetime import datetime
            import autostat, autostat_page
            del autostat.COMPILED_OUTPUT_VERSION
            autostat.AUTOSTAT_VERSION = 5
            autostat_page.build_autostat_exports = lambda *args, **kwargs: None
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'autostat', 'autostat_page'} <= refreshed
            contents = [
                b'cardnumber,email,phone\\n0001,first@example.com,\\n',
                b'cardnumber,email,phone\\n0001,later@example.com,09171234567\\n0002,second@example.com,09181234567\\n',
            ]
            sources = [autostat.SourceSelection(
                f'source-{index}.csv', data, autostat.inspect_contact_source(data, '.csv')[0],
            ) for index, data in enumerate(contents)]
            bundle = autostat_page.build_autostat_exports(sources, timestamp=datetime(2026, 10, 28))
            assert [file.filename for file in bundle.files] == ['AUTOSTAT_Email.csv', 'AUTOSTAT_Viber.csv']
            assert [file.count for file in bundle.files] == [2, 2]
            assert len(bundle.summary) == 2
            assert autostat.AUTOSTAT_VERSION == 13
            assert autostat_page.AUTOSTAT_PAGE_VERSION == 10
            assert refresh_application_modules() == set()
        """)

    def test_cached_old_configuration_is_refreshed_before_app_import(self):
        self.run_isolated("""
            import automation, export_config, exports, folder_output
            error_class = automation.AutomationError
            for name in ('GEO_OUTPUT_FILENAMES', 'GEO_ZIP_FILENAME'):
                delattr(export_config, name)
            try:
                from export_config import GEO_OUTPUT_FILENAMES
            except ImportError:
                pass
            else:
                raise AssertionError('Expected the old configuration import to fail')
            import app
            assert app.GEO_OUTPUT_FILENAMES is export_config.GEO_OUTPUT_FILENAMES
            assert exports.GEO_OUTPUT_FILENAMES is export_config.GEO_OUTPUT_FILENAMES
            assert app.AutomationError is error_class
            assert folder_output.AutomationError is error_class
            from app_startup import refresh_application_modules
            assert refresh_application_modules() == set()
        """)

    def test_cached_pre_geo_modules_refresh_together_and_export(self):
        self.run_isolated("""
            import csv
            from io import StringIO
            import automation, export_config, exports, folder_output
            for name in ('GEO_OUTPUT_FILENAMES', 'GEO_ZIP_FILENAME'):
                delattr(export_config, name)
            for name in ('MAPPING_VERSION', 'OPTIONAL_FIELD_MAPPING'):
                delattr(automation, name)
            for name in ('GEO_GROUP_LABELS', 'GEO_REVIEW_GROUP_KEYS', 'build_geo_exports'):
                delattr(exports, name)
            import app
            assert exports.AutomationError is automation.AutomationError
            assert folder_output.AutomationError is automation.AutomationError
            assert app.AutomationError is automation.AutomationError
            records = [{
                'email': 'ana@example.com', 'phone': '09171234567',
                'primary_key': '0001', 'AREA': 'Region 3', 'ChristmasAmnesty': 'QUALIFIED',
                'product': 'CARDS', 'level': 'Level 6',
            }]
            bundle = app.build_geo_exports(
                automation.create_default_csv_template(), records, '', 1, '.csv',
            )
            assert len(bundle.files) == 2
            assert all(file.group_key == 'r3_qualified_cards' for file in bundle.files)
            for file in bundle.files:
                rows = list(csv.DictReader(StringIO(file.data.decode('cp1252'))))
                assert rows[0]['AREA'] == 'Region 3'
                assert rows[0]['primary_key'] == '0001'
        """)

    def test_cached_level_grouping_and_old_bootstrap_refresh_to_product_groups(self):
        self.run_isolated("""
            import app_startup, automation, export_config, exports
            error_class = automation.AutomationError
            export_config.GEO_OUTPUT_FILENAMES = {
                'qualified_cards_level_1': 'Old Level 1',
                'qualified_cards_level_6': 'Old Level 6',
            }
            del exports.GEO_PRODUCT_GROUP_LABELS
            app_startup.refresh_application_modules = lambda: set()
            import app
            assert 'mindanao_qualified_cards' in app.GEO_OUTPUT_FILENAMES
            assert 'r3_qualified_cards' in app.GEO_OUTPUT_FILENAMES
            assert 'qualified_cards_level_1' not in app.GEO_OUTPUT_FILENAMES
            assert app.AutomationError is error_class
            records = [
                {'email': 'ana@example.com', 'phone': '09171234567', 'AREA': 'Mindanao',
                 'ChristmasAmnesty': 'QUALIFIED', 'product': 'CARDS', 'level': level}
                for level in (1, 2, 3, 4, 5, 6)
            ]
            bundle = app.build_geo_exports(automation.create_default_csv_template(), records, '', 1, '.csv')
            assert len(bundle.files) == 2
            assert all(file.group_key == 'mindanao_qualified_cards' and file.count == 6 for file in bundle.files)
        """)

    def test_cached_modules_refresh_for_filename_config_and_email_counts(self):
        self.run_isolated("""
            import automation, exports, folder_output
            del automation.MappedData.valid_email_rows
            del automation.MappedData.missing_email_rows
            del exports.GEO_FILENAME_CONFIG_VERSION
            import app
            assert hasattr(automation.MappedData, 'valid_email_rows')
            assert hasattr(automation.MappedData, 'missing_email_rows')
            assert app.GEO_FILENAME_CONFIG_VERSION == 3
            assert exports.AutomationError is automation.AutomationError
            assert folder_output.AutomationError is automation.AutomationError
        """)

    def test_cached_modules_refresh_for_geo_doclass_and_export_preview(self):
        self.run_isolated("""
            import csv
            from io import StringIO
            import automation, exports, folder_output
            del automation.geo_doclass
            del automation.DERIVED_TEMPLATE_HEADERS
            import app
            assert app.MAPPING_VERSION == 8
            assert 'doclass' in app.DERIVED_TEMPLATE_HEADERS
            assert exports.geo_doclass is automation.geo_doclass
            assert folder_output.AutomationError is automation.AutomationError
            bundle = exports.build_sheet_exports(
                automation.create_default_csv_template(),
                [{'email': 'first@example.com', 'AREA': 'Mindanao'}],
                '', 1, template_extension='.csv', base_name='GEO', label='GEO',
            )
            rows = list(csv.DictReader(StringIO(bundle.files[0].data.decode('cp1252'))))
            assert rows[0]['doclass'] == 'SBC GEO Mindanao'
            assert bundle.files[0].preview[0]['doclass'] == 'SBC GEO Mindanao'
            from app_startup import refresh_application_modules
            assert refresh_application_modules() == set()
        """)

    def test_cached_exports_refresh_for_grouping_other_source_sheets(self):
        self.run_isolated("""
            import automation, exports, folder_output
            error_class = automation.AutomationError
            del exports.GROUPED_SHEET_EXPORT_VERSION
            import app
            assert app.GROUPED_SHEET_EXPORT_VERSION == 1
            assert exports.AutomationError is error_class
            assert folder_output.AutomationError is error_class
            bundle = app.build_negative_exports(
                automation.create_default_csv_template(),
                [{'email': 'first@example.com', 'product': 'CARDS', 'level': 1,
                  'ChristmasAmnesty': 'QUALIFIED'}],
                '', 1, '.csv', source_sheet='POSITIVE',
            )
            assert all(file.filename.startswith('POSITIVE - ') for file in bundle.files)
            assert bundle.counts['qualified_cards_level_1'] == 1
            from app_startup import refresh_application_modules
            assert refresh_application_modules() == set()
        """)

    def test_cached_modules_refresh_for_partial_source_headers(self):
        self.run_isolated("""
            from io import BytesIO
            from openpyxl import Workbook
            import automation, exports
            del automation.SheetInfo.available_fields
            del automation.inspect_source_sheet
            import app
            assert app.MAPPING_VERSION == 8
            assert exports.AutomationError is automation.AutomationError
            workbook = Workbook()
            workbook.active.append(['primary_key', 'email', 'PHONE'])
            workbook.active.append(['0001', 'first@example.com', '09171234567'])
            output = BytesIO()
            workbook.save(output)
            workbook.close()
            data = output.getvalue()
            info = automation.inspect_workbook(data)[0]
            assert set(info.available_fields) == {'primary_key', 'email', 'phone'}
            mapped = app.read_source(data, info.name, 1, allow_missing=True)
            assert mapped.records[0]['primary_key'] == '0001'
            assert mapped.records[0]['phone'] == '09171234567'
            assert mapped.records[0]['product'] is None
        """)


    def test_cached_modules_refresh_for_b2_matching_and_workbook_dimensions(self):
        self.run_isolated("""
            import automation, autostat, b2_matching, b2_page, export_config
            del automation.WORKBOOK_READER_VERSION
            del export_config.AUTOSTAT_B2_FILENAMES
            del b2_matching.build_b2_exports
            import app
            assert app.MAPPING_VERSION == 8
            assert automation.WORKBOOK_READER_VERSION == 1
            assert b2_page.build_b2_exports is b2_matching.build_b2_exports
            assert export_config.AUTOSTAT_B2_FILENAMES['CARDS'] == 'NEGATIVE B2 - CARDS'
            assert autostat.AutomationError is automation.AutomationError
            from app_startup import refresh_application_modules
            assert refresh_application_modules() == set()
        """)

    def test_cached_b2_page_refreshes_for_blasting_template(self):
        self.run_isolated("""
            import inspect
            import b2_matching, b2_page
            del b2_matching.DATAGRID_ROUTING_VERSION
            del b2_matching.B2_TEMPLATE_VERSION
            b2_page.render_b2_page = lambda: None
            import app
            assert b2_matching.is_datagrid_sheet(['Account No.', 'Financing/Card No.', 'Status Code'])
            assert b2_matching.B2_MATCHING_VERSION == 6
            assert len(b2_page.B2_TEMPLATE_HEADERS) == 16
            assert 'Time' not in b2_page.B2_TEMPLATE_HEADERS
            assert 'timestamp' not in inspect.signature(b2_matching.build_b2_exports).parameters
            from app_startup import refresh_application_modules
            assert refresh_application_modules() == set()
        """)

    def test_cached_b2_modules_refresh_for_upload_role_detection(self):
        self.run_isolated("""
            import b2_matching, b2_page
            del b2_matching.B2_UPLOAD_ROLES_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'b2_matching', 'b2_page'} <= refreshed
            assert b2_page.is_sbc_database_sheet(['Account Key', 'Contract Number', 'Bucket Code', 'Product Category Code'])
            assert b2_page.validate_b2_source is b2_matching.validate_b2_source
            assert refresh_application_modules() == set()
        """)

    def test_cached_autostat_page_refreshes_after_b2_moves_to_its_own_page(self):
        self.run_isolated("""
            import autostat_page
            del autostat_page.AUTOSTAT_PAGE_VERSION
            autostat_page.render_autostat = lambda: None
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert 'autostat_page' in refreshed
            assert autostat_page.AUTOSTAT_PAGE_VERSION == 10
            assert autostat_page.render_autostat.__name__ == 'render_autostat'
            assert refresh_application_modules() == set()
        """)


    def test_cached_b2_activity_templates_and_upload_page_refresh_together(self):
        self.run_isolated("""
            import b2_autostat, autostat_page
            del b2_autostat.B2_COLLECTOR_TERMS_VERSION
            del autostat_page.B2_UPLOAD_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'b2_autostat', 'autostat_page'} <= refreshed
            assert autostat_page.build_b2_autostat_exports is b2_autostat.build_b2_autostat_exports
            assert autostat_page.B2_UPLOAD_VERSION == 1
            assert b2_autostat.B2_AUTOSTAT_VERSION == 5
            data = b'cardnumber,primary_key,email,phone,Collector,Terms\\nCARD,000123,one@example.com,09171234567,OLD,AGENT1\\n'
            import autostat
            source = autostat.SourceSelection('B2.csv', data, autostat.inspect_contact_source(data, '.csv')[0])
            selection = b2_autostat.B2ActivitySource(source, *b2_autostat.b2_metadata_columns(source.sheet.headers))
            bundle = b2_autostat.build_b2_autostat_exports([selection], timestamp=None)
            assert bundle.files[0].preview[0]['Account No.'] == '000123'
            assert bundle.files[1].preview[0]['Account number'] == '000123'
            assert bundle.files[2].preview[0]['Account number'] == '000123'
            assert bundle.files[0].preview[0]['Collector'] == 'AGENT1'
            assert refresh_application_modules() == set()
        """)

    def test_cached_geo_financing_mapping_and_negative_b2_upload_panel_refresh(self):
        self.run_isolated("""
            import geo_autostat, autostat_page
            del geo_autostat.GEO_FINANCING_ACCOUNT_VERSION
            del autostat_page.B2_ACTIVITY_PANEL_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'geo_autostat', 'autostat_page'} <= refreshed
            assert callable(autostat_page.render_b2_autostat_uploads)
            source = geo_autostat.mapped_contact_source('source.csv', b'', 'CSV', 1, 3,
                {'primary_key': 0, 'cardnumber': 1, 'email': 2})
            assert source.sheet.account_column == 1
            assert geo_autostat.GEO_AUTOSTAT_VERSION == 4
            assert refresh_application_modules() == set()
        """)

    def test_cached_pages_refresh_for_the_shared_clear_all_files_action(self):
        self.run_isolated("""
            import branding, b2_page, autostat_page, upload_state
            del upload_state.CLEAR_UPLOADS_VERSION
            del branding.render_clear_all_files
            del b2_page.upload_widget_key
            del autostat_page.upload_widget_key
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'upload_state', 'branding', 'b2_page', 'autostat_page'} <= refreshed
            assert branding.render_clear_all_files is upload_state.render_clear_all_files
            assert b2_page.upload_widget_key is upload_state.upload_widget_key
            assert autostat_page.upload_widget_key is upload_state.upload_widget_key
            assert refresh_application_modules() == set()
        """)


    def test_cached_b4_matching_agent_names_and_uploads_refresh_together(self):
        self.run_isolated("""
            import export_config, b2_matching, b2_page, upload_state
            del export_config.AUTOSTAT_B4_FILENAMES
            del b2_matching.B4_MATCHING_VERSION
            del b2_page.B4_PAGE_VERSION
            del upload_state.B4_UPLOADS_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'export_config', 'b2_matching', 'b2_page', 'upload_state'} <= refreshed
            assert b2_page.build_b4_exports is b2_matching.build_b4_exports
            assert b2_matching.B2_TEMPLATE_VERSION == 2
            assert export_config.AUTOSTAT_B4_FILENAMES['CARDS'] == 'NEGATIVE B4 - CARDS'
            assert export_config.COLLECTION_AGENTS == (('Almera Ranceli', 'AMRANCELI'), ('Kyarill Palima', 'KPALIMA'))
            assert 'b4_database_upload' in upload_state.UPLOAD_KEYS
            assert 'b4_grid_upload' in upload_state.UPLOAD_KEYS
            assert refresh_application_modules() == set()
        """)


    def test_cached_geo_reco_exports_refresh_for_consistent_output_format(self):
        self.run_isolated("""
            import exports
            del exports.CAMPAIGN_OUTPUT_FORMAT_VERSION
            del exports.campaign_output_files
            from app_startup import refresh_application_modules
            assert 'exports' in refresh_application_modules()
            import app
            assert app.campaign_output_files is exports.campaign_output_files
            assert refresh_application_modules() == set()
        """)

    def test_cached_geo_reco_exports_refresh_for_native_excel_types(self):
        self.run_isolated("""
            import automation, exports
            del automation.write_typed_template_sheet
            del exports.TYPED_EXCEL_EXPORT_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'automation', 'exports'} <= refreshed
            assert exports.EXCEL_EXPORT_VERSION == 2
            assert 'records' in exports.ExportFile.__dataclass_fields__
            from io import BytesIO
            from openpyxl import load_workbook
            bundle = exports.build_sheet_exports(automation.create_default_csv_template(),
                [{'primary_key': '000123', 'email': 'one@example.com', 'OB': 123.5}], '', 1,
                base_name='RECO', label='RECO', template_extension='.csv')
            workbook = load_workbook(BytesIO(exports.build_export_excel(bundle.files[0])))
            assert workbook.active['B2'].value == '000123'
            assert workbook.active['B2'].number_format == '@'
            assert workbook.active['G2'].value == 123.5
            assert workbook.active['G2'].data_type == 'n'
            assert workbook.active['G2'].number_format == 'General'
            workbook.close()
            assert refresh_application_modules() == set()
        """)


    def test_cached_geo_autostat_refreshes_for_separate_email_and_viber_downloads(self):
        self.run_isolated("""
            import geo_autostat, autostat
            del geo_autostat.GEO_AUTOSTAT_CHANNEL_DOWNLOADS_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert 'geo_autostat' in refreshed
            assert geo_autostat.build_autostat_channel_excel is autostat.build_autostat_channel_excel
            assert geo_autostat.build_autostat_excel_zip is autostat.build_autostat_excel_zip
            assert geo_autostat.GEO_AUTOSTAT_VERSION == 4
            assert refresh_application_modules() == set()
        """)


    def test_cached_modules_refresh_for_reco_only_agent_assignment_fields(self):
        self.run_isolated("""
            from datetime import date
            import automation, agent_cycle, exports, agent_cycle_page
            del automation.ASSIGNMENT_TEMPLATE_HEADERS
            del exports.AGENT_ASSIGNMENT_EXPORT_VERSION
            del exports.RECO_ONLY_AGENTS_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'automation', 'agent_cycle', 'exports', 'agent_cycle_page'} <= refreshed
            agents = agent_cycle.cycle_agents(agent_cycle.default_cycle(), date(2026, 10, 31))
            assert agents[0] == ('Rosalie Vinoya', 'ROSVINOYA')
            bundle = exports.build_negative_exports(automation.create_default_csv_template(),
                [{'primary_key': '0001', 'email': 'one@example.com', 'product': 'CARDS', 'level': 1}],
                '', 1, '.csv', agents=agents)
            assert bundle.files[0].preview[0]['Monthly'] == 'Rosalie Vinoya'
            assert bundle.files[0].preview[0]['Terms'] == 'ROSVINOYA'
            assert b'Rosalie Vinoya,ROSVINOYA' in bundle.files[0].data
            geo = exports.build_geo_exports(automation.create_default_csv_template(),
                [{'primary_key': '0001', 'email': 'one@example.com', 'product': 'CARDS', 'AREA': 'Mindanao'}],
                '', 1, '.csv', agents=agents)
            assert all('Terms' not in row and 'Monthly' not in row for file in geo.files for row in file.preview)
            assert agent_cycle_page.AGENT_CYCLE_PAGE_VERSION == 2
            assert refresh_application_modules() == set()
        """)


    def test_cached_bucket_pages_refresh_for_automatic_downloads_and_bucket_errors(self):
        self.run_isolated("""
            import b2_matching, b2_page
            del b2_matching.B2_BUCKET_DIAGNOSTICS_VERSION
            del b2_page.B2_COMPACT_DOWNLOADS_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'b2_matching', 'b2_page'} <= refreshed
            assert b2_matching.B2_MATCHING_VERSION == 6
            assert b2_page.B2_COMPACT_DOWNLOADS_VERSION == 1
            assert b2_page.B4_PAGE_VERSION == 2
            assert callable(b2_page.no_matching_rows_message)
            assert refresh_application_modules() == set()
        """)


    def test_cached_autostat_pages_refresh_for_b4_templates_uploads_and_clear_all(self):
        self.run_isolated("""
            import b2_autostat, autostat_page, upload_state
            del b2_autostat.B4_AUTOSTAT_VERSION
            del autostat_page.B4_ACTIVITY_PANEL_VERSION
            del upload_state.B4_AUTOSTAT_UPLOADS_VERSION
            from app_startup import refresh_application_modules
            refreshed = refresh_application_modules()
            assert {'b2_autostat', 'autostat_page', 'upload_state'} <= refreshed
            assert callable(b2_autostat.build_b4_autostat_exports)
            assert callable(autostat_page.render_b4_autostat_uploads)
            assert {'autostat_b4_uploads', 'b4_autostat_uploads'} <= set(upload_state.UPLOAD_KEYS)
            assert {'autostat_b4_result', 'b4_autostat_result'} <= upload_state.FILE_STATE_KEYS
            assert autostat_page.AUTOSTAT_PAGE_VERSION == 10
            assert refresh_application_modules() == set()
        """)


if __name__ == "__main__":
    unittest.main()
