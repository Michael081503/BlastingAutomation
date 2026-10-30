from __future__ import annotations

import csv
import unittest
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from automation import DEFAULT_TEMPLATE_HEADERS, create_default_csv_template
from autostat import SourceSelection, build_autostat_exports, inspect_contact_source
from exports import build_excel_zip, build_export_excel, build_geo_exports, build_negative_exports, build_sheet_exports


class ExcelExportTests(unittest.TestCase):
    def test_excel_download_preserves_full_rows_and_literal_identifiers_through_autostat(self):
        accounts = ["5156041234567897", "001234567890123456", "123456789012345678901234", "=2+2",
                    *[f"000000{index:012d}" for index in range(110)]]
        records = [{
            "primary_key": f"KEY-{index}", "cardnumber": account, "indexoneuid": f"UID-{index}",
            "email": f"account{index}@example.com", "phone": "09171234567", "OB": index,
        } for index, account in enumerate(accounts)]
        for encoding in ("cp1252", "utf-8-sig"):
            with self.subTest(encoding=encoding):
                bundle = build_sheet_exports(
                    create_default_csv_template(), records, "", 1, base_name="Accounts", label="Accounts",
                    template_extension=".csv", encoding=encoding,
                )
                original_csvs = [file.data for file in bundle.files]
                with ZipFile(BytesIO(build_excel_zip(bundle.files))) as archive:
                    self.assertEqual(archive.namelist(), ["Accounts_Email.xlsx", "Accounts_Viber.xlsx"])
                    for file in bundle.files:
                        excel_data = archive.read(Path(file.filename).with_suffix(".xlsx").name)
                        workbook = load_workbook(BytesIO(excel_data))
                        sheet = workbook.active
                        raw_rows = list(csv.reader(StringIO(file.data.decode(encoding))))
                        expected = [tuple(raw_rows[0])]
                        ob_column = file.headers.index("OB")
                        expected.extend(tuple(int(value) if column == ob_column else value or None
                                              for column, value in enumerate(row)) for row in raw_rows[1:])
                        self.assertEqual(list(sheet.values), expected)
                        self.assertEqual(sheet.max_row, len(accounts) + 1)
                        for row in sheet.iter_rows(min_row=2):
                            for column in ("primary_key", "cardnumber", "indexoneuid"):
                                cell = row[file.headers.index(column)]
                                self.assertEqual(cell.data_type, "s")
                                self.assertEqual(cell.number_format, "@")
                        workbook.close()
                        selection = SourceSelection("Accounts.xlsx", excel_data, inspect_contact_source(excel_data, ".xlsx")[0])
                        logs = build_autostat_exports([selection], timestamp=None)
                        log = next(log for log in logs.files if log.count)
                        actual = list(csv.DictReader(StringIO(log.data.decode("cp1252"))))
                        self.assertEqual([row["Account No."] for row in actual], list(reversed(accounts)))
                        self.assertEqual(sum(row["Account precision issues"] for row in logs.summary), 0)
                self.assertEqual([file.data for file in bundle.files], original_csvs)

    def test_geo_reco_excel_keeps_numeric_and_date_types_and_general_text_with_reordered_headers(self):
        headers = [*reversed(DEFAULT_TEMPLATE_HEADERS), "Unused column"]
        template_output = StringIO(newline="")
        csv.writer(template_output).writerow(headers)
        template = template_output.getvalue().encode()
        records = [
            {"primary_key": "0005156041234567897", "cardnumber": "000001234567890123",
             "indexoneuid": "000045", "email": "first@example.com", "phone": "09171234567",
             "FullName": "=literal text", "OB": 1234.5, "level": 1, "product": "CARDS",
             "endodate": datetime(2026, 10, 30, 13, 14, 15), "AREA": "Mindanao", "ChristmasAmnesty": "QUALIFIED"},
            {"primary_key": "000002", "cardnumber": "000000000000000002", "indexoneuid": "000046",
             "email": "second@example.com", "phone": "09181234567", "FullName": "Second Person",
             "OB": Decimal("-2.25"), "level": 1, "product": "CARDS", "endodate": date(2026, 10, 29),
             "AREA": "Mindanao", "ChristmasAmnesty": "QUALIFIED"},
        ]
        for builder in (build_geo_exports, build_negative_exports):
            with self.subTest(workflow=builder.__name__):
                bundle = builder(template, records, "", 1, template_extension=".csv")
                original_csvs = [file.data for file in bundle.files]
                with ZipFile(BytesIO(build_excel_zip(bundle.files))) as archive:
                    for file in bundle.files:
                        workbook = load_workbook(BytesIO(archive.read(Path(file.filename).with_suffix(".xlsx").name)))
                        sheet = workbook.active
                        self.assertEqual([cell.value for cell in sheet[1]], headers)
                        self.assertTrue(all(cell.number_format == "General" for cell in sheet[1]))
                        first = dict(zip(headers, sheet[2]))
                        second = dict(zip(headers, sheet[3]))
                        for header in ("primary_key", "cardnumber", "indexoneuid"):
                            self.assertEqual(first[header].value, records[0][header])
                            self.assertEqual((first[header].data_type, first[header].number_format), ("s", "@"))
                        self.assertEqual((first["OB"].value, first["OB"].data_type, first["OB"].number_format),
                                         (1234.5, "n", "General"))
                        self.assertEqual((second["OB"].value, second["OB"].data_type), (-2.25, "n"))
                        self.assertEqual((first["Level"].value, first["Level"].data_type, first["Level"].number_format),
                                         (1, "n", "General"))
                        self.assertEqual((first["EndoDate"].value, first["EndoDate"].data_type),
                                         (datetime(2026, 10, 30, 13, 14, 15), "d"))
                        self.assertEqual(second["EndoDate"].value, datetime(2026, 10, 29))
                        self.assertEqual(first["FullName"].value, "=literal text")
                        self.assertEqual(first["FullName"].data_type, "s")
                        for header in ("FullName", "Product", "AREA", "doclass", "email"):
                            self.assertEqual(first[header].number_format, "General")
                        self.assertEqual(first["doclass"].value, "SBC GEO Mindanao")
                        if file.channel == "viber":
                            self.assertEqual((first["phone"].value, first["phone"].data_type, first["phone"].number_format),
                                             ("09171234567", "s", "@"))
                            self.assertIsNone(first["email"].value)
                        else:
                            self.assertIsNone(first["phone"].value)
                        self.assertIsNone(first["Unused column"].value)
                        workbook.close()
                self.assertEqual([file.data for file in bundle.files], original_csvs)

    def test_empty_excel_channel_keeps_headers_with_general_format(self):
        bundle = build_sheet_exports(create_default_csv_template(),
            [{"primary_key": "0001", "email": "one@example.com", "OB": 100}], "", 1,
            base_name="RECO", label="RECO", template_extension=".csv")
        file = next(file for file in bundle.files if file.channel == "viber")
        workbook = load_workbook(BytesIO(build_export_excel(file)))
        self.assertEqual(workbook.active.max_row, 1)
        self.assertEqual([cell.value for cell in workbook.active[1]], list(DEFAULT_TEMPLATE_HEADERS))
        self.assertTrue(all(cell.number_format == "General" for cell in workbook.active[1]))
        workbook.close()

    def test_main_page_offers_excel_zip_and_individual_workbooks_without_building_them_on_rerun(self):
        # A real source workbook exercises the page's normal preparation flow.
        workbook = Workbook()
        workbook.active.append(["cardnumber", "email", "OB", "level", "endo date"])
        workbook.active.append(["5156041234567897", "first@example.com", 1000, 1, datetime(2026, 10, 30, 13, 14, 15)])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        upload = BytesIO(output.getvalue())
        upload.name = "accounts.xlsx"
        root = Path(__file__).resolve().parents[1]
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: upload if kw.get("key") == "source_upload" else None), \
                patch("exports.build_excel_zip", wraps=build_excel_zip) as zip_builder, \
                patch("exports.build_export_excel", wraps=build_export_excel) as excel_builder, \
                patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(root / "app.py")).run(timeout=20)
            next(button for button in app.button if button.label == "Preview and prepare files").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            zip_builder.assert_not_called()
            excel_builder.assert_not_called()
            zip_call = next(call for call in download.call_args_list if call.kwargs.get("key") == "excel_zip")
            self.assertEqual(zip_call.kwargs["type"], "primary")
            self.assertTrue(callable(zip_call.args[1]))
            self.assertTrue(zip_call.args[2].endswith("_Excel.zip"))
            with ZipFile(BytesIO(zip_call.args[1]())) as archive:
                self.assertEqual(len(archive.namelist()), 2)
                self.assertTrue(all(name.endswith(".xlsx") for name in archive.namelist()))
            single = next(call for call in download.call_args_list if call.kwargs.get("key", "").startswith("excel_download_"))
            workbook = load_workbook(BytesIO(single.args[1]()))
            headers = [cell.value for cell in workbook.active[1]]
            cell = workbook.active.cell(2, headers.index("cardnumber") + 1)
            self.assertEqual(cell.value, "5156041234567897")
            self.assertEqual(cell.data_type, "s")
            self.assertEqual(cell.number_format, "@")
            for header, value, cell_type in (("OB", 1000, "n"), ("Level", 1, "n"),
                                            ("EndoDate", datetime(2026, 10, 30, 13, 14, 15), "d")):
                cell = workbook.active.cell(2, headers.index(header) + 1)
                self.assertEqual((cell.value, cell.data_type), (value, cell_type))
                self.assertNotEqual(cell.number_format, "@")
            workbook.close()


    def test_geo_individual_excel_download_shows_xlsx_and_retains_both_source_identifiers(self):
        workbook = Workbook()
        workbook.active.title = "NEGATIVE"
        workbook.active.append(["account", "card_no", "old_ic", "email", "Phone 1", "OB", "Product", "placement", "Area", "Christmas Amnesty", "assign_date"])
        workbook.active.append(["00001234567890123456", "5156041234567897", "000019", "one@example.com", "09171234567",
                                1234.5, "MC", 1, "Mindanao", "QUALIFIED", datetime(2026, 10, 30)])
        data = BytesIO()
        workbook.save(data)
        workbook.close()
        upload = BytesIO(data.getvalue())
        upload.name = "GEO.xlsx"
        root = Path(__file__).resolve().parents[1]
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: upload if kw.get("key") == "source_upload" else None), \
                patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(root / "app.py")).run(timeout=20)
            next(button for button in app.button if button.label == "Preview and prepare files").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            picker = next(widget for widget in app.selectbox if widget.label == "File to preview")
            self.assertTrue(all(name.endswith(".xlsx") for name in picker.options))
            calls = download.call_args_list
            single_index = next(index for index, call in enumerate(calls) if call.kwargs.get("key", "").startswith("excel_download_"))
            csv_index = next(index for index, call in enumerate(calls) if call.kwargs.get("key", "").startswith("download_"))
            self.assertLess(single_index, csv_index)
            call = calls[single_index]
            self.assertEqual(call.args[0], "Download " + call.args[2])
            self.assertEqual(call.kwargs["type"], "primary")
            saved = load_workbook(BytesIO(call.args[1]()))
            headers = [cell.value for cell in saved.active[1]]
            for header, expected in (("primary_key", "00001234567890123456"), ("cardnumber", "5156041234567897"), ("indexoneuid", "000019")):
                cell = saved.active.cell(2, headers.index(header) + 1)
                self.assertEqual((cell.value, cell.data_type, cell.number_format), (expected, "s", "@"))
            self.assertEqual(saved.active.cell(2, headers.index("OB") + 1).data_type, "n")
            self.assertEqual(saved.active.cell(2, headers.index("EndoDate") + 1).data_type, "d")
            reopened_data = BytesIO()
            saved.save(reopened_data)
            saved.close()
            reopened = load_workbook(BytesIO(reopened_data.getvalue()))
            for header, expected in (("primary_key", "00001234567890123456"), ("cardnumber", "5156041234567897"), ("indexoneuid", "000019")):
                cell = reopened.active.cell(2, headers.index(header) + 1)
                self.assertEqual((cell.value, cell.data_type, cell.number_format), (expected, "s", "@"))
            reopened.close()


if __name__ == "__main__":
    unittest.main()
