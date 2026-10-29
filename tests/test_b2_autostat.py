from __future__ import annotations

import csv
import unittest
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from automation import AutomationError
from autostat import (
    SourceSelection, build_autostat_excel,
    build_autostat_excel_zip, inspect_contact_source,
)
from b2_autostat import B2ActivitySource, b2_metadata_columns, build_b2_autostat_exports
from folder_output import MANILA_TIMEZONE


ROOT = Path(__file__).resolve().parents[1]
NOW = datetime(2026, 10, 30, 16, 2, 3, tzinfo=MANILA_TIMEZONE)
EMAIL_HEADERS = ("Name", "Account No.", "STATUS CODE", "REMARKS", "REMARKS BY", "REMARKS DATE", "Email", "Collector")
PHONE_HEADERS = ("Account number", "Name", "STATUS CODE", "REMARKS", "REMARKS BY", "REMARKS DATE", "Mobile Number")


def source(data, name="NEGATIVE B2 - CARDS.csv"):
    sheet = inspect_contact_source(data, Path(name).suffix)[0]
    return B2ActivitySource(SourceSelection(name, data, sheet), *b2_metadata_columns(sheet.headers))


def rows(file):
    return list(csv.DictReader(StringIO(file.data.decode("utf-8-sig" if file.data.startswith(b"\xef\xbb\xbf") else "cp1252"))))


class Upload(BytesIO):
    def __init__(self, data, name="NEGATIVE B2 - CARDS.csv"):
        super().__init__(data)
        self.name = name


class B2AutostatTests(unittest.TestCase):
    def test_exact_layouts_and_channel_text_use_source_data_and_current_settings(self):
        selection = source(
            b"primary_key,cardnumber,FullName,email,phone,Collector\n"
            b"0005156041234567897.0,DIFFERENT-CARD,Sample Person,User@EXAMPLE.COM,+63 (917) 123-4567,COLLECTOR1\n"
        )
        bundle = build_b2_autostat_exports([selection], timestamp=NOW)
        self.assertEqual([file.filename for file in bundle.files],
                         ["AUTOSTAT_B2_Email.csv", "AUTOSTAT_B2_SMS.csv", "AUTOSTAT_B2_Viber.csv"])
        self.assertEqual([file.headers for file in bundle.files], [EMAIL_HEADERS, PHONE_HEADERS, PHONE_HEADERS])
        email, sms, viber = [rows(file)[0] for file in bundle.files]
        self.assertEqual(email, {
            "Name": "Sample Person", "Account No.": "0005156041234567897",
            "STATUS CODE": "EMAIL BLAST SENT - WAITING FOR REPLY",
            "REMARKS": "User@example.com | EMAIL SENT | COLLECTOR1", "REMARKS BY": "PBBALUYOT",
            "REMARKS DATE": "2026-10-30 16:02:03", "Email": "User@example.com", "Collector": "COLLECTOR1",
        })
        for row, channel, remark in (
            (sms, "SMS", "09171234567 |  SMS SENT"),
            (viber, "VIBER", "09171234567 | VIBER SENT"),
        ):
            self.assertEqual(row, {
                "Account number": "0005156041234567897", "Name": "Sample Person",
                "STATUS CODE": f"{channel} BLAST SENT - WAITING FOR REPLY", "REMARKS": remark,
                "REMARKS BY": "PBBALUYOT", "REMARKS DATE": "2026-10-30 16:02:03", "Mobile Number": "09171234567",
            })
        with ZipFile(BytesIO(bundle.zip_data)) as archive:
            self.assertEqual(archive.namelist(), [file.filename for file in bundle.files])
            for file in bundle.files:
                self.assertEqual(archive.read(file.filename), file.data)

    def test_primary_key_controls_accounts_deduplication_and_missing_accounts_in_all_outputs(self):
        selection = source(
            b"cardnumber, PRIMARY_KEY ,email,phone\n"
            b"SAME-CARD,000001,one@example.com,09171234567\n"
            b"SAME-CARD,000002,two@example.com,09181234567\n"
            b"OTHER-CARD,000001,duplicate@example.com,09191234567\n"
            b"DO-NOT-SUBSTITUTE,,missing@example.com,09201234567\n"
            b"5.15604E+15,000003,three@example.com,09211234567\n"
        )
        self.assertEqual(selection.source.sheet.account_column, 0)  # Generic detection prefers cardnumber.
        bundle = build_b2_autostat_exports([selection], timestamp=None)
        self.assertEqual([file.count for file in bundle.files], [3, 3, 3])
        self.assertEqual(bundle.summary[0]["Missing account"], 1)
        for file in bundle.files:
            header = "Account No." if file.filename.endswith("Email.csv") else "Account number"
            self.assertEqual([row[header] for row in rows(file)], ["000001", "000002", "000003"])
        workbook = load_workbook(BytesIO(build_autostat_excel(bundle)))
        for sheet in workbook:
            column = "B" if sheet.title == "B2_Email" else "A"
            self.assertEqual([cell.value for cell in sheet[column][1:]], ["000001", "000002", "000003"])
            self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in sheet[column][1:]))
        workbook.close()

    def test_ambiguous_primary_key_headers_stop_instead_of_using_cardnumber(self):
        selection = source(b"cardnumber,primary_key,Primary Key,email\nCARD,ACCOUNT1,ACCOUNT2,one@example.com\n")
        with self.assertRaisesRegex(AutomationError, "more than one primary_key"):
            build_b2_autostat_exports([selection], timestamp=None)

    def test_validation_deduplication_source_identity_metadata_and_ob_order(self):
        first = source(
            b"cardnumber,indexoneuid,Name,email,phone,Collector,OB\n"
            b"0001,ID1,Invalid row,bad,123,WRONG,9000\n"
            b"0001,ID1,First valid,one@example.com,09171234567,FIRST,100\n"
            b"0001,ID1,Duplicate,two@example.com,09181234567,WRONG,8000\n"
            b"0001,ID2,Second identity,identity@example.com,09191234567,SECOND,500\n"
            b"0002,ID3,Phone only,,09201234567,THIRD,300\n"
            b",ID4,Orphan,orphan@example.com,09211234567,ORPHAN,10000\n"
        )
        second = source(
            b"cardnumber,indexoneuid,FullName,email,phone,Collector,OB\n"
            b"0001,ID1,Across files,other@example.com,09221234567,WRONG,10000\n"
            b"0003,ID5,Email only,extra@example.com,,FOURTH,600\n",
            "NEGATIVE B2 - BEL.csv",
        )
        bundle = build_b2_autostat_exports([first, second], timestamp=None, remark_by=" OTHER ")
        self.assertEqual([file.count for file in bundle.files], [3, 3, 3])
        self.assertEqual([row["Name"] for row in rows(bundle.files[0])], ["Email only", "Second identity", "First valid"])
        self.assertEqual([row["Collector"] for row in rows(bundle.files[0])], ["FOURTH", "SECOND", "FIRST"])
        self.assertEqual([row["Name"] for row in rows(bundle.files[1])], ["Second identity", "Phone only", "First valid"])
        for file in bundle.files:
            self.assertTrue(all(row["REMARKS DATE"] == "" and row["REMARKS BY"] == "OTHER" for row in rows(file)))
        self.assertEqual(bundle.summary[0]["Invalid emails"], 1)
        self.assertEqual(bundle.summary[0]["Invalid phones"], 1)
        self.assertEqual(bundle.summary[0]["Missing account"], 1)
        self.assertEqual(bundle.summary[0]["Duplicate Email rows removed"], 1)

    def test_optional_metadata_and_empty_channels_keep_template_headers(self):
        selection = source(b"Account number,Mobile Number\n000001,09171234567\n")
        bundle = build_b2_autostat_exports([selection], timestamp=None)
        self.assertEqual([file.count for file in bundle.files], [0, 1, 1])
        self.assertEqual(bundle.files[0].headers, EMAIL_HEADERS)
        self.assertEqual(rows(bundle.files[1])[0]["Name"], "")
        email_only = source(b"Account No.,Email\n000002,only@example.com\n")
        for collector in ("", "DEFAULT"):
            result = build_b2_autostat_exports([email_only], timestamp=None, default_collector=collector)
            row = rows(result.files[0])[0]
            self.assertEqual(row["Collector"], collector)
            self.assertEqual(row["REMARKS"], f"only@example.com | EMAIL SENT | {collector}")
            self.assertEqual([file.count for file in result.files], [1, 0, 0])

    def test_scientific_accounts_and_uncomputed_formulas_fail_without_output(self):
        with self.assertRaisesRegex(AutomationError, "scientific notation"):
            build_b2_autostat_exports([source(b"cardnumber,email\n5.15604E+15,a@example.com\n")], timestamp=None)
        workbook = Workbook()
        workbook.active.append(["Account No.", "Name", "Email"])
        workbook.active.append(["00001", '=CONCAT("Sample","Name")', "a@example.com"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        with self.assertRaisesRegex(AutomationError, "formula has no saved value"):
            build_b2_autostat_exports([source(output.getvalue(), "B2.xlsx")], timestamp=None)
        with self.assertRaisesRegex(AutomationError, "Remarks By"):
            build_b2_autostat_exports([source(b"Account,email\n1,a@example.com\n")], timestamp=None, remark_by=" ")

    def test_excel_and_zip_preserve_full_identifiers_mobile_numbers_and_literal_names(self):
        selection = source(b"cardnumber,FullName,email,phone\n0005156041234567897,=literal,one@example.com,09171234567\n")
        bundle = build_b2_autostat_exports([selection], timestamp=NOW)
        workbook = load_workbook(BytesIO(build_autostat_excel(bundle)))
        self.assertEqual(workbook.sheetnames, ["B2_Email", "B2_SMS", "B2_Viber"])
        for sheet in workbook:
            account_column = "B" if sheet.title == "B2_Email" else "A"
            cell = sheet[f"{account_column}2"]
            self.assertEqual((cell.value, cell.data_type, cell.number_format), ("0005156041234567897", "s", "@"))
            name_cell = sheet["A2" if sheet.title == "B2_Email" else "B2"]
            self.assertEqual((name_cell.value, name_cell.data_type), ("=literal", "s"))
        workbook.close()
        with ZipFile(BytesIO(build_autostat_excel_zip(bundle))) as archive:
            for file in bundle.files:
                name = Path(file.filename).with_suffix(".xlsx").name
                book = load_workbook(BytesIO(archive.read(name)))
                self.assertEqual(tuple(cell.value for cell in book.active[1]), file.headers)
                book.close()


class B2AutostatPageTests(unittest.TestCase):
    def test_b2_account_control_and_prepared_files_use_primary_key(self):
        upload = Upload(b"cardnumber,primary_key,email,phone\nCARD-1,000001,one@example.com,09171234567\n")
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: [upload] if kw.get("key") == "autostat_b2_uploads" else []):
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            app.radio(key="autostat_upload_type").set_value("B2").run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            account = next(widget for widget in app.selectbox if widget.label == "Account number column")
            self.assertEqual(account.value, 1)
            self.assertTrue(account.disabled)
            app.button(key="autostat_b2_prepare").click().run(timeout=20)
            self.assertEqual(len(app.error), 0)
            for file in app.session_state["autostat_b2_result"]["bundle"].files:
                header = "Account No." if file.filename.endswith("Email.csv") else "Account number"
                self.assertEqual(rows(file)[0][header], "000001")

    def test_b2_upload_prepare_download_save_invalidation_and_switching(self):
        uploads = [Upload(b"cardnumber,FullName,email,phone,Collector\n00001,Sample,one@example.com,09171234567,COLLECTOR\n")]
        def uploader(label, **kwargs):
            return uploads if kwargs.get("key") == "autostat_b2_uploads" else []

        with patch("streamlit.file_uploader", side_effect=uploader), TemporaryDirectory(prefix="b2-autostat-test-") as directory:
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            app.radio(key="autostat_upload_type").set_value("B2").run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.text_input(key="autostat_b2_remark_by").value, "PBBALUYOT")
            self.assertFalse(app.button(key="autostat_b2_prepare").disabled)
            app.button(key="autostat_b2_prepare").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state["autostat_b2_result"]
            self.assertEqual([file.count for file in result["bundle"].files], [1, 1, 1])
            self.assertEqual(len(app.get("download_button")), 5)
            self.assertEqual([(metric.label, metric.value) for metric in app.metric],
                             [("Email accounts", "1"), ("SMS accounts", "1"), ("Viber accounts", "1")])
            self.assertEqual(result["timestamp"].tzinfo, MANILA_TIMEZONE)
            app.run(timeout=20)
            self.assertEqual(app.session_state["autostat_b2_result"]["timestamp"], result["timestamp"])
            app.text_input(key="autostat_b2_output_folder").set_value(directory).run(timeout=20)
            app.button(key="autostat_b2_save").click().run(timeout=20)
            self.assertEqual(len(app.error), 0)
            saved = Path(app.session_state["autostat_b2_result"]["saved_folder"])
            self.assertEqual(len(list(saved.iterdir())), 3)
            for file in result["bundle"].files:
                book = load_workbook(saved / Path(file.filename).with_suffix(".xlsx").name)
                self.assertEqual(tuple(cell.value for cell in book.active[1]), file.headers)
                book.close()
            app.checkbox(key="autostat_b2_blank_time").set_value(True).run(timeout=20)
            self.assertNotIn("autostat_b2_result", app.session_state)
            app.selectbox(key="autostat_b2_output_format").select("CSV (.csv)").run(timeout=20)
            app.button(key="autostat_b2_prepare").click().run(timeout=20)
            self.assertTrue(all(rows(file)[0]["REMARKS DATE"] == "" for file in app.session_state["autostat_b2_result"]["bundle"].files))
            app.text_input(key="autostat_b2_collector").set_value("FALLBACK").run(timeout=20)
            self.assertNotIn("autostat_b2_result", app.session_state)
            app.button(key="autostat_b2_prepare").click().run(timeout=20)
            uploads[0] = Upload(b"cardnumber,email\n00002,updated@example.com\n")
            app.run(timeout=20)
            self.assertNotIn("autostat_b2_result", app.session_state)
            app.button(key="autostat_b2_prepare").click().run(timeout=20)
            self.assertEqual(rows(app.session_state["autostat_b2_result"]["bundle"].files[0])[0]["Account No."], "00002")
            uploads.clear()
            app.run(timeout=20)
            self.assertNotIn("autostat_b2_result", app.session_state)
            self.assertEqual(len(app.get("download_button")), 0)
            app.radio(key="autostat_upload_type").set_value("Standard").run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.text_input(key="autostat_remark_by").value, "MBURGOS")


if __name__ == "__main__":
    unittest.main()
