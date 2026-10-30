from __future__ import annotations

import csv
import unittest
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from streamlit.testing.v1 import AppTest

from automation import (
    DEFAULT_TEMPLATE_HEADERS,
    FIELD_MAPPING,
    TARGET_HEADERS,
    AutomationError,
    create_default_csv_template,
    create_default_template,
    csv_template_headers,
    default_source_sheet,
    fill_csv_template,
    fill_excel_template,
    inspect_workbook,
    read_source,
)
from export_config import CHANNEL_SUFFIXES, GEO_OUTPUT_FILENAMES, OUTPUT_FILENAMES
from email_validation import validated_email
from exports import (
    GEO_REVIEW_GROUP_KEYS, build_geo_exports, build_negative_exports, build_sheet_exports,
    build_export_csv,
    channel_filename, export_filename, separate_channels, separate_geo, separate_negative,
)
from folder_output import save_files_to_folder


APP_PATH = Path(__file__).resolve().parents[1] / "app.py"


def save(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def sample_row(**updates):
    row = {
        "debtor_email": "ana@example.com",
        "OLD I.C": "000123",
        "Account no.": "0000000456",
        "financing/Card no.": "000123456789012345",
        "Name": "Ana Cruz",
        "endo date": datetime(2026, 10, 26),
        "level": 2,
        "product": "Card",
        "OB (Masterlist)": 1250.75,
        "phone_1": "09171234567",
        "Christmas Amnesty": None,
    }
    row.update(updates)
    return row


def source_file(rows=None, *, headers=None, title_rows=0, other_sheet=True):
    headers = list(FIELD_MAPPING) if headers is None else headers
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary" if other_sheet else "NEGATIVE"
    if other_sheet:
        sheet.append(["Not debtor data"])
        sheet = workbook.create_sheet("NEGATIVE")
    for _ in range(title_rows):
        sheet.append(["Masterlist export"])
    sheet.append(headers)
    for row in [sample_row()] if rows is None else rows:
        sheet.append([row.get(header) for header in headers])
    return save(workbook)


def loaded(data):
    return load_workbook(BytesIO(data))


def geo_source_file(rows):
    aliases = {
        "name": "Name", "old_ic": "OLD I.C", "account": "Account no.",
        "card_no": "financing/Card no.", "assign_date": "endo date", "placement": "level",
        "Product": "product", "OB (MASTERLIST}": "OB (Masterlist)",
        "Phone 1": "phone_1", "email": "debtor_email", "Area": "Area",
        "Christmas Amnesty": "Christmas Amnesty",
    }
    return source_file(
        [{header: row.get(source) for header, source in aliases.items()} for row in rows],
        headers=list(aliases),
    )


class EmailValidationTests(unittest.TestCase):
    def test_trim_spaces_without_changing_valid_dots_or_plus_tags(self):
        self.assertEqual(
            validated_email(" \tJohn.Smith+campaign@EXAMPLE.COM\n"),
            "John.Smith+campaign@example.com",
        )

    def test_malformed_and_non_text_addresses_are_rejected(self):
        for value in (
            "john smith@example.com", "john\tsmith@example.com",
            "john..smith@example.com", ".john@example.com", "john.@example.com",
            "john@example..com", "john@.example.com", "john@example.com.",
            "johnexample.com", "john@@example.com", "john@", "@example.com",
            "john@example", "john@example.invalid", "john@example.test",
            "john@example.com;other@example.com", "a" * 65 + "@example.com",
            1234, None, "", " \t",
        ):
            with self.subTest(value=value):
                self.assertIsNone(validated_email(value))

    def test_source_counts_invalid_emails_and_keeps_phone_accounts(self):
        source = source_file([
            sample_row(debtor_email="ana..cruz@example.com"),
            sample_row(debtor_email="ana cruz@example.com", phone_1=None),
            sample_row(debtor_email=" Ana.Cruz+tag@EXAMPLE.COM ", phone_1=None),
            sample_row(debtor_email=None),
            {},
            {"debtor_email": "not-an-email"},
        ])
        result = read_source(source, "NEGATIVE", 1)
        self.assertEqual(result.invalid_email_rows, 3)
        self.assertEqual(result.valid_email_rows, 1)
        self.assertEqual(result.missing_email_rows, 1)
        self.assertEqual(result.no_contact_rows_skipped, 2)
        self.assertEqual(result.blank_rows_skipped, 1)
        self.assertEqual(result.source_rows, [2, 4, 5])
        self.assertEqual([row["email"] for row in result.records], [None, "Ana.Cruz+tag@example.com", None])
        # Reading and cleaning the source must not modify the original workbook.
        workbook = loaded(source)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["NEGATIVE"].cell(2, 1).value, "ana..cruz@example.com")

    def test_invalid_addresses_cannot_enter_grouped_or_single_exports(self):
        base = read_source(source_file(), "NEGATIVE", 1).records[0]
        cases = [
            ("invalid-with-phone", "bad..address@example.com", "09171234567", "QUALIFIED", "CARDS", 1),
            ("invalid-only", "bad address@example.com", None, "QUALIFIED", "CARDS", 1),
            ("trimmed", " John.Smith+tag@EXAMPLE.COM ", None, None, "BEL", 5),
            ("valid", "ana.cruz@example.com", "09181234567", "QUALIFIED", "OTHER", 1),
        ]
        records = [
            {**base, "primary_key": key, "email": email, "phone": phone,
             "ChristmasAmnesty": status, "product": product, "level": level}
            for key, email, phone, status, product, level in cases
        ]
        original = [dict(row) for row in records]
        for grouped in (False, True):
            for extension in (".xlsx", ".csv"):
                with self.subTest(grouped=grouped, extension=extension):
                    template = (create_default_template() if extension == ".xlsx"
                                else (",".join(TARGET_HEADERS) + "\n").encode())
                    if grouped:
                        bundle = build_negative_exports(template, records, "Template", 1, extension)
                    else:
                        bundle = build_sheet_exports(
                            template, records, "Template", 1, base_name="Campaign", label="Test",
                            template_extension=extension,
                        )
                    self.assertEqual(bundle.no_contact_rows_skipped, 1)
                    self.assertEqual(sum(bundle.counts.values()), 3)
                    actual = {"email": {}, "viber": {}}
                    with ZipFile(BytesIO(bundle.zip_data)) as archive:
                        for file in bundle.files:
                            self.assertEqual(archive.read(file.filename), file.data)
                            self.assertTrue(file.filename.endswith(".csv"))
                            self.assertFalse(file.data.startswith(b"\xef\xbb\xbf"))
                            rows = list(csv.reader(StringIO(file.data.decode("utf-8-sig"))))
                            for row in rows[1:]:
                                values = dict(zip(rows[0], row))
                                if values.get("primary_key"):
                                    actual[file.channel][values["primary_key"]] = values["email"]
                    self.assertEqual(actual["email"], {
                        "trimmed": "John.Smith+tag@example.com", "valid": "ana.cruz@example.com",
                    })
                    self.assertEqual(set(actual["viber"]), {"invalid-with-phone", "valid"})
                    self.assertTrue(all(not value for value in actual["viber"].values()))
        self.assertEqual(records, original)


class MappingTests(unittest.TestCase):
    def test_template_headers_are_also_accepted_as_source_headers(self):
        row = {target: sample_row()[source] for source, target in FIELD_MAPPING.items()}
        data = source_file([row], headers=list(TARGET_HEADERS), other_sheet=False)
        info = inspect_workbook(data)[0]
        self.assertEqual(info.matched_fields, len(FIELD_MAPPING))
        self.assertEqual(read_source(data, 'NEGATIVE', 1).records,
                         read_source(source_file(), 'NEGATIVE', 1).records)

    def test_partial_source_copies_present_values_and_leaves_missing_fields_blank(self):
        headers = ['debtor_email', 'OLD I.C', 'Account no.', 'financing/Card no.', 'Name']
        data = source_file([sample_row(), sample_row(debtor_email='bad..email@example.com')], headers=headers)
        with self.assertRaisesRegex(AutomationError, 'Source headers are missing'):
            read_source(data, 'NEGATIVE', 1)
        result = read_source(data, 'NEGATIVE', 1, allow_missing=True)
        self.assertEqual(len(result.records), 1)
        self.assertEqual(result.invalid_email_rows, 1)
        self.assertEqual(result.no_contact_rows_skipped, 1)
        record = result.records[0]
        self.assertEqual(record['primary_key'], sample_row()['Account no.'])
        self.assertEqual(record['cardnumber'], sample_row()['financing/Card no.'])
        self.assertEqual(record['FullName'], 'Ana Cruz')
        for field in ('endodate', 'level', 'product', 'OB', 'phone', 'ChristmasAmnesty'):
            self.assertIsNone(record[field])

    def test_partial_source_requires_account_and_contact_and_accepts_phone_aliases(self):
        for phone_header in ('PHONE', 'phone number', 'Mobile Number', 'contact number'):
            with self.subTest(phone=phone_header):
                data = source_file([{'primary_key': '0001', phone_header: '09171234567'}], headers=['primary_key', phone_header])
                row = read_source(data, 'NEGATIVE', 1, allow_missing=True).records[0]
                self.assertEqual(row['primary_key'], '0001')
                self.assertEqual(row['phone'], '09171234567')
                self.assertIsNone(row['email'])
        for headers, values, error in (
            (['email'], {'email': 'first@example.com'}, 'No account-number column'),
            (['cardnumber'], {'cardnumber': '0001'}, 'No email or phone column'),
        ):
            with self.subTest(headers=headers):
                with self.assertRaisesRegex(AutomationError, error):
                    read_source(source_file([values], headers=headers), 'NEGATIVE', 1, allow_missing=True)

    def test_best_data_sheet_is_selected_when_negative_is_absent(self):
        workbook = loaded(source_file())
        workbook['NEGATIVE'].title = 'Accounts'
        sheets = inspect_workbook(save(workbook))
        self.assertEqual(default_source_sheet(sheets), 'Accounts')

    def test_negative_default_and_detect_header_below_title(self):
        sheets = inspect_workbook(source_file(title_rows=3))
        self.assertEqual(default_source_sheet(sheets), "NEGATIVE")
        self.assertEqual(sheets[1].header_row, 4)
        self.assertEqual(sheets[1].matched_fields, len(FIELD_MAPPING))

    def test_all_mappings_reordered_source_and_template(self):
        source = source_file(headers=list(reversed(FIELD_MAPPING)))
        records = read_source(source, "NEGATIVE", 1).records
        template = Workbook()
        template.active.title = "Destination"
        headers = ["Extra"] + list(reversed(TARGET_HEADERS))
        template.active.append(headers)
        result = loaded(fill_excel_template(save(template), records, "Destination", 1))
        self.addCleanup(result.close)
        sheet = result.active
        self.assertEqual([cell.value for cell in sheet[1]], headers)
        actual = dict(zip(headers, [cell.value for cell in sheet[2]]))
        expected = {target: sample_row()[source] for source, target in FIELD_MAPPING.items()}
        expected["Extra"] = None
        self.assertEqual(actual, expected)
        for column in ("primary_key", "cardnumber", "phone", "indexoneuid"):
            self.assertEqual(sheet.cell(2, headers.index(column) + 1).data_type, "s")
            self.assertEqual(sheet.cell(2, headers.index(column) + 1).number_format, "@")

    def test_header_variations_and_ob_alias(self):
        workbook = loaded(source_file())
        sheet = workbook["NEGATIVE"]
        for cell in sheet[1]:
            cell.value = "  " + str(cell.value).upper().replace(" ", "\n") + "  "
        sheet.cell(1, 9, "OB")
        data = read_source(save(workbook), "NEGATIVE", 1)
        self.assertEqual(data.records[0]["OB"], 1250.75)
        self.assertEqual(data.records[0]["email"], "ana@example.com")

    def test_sample_product_type_header_and_sixteen_column_template(self):
        workbook = loaded(source_file())
        workbook["NEGATIVE"].cell(1, 8, "Product Type")
        source = save(workbook)
        self.assertEqual(inspect_workbook(source)[1].matched_fields, len(FIELD_MAPPING))
        records = read_source(source, "NEGATIVE", 1).records
        template_headers = [
            "email", "primary_key", "indexoneuid", "cardnumber", "FullName", "phone", "OB",
            "Monthly", "Terms", "doclass", "Level", "Product", "EndoDate", "AREA",
            "ChristmasAmnesty", "NewEndo",
        ]
        workbook = Workbook()
        workbook.active.append(template_headers)
        result = loaded(fill_excel_template(save(workbook), records, "Sheet", 1))
        self.addCleanup(result.close)
        self.assertEqual([cell.value for cell in result.active[1]], template_headers)
        output = dict(zip(template_headers, [cell.value for cell in result.active[2]]))
        self.assertEqual(output["Product"], "Card")
        self.assertEqual(output["Level"], 2)
        self.assertEqual(output["EndoDate"], datetime(2026, 10, 26))
        for header in ("Monthly", "Terms", "doclass", "AREA", "ChristmasAmnesty", "NewEndo"):
            self.assertIsNone(output[header])

    def test_numeric_identifiers_keep_zero_display_masks(self):
        workbook = loaded(source_file())
        sheet = workbook["NEGATIVE"]
        sheet.cell(2, 3, 456).number_format = "0000000000"
        sheet.cell(2, 10, 9171234567).number_format = "00000000000"
        row = read_source(save(workbook), "NEGATIVE", 1).records[0]
        self.assertEqual(row["primary_key"], "0000000456")
        self.assertEqual(row["phone"], "09171234567")

    def test_blank_rows_skipped_duplicates_and_missing_email_kept(self):
        source = source_file([sample_row(), {}, sample_row(), sample_row(debtor_email=None)])
        result = read_source(source, "NEGATIVE", 1)
        self.assertEqual(len(result.records), 3)
        self.assertEqual(result.source_rows, [2, 4, 5])
        self.assertEqual(result.blank_rows_skipped, 1)
        self.assertIsNone(result.records[-1]["email"])

    def test_remove_only_accounts_missing_both_email_and_phone(self):
        contacts = [
            (None, None), (None, "  "), ("a@example.com", None),
            (None, "09171234567"), ("a@example.com", "09171234567"),
            ("  ", "\t"), ("", "09171234567"), ("a@example.com", "  "),
        ]
        source = source_file([sample_row(debtor_email=email, phone_1=phone) for email, phone in contacts])
        result = read_source(source, "NEGATIVE", 1)
        self.assertEqual(result.no_contact_rows_skipped, 3)
        self.assertEqual(result.source_rows, [4, 5, 6, 8, 9])
        self.assertEqual(len(result.records), 5)
        self.assertEqual(result.blank_rows_skipped, 0)

    def test_empty_rows_and_contactless_accounts_counted_separately(self):
        result = read_source(source_file([{}, sample_row(debtor_email=None, phone_1=None)]), "NEGATIVE", 1)
        self.assertEqual(result.records, [])
        self.assertEqual(result.blank_rows_skipped, 1)
        self.assertEqual(result.no_contact_rows_skipped, 1)

    def test_hidden_rows_are_included(self):
        workbook = loaded(source_file())
        workbook["NEGATIVE"].row_dimensions[2].hidden = True
        self.assertEqual(len(read_source(save(workbook), "NEGATIVE", 1).records), 1)

    def test_missing_and_duplicate_headers_fail_with_names(self):
        with self.assertRaisesRegex(AutomationError, "phone_1"):
            read_source(source_file(headers=[header for header in FIELD_MAPPING if header != "phone_1"]), "NEGATIVE", 1)
        workbook = loaded(source_file())
        workbook["NEGATIVE"].cell(1, len(FIELD_MAPPING) + 1, "PHONE 1")
        with self.assertRaisesRegex(AutomationError, "ambiguous.*phone_1"):
            read_source(save(workbook), "NEGATIVE", 1)

    def test_christmas_amnesty_qualified_or_blank(self):
        statuses = ["QUALIFIED", " qualified ", None, "", "NOT QUALIFIED"]
        source = source_file([sample_row(**{"Christmas Amnesty": status}) for status in statuses])
        records = read_source(source, "NEGATIVE", 1).records
        self.assertEqual([row["ChristmasAmnesty"] for row in records], ["QUALIFIED", "QUALIFIED", None, None, None])

    def test_missing_amnesty_header_cannot_silently_become_blank(self):
        source = source_file(headers=[header for header in FIELD_MAPPING if header != "Christmas Amnesty"])
        with self.assertRaisesRegex(AutomationError, "Christmas Amnesty"):
            read_source(source, "NEGATIVE", 1)

    def test_missing_formula_cache_and_excel_errors_fail(self):
        for value, message in [("=1+1", "formula has no saved value"), ("#N/A", "Excel error")]:
            with self.subTest(value=value):
                with self.assertRaisesRegex(AutomationError, message):
                    read_source(source_file([sample_row(**{"OB (Masterlist)": value})]), "NEGATIVE", 1)

    def test_selected_sheet_only(self):
        workbook = loaded(source_file())
        sheet = workbook.create_sheet("POSITIVE")
        sheet.append(list(FIELD_MAPPING))
        sheet.append(list(sample_row(Name="Ben Santos").values()))
        source = save(workbook)
        self.assertEqual(read_source(source, "NEGATIVE", 1).records[0]["FullName"], "Ana Cruz")
        self.assertEqual(read_source(source, "POSITIVE", 1).records[0]["FullName"], "Ben Santos")

    def test_bad_workbook_is_actionable(self):
        with self.assertRaisesRegex(AutomationError, "could not be read"):
            inspect_workbook(b"not an Excel workbook")


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.records = read_source(source_file(), "NEGATIVE", 1).records

    def test_default_templates_include_all_sixteen_campaign_headers(self):
        expected = [
            "email", "primary_key", "indexoneuid", "cardnumber", "FullName", "phone", "OB",
            "Monthly", "Terms", "doclass", "Level", "Product", "EndoDate", "AREA",
            "ChristmasAmnesty", "NewEndo",
        ]
        self.assertEqual(csv_template_headers(create_default_csv_template()), expected)
        workbook = loaded(create_default_template())
        self.addCleanup(workbook.close)
        self.assertEqual([cell.value for cell in workbook.active[1]], expected)

    def test_template_styles_other_sheets_and_stale_rows(self):
        workbook = loaded(create_default_template())
        sheet = workbook.active
        sheet.insert_rows(1)
        sheet.cell(1, 1, "Template title")
        sheet.cell(3, 5, "Old debtor").fill = PatternFill("solid", fgColor="FFE699")
        sheet.cell(4, 5, "Stale debtor")
        sheet.cell(4, 1).hyperlink = "https://example.test/old"
        sheet.row_dimensions[3].height = 25
        sheet.row_dimensions[3].hidden = True
        sheet.auto_filter.ref = "A2:J4"
        workbook.create_sheet("Instructions").cell(1, 1, "Keep this")
        result = loaded(fill_excel_template(save(workbook), self.records, "Template", 2))
        self.addCleanup(result.close)
        sheet = result["Template"]
        self.assertEqual(sheet.cell(1, 1).value, "Template title")
        self.assertEqual(sheet.cell(3, 5).value, "Ana Cruz")
        self.assertEqual(sheet.cell(3, 5).fill.fgColor.rgb, "00FFE699")
        self.assertEqual(sheet.row_dimensions[3].height, 25)
        self.assertFalse(sheet.row_dimensions[3].hidden)
        self.assertIsNone(sheet.cell(4, 5).value)
        self.assertIsNone(sheet.cell(4, 1).hyperlink)
        self.assertEqual(sheet.auto_filter.ref, "A2:P3")
        self.assertEqual(result["Instructions"].cell(1, 1).value, "Keep this")

    def test_table_range_grows_to_fit_output(self):
        workbook = loaded(create_default_template())
        sheet = workbook.active
        sheet.append(["old"] * 10)
        table = Table(displayName="Debtors", ref="A1:J2")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
        result = loaded(fill_excel_template(save(workbook), self.records * 3, "Template", 1))
        self.addCleanup(result.close)
        self.assertEqual(result.active.tables["Debtors"].ref, "A1:J4")

    def test_untrusted_string_remains_text_in_excel(self):
        self.records[0]["FullName"] = "=HYPERLINK(\"https://example.test\")"
        result = loaded(fill_excel_template(create_default_template(), self.records, "Template", 1))
        self.addCleanup(result.close)
        self.assertEqual(result.active.cell(2, 5).data_type, "s")
        self.assertEqual(result.active.cell(2, 5).value, self.records[0]["FullName"])

    def test_csv_keeps_order_quoting_unicode_and_leading_zeros(self):
        headers = ["Extra"] + list(reversed(TARGET_HEADERS))
        template = (";".join(headers) + "\r\n").encode("utf-8-sig")
        self.records[0]["FullName"] = 'Peña, Ana "A"'
        output = fill_csv_template(template, self.records)
        self.assertFalse(output.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(StringIO(output.decode("cp1252"))))
        self.assertEqual(rows[0], headers)
        row = dict(zip(headers, rows[1]))
        self.assertEqual(row["primary_key"], "0000000456")
        self.assertEqual(row["FullName"], 'Peña, Ana "A"')
        self.assertEqual(row["Extra"], "")
        self.assertEqual(row["endodate"], "2026-10-26 00:00:00")

    def test_every_template_delimiter_produces_comma_delimited_csv(self):
        for delimiter in (",", ";", "\t", "|"):
            with self.subTest(delimiter=delimiter):
                template = StringIO(newline="")
                csv.writer(template, delimiter=delimiter).writerow(DEFAULT_TEMPLATE_HEADERS)
                records = [{**self.records[0], "FullName": 'Pe\u00f1a, Ana "A"\nSecond line'}]
                output = fill_csv_template(template.getvalue().encode("utf-8"), records)
                rows = list(csv.reader(StringIO(output.decode("cp1252"))))
                self.assertEqual(rows[0], list(DEFAULT_TEMPLATE_HEADERS))
                self.assertEqual(len(rows), 2)
                self.assertEqual(rows[1][DEFAULT_TEMPLATE_HEADERS.index("FullName")], records[0]["FullName"])
                self.assertEqual(rows[1][DEFAULT_TEMPLATE_HEADERS.index("primary_key")], "0000000456")
                self.assertIn(b"\r\n", output)

    def test_unsupported_characters_require_utf8_without_losing_data(self):
        records = [{**self.records[0], "FullName": "Ana \u674e"}]
        template = create_default_csv_template()
        with self.assertRaisesRegex(AutomationError, "CSV UTF-8"):
            fill_csv_template(template, records)
        output = fill_csv_template(template, records, encoding="utf-8-sig")
        self.assertTrue(output.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.DictReader(StringIO(output.decode("utf-8-sig"))))
        self.assertEqual(rows[0]["FullName"], records[0]["FullName"])

    def test_missing_template_headers_and_merged_data_fail(self):
        workbook = Workbook()
        workbook.active.append(["email"])
        with self.assertRaisesRegex(AutomationError, "Template headers are missing"):
            fill_excel_template(save(workbook), self.records, "Sheet", 1)
        workbook = loaded(create_default_template())
        workbook.active.merge_cells("A2:B2")
        with self.assertRaisesRegex(AutomationError, "merged cells"):
            fill_excel_template(save(workbook), self.records, "Template", 1)


class FakeUpload(BytesIO):
    def __init__(self, name, data):
        super().__init__(data)
        self.name = name


class SeparationTests(unittest.TestCase):
    def test_other_sheet_keeps_every_negative_group_and_template_layout(self):
        rows = [sample_row(product=product, level=level, **{"Christmas Amnesty": status})
                for status, product, level in (
                    ("QUALIFIED", "CARDS", 1), ("QUALIFIED", "CARDS", 6), ("QUALIFIED", "BEL", 3),
                    (None, "CARDS", 1), (None, "CARDS", 6), (None, "BEL", 2), (None, "PL", 4),
                    ("QUALIFIED", "PL", 2), (None, "CARDS", 2),
                )]
        records = read_source(source_file(rows), "NEGATIVE", 1).records
        for extension, template, template_sheet in (
            (".csv", create_default_csv_template(), ""), (".xlsx", create_default_template(), "Template"),
        ):
            with self.subTest(template=extension):
                negative = build_negative_exports(template, records, template_sheet, 1, extension)
                other = build_negative_exports(template, records, template_sheet, 1, extension, source_sheet="POSITIVE")
                self.assertEqual(other.counts, negative.counts)
                self.assertEqual(len(other.files), 18)
                self.assertEqual(other.zip_filename, "POSITIVE - " + negative.zip_filename)
                with ZipFile(BytesIO(other.zip_data)) as archive:
                    for actual, original in zip(other.files, negative.files):
                        self.assertEqual(actual.filename, "POSITIVE - " + original.filename)
                        self.assertEqual(actual.group_key, original.group_key)
                        self.assertEqual(actual.channel, original.channel)
                        self.assertEqual(actual.data, original.data)
                        self.assertEqual(archive.read(actual.filename), actual.data)

    def records(self, cases):
        rows = [sample_row(**{"product": product, "level": level, "Christmas Amnesty": status})
                for status, product, level in cases]
        return read_source(source_file(rows), "NEGATIVE", 1).records

    def test_status_product_and_level_boundaries_account_for_every_row(self):
        records = self.records([
            ("QUALIFIED", "CARDS", "LEVEL 1"),
            (" qualified ", " cards ", "Level 6"),
            ("QUALIFIED", "CARDS", 1),
            ("QUALIFIED", "CARDS", 6.0),
            ("QUALIFIED", "BEL", "Level 3"),
            ("QUALIFIED", "BEL", None),
            (None, "CARDS", "Level 1"),
            ("NOT QUALIFIED", "BEL", "Level 6"),
            ("QUALIFIED", "CARDS", "Level 2"),
            ("QUALIFIED", "SALAD", "Level 1"),
            (None, "CARDS", "Level 6"),
            (None, "SLA", "Level 1"),
            (None, "PL", None),
            (None, "CARDS", "Level 2"),
        ])
        groups = separate_negative(records)
        self.assertEqual({key: len(rows) for key, rows in groups.items()}, {
            "qualified_cards_level_1": 2, "qualified_cards_level_6": 2,
            "qualified_bel": 2, "qualified_review": 2,
            "blank_cards_level_1": 1, "blank_cards_level_6": 1,
            "blank_bel": 1, "blank_other_products": 2, "blank_review": 1,
        })
        self.assertEqual(sum(map(len, groups.values())), len(records))
        self.assertTrue(all(row["ChristmasAmnesty"] is None
                            for key, rows in groups.items() if key.startswith("blank_") for row in rows))
        self.assertTrue(all(row["ChristmasAmnesty"] == "QUALIFIED" for row in groups["qualified_bel"]))

    def test_zip_has_separate_csv_files_and_configured_filenames(self):
        records = self.records([
            ("QUALIFIED", "CARDS", "Level 1"),
            ("QUALIFIED", "CARDS", "Level 6"),
            ("QUALIFIED", "BEL", "Level 4"),
            (None, "PL", "Level 6"),
        ])
        with patch.dict(OUTPUT_FILENAMES, {"qualified_cards_level_1": "My custom file"}):
            bundle = build_negative_exports(create_default_template(), records, "Template", 1)
        self.assertEqual(len(bundle.files), 8)
        self.assertEqual(bundle.files[0].filename, "My custom file_Email.csv")
        self.assertEqual(bundle.files[1].filename, "My custom file_Viber.csv")
        with ZipFile(BytesIO(bundle.zip_data)) as archive:
            self.assertEqual(archive.namelist(), [file.filename for file in bundle.files])
            for file in bundle.files:
                self.assertTrue(file.filename.endswith(".csv"))
                rows = list(csv.reader(StringIO(archive.read(file.filename).decode("utf-8-sig"))))
                actual = dict(zip(rows[0], rows[1]))
                self.assertEqual(len(rows), 2)
                self.assertEqual(actual["ChristmasAmnesty"], "" if file.group_key.startswith("blank_") else "QUALIFIED")
                self.assertEqual(actual["primary_key"], "0000000456")
                self.assertEqual(actual["phone"] if file.channel == "email" else actual["email"], "")

    def test_empty_groups_omitted_and_missing_emails_kept(self):
        records = self.records([("QUALIFIED", "BEL", "Level 6")])
        records[0]["email"] = None
        bundle = build_negative_exports(create_default_template(), records, "Template", 1)
        self.assertEqual([file.key for file in bundle.files], ["qualified_bel_email", "qualified_bel_viber"])
        self.assertEqual([file.count for file in bundle.files], [0, 1])
        self.assertEqual(bundle.files[1].missing_emails, 1)
        self.assertEqual(bundle.counts["blank_cards_level_1"], 0)

    def test_direct_grouped_exports_exclude_contactless_accounts_from_both_statuses(self):
        records = self.records([
            ("QUALIFIED", "CARDS", "Level 1"),
            ("QUALIFIED", "BEL", "Level 6"),
            (None, "CARDS", "Level 6"),
            (None, "SLA", "Level 6"),
        ])
        records[0].update(email=None, phone="  ")
        records[1]["phone"] = None  # Email alone must suffice.
        records[2]["email"] = None  # Phone alone must suffice.
        records[3].update(email="\t", phone=None)
        bundle = build_negative_exports(create_default_template(), records, "Template", 1)
        self.assertEqual(bundle.no_contact_rows_skipped, 2)
        self.assertEqual(sum(bundle.counts.values()), 2)
        self.assertEqual([file.key for file in bundle.files], [
            "qualified_bel_email", "qualified_bel_viber", "blank_cards_level_6_email", "blank_cards_level_6_viber",
        ])
        self.assertEqual([file.count for file in bundle.files], [1, 0, 0, 1])
        self.assertEqual(bundle.files[3].missing_emails, 1)

    def test_grouped_csv_has_status_and_template_order(self):
        records = self.records([("QUALIFIED", "CARDS", "Level 1"), (None, "BEL", "Level 6")])
        headers = list(reversed(TARGET_HEADERS))
        template = (",".join(headers) + "\n").encode()
        bundle = build_negative_exports(template, records, "", 1, ".csv")
        for file in bundle.files:
            rows = list(csv.reader(StringIO(file.data.decode("utf-8-sig"))))
            self.assertEqual(rows[0], headers)
            self.assertEqual(rows[1][headers.index("ChristmasAmnesty")], "" if file.group_key.startswith("blank_") else "QUALIFIED")
            self.assertEqual(rows[1][headers.index("phone" if file.channel == "email" else "email")], "")
            self.assertTrue(file.filename.endswith(".csv"))

    def test_invalid_or_duplicate_filenames_fail_before_creating_zip(self):
        for name in ("", "../bad", "file/name", "CON", "bad:name"):
            with self.subTest(name=name), self.assertRaises(AutomationError):
                export_filename(name, ".xlsx")
        records = self.records([("QUALIFIED", "BEL", "Level 6")])
        with patch.dict(OUTPUT_FILENAMES, {"blank_other_products": OUTPUT_FILENAMES["qualified_bel"]}):
            with self.assertRaisesRegex(AutomationError, "unique"):
                build_negative_exports(create_default_template(), records, "Template", 1)


class ChannelTests(unittest.TestCase):
    def records(self):
        base = read_source(source_file(), "NEGATIVE", 1).records[0]
        return [
            {**base, "primary_key": "0001", "email": "email-only@example.com", "phone": None},
            {**base, "primary_key": "0002", "email": None, "phone": "09171234567"},
            {**base, "primary_key": "0003", "email": "both@example.com", "phone": "09181234567"},
            {**base, "primary_key": "0004", "email": " \t", "phone": None},
        ]

    def test_email_only_phone_only_both_and_neither_routing_without_source_mutation(self):
        records = self.records()
        original = [dict(record) for record in records]
        channels = separate_channels(records)
        self.assertEqual([row["primary_key"] for row in channels["email"]], ["0001", "0003"])
        self.assertEqual([row["primary_key"] for row in channels["viber"]], ["0002", "0003"])
        self.assertTrue(all(row["phone"] is None for row in channels["email"]))
        self.assertTrue(all(row["email"] is None for row in channels["viber"]))
        self.assertEqual(records, original)

    def test_excel_template_exports_csv_with_all_sixteen_headers_and_other_values(self):
        headers = [
            "email", "primary_key", "indexoneuid", "cardnumber", "FullName", "phone", "OB",
            "Monthly", "Terms", "doclass", "Level", "Product", "EndoDate", "AREA",
            "ChristmasAmnesty", "NewEndo",
        ]
        template = Workbook()
        template.active.append(["Template title"])
        template.active.append(headers)
        template.active.append(["Stale template data"] * len(headers))
        template.create_sheet("Instructions").cell(1, 1, "Keep these instructions")
        records = self.records()
        bundle = build_sheet_exports(
            save(template), records, "Sheet", 2, base_name="Campaign.xlsx", label="Test sheet",
        )
        self.assertEqual(bundle.counts, {"all": 3})
        self.assertEqual(bundle.no_contact_rows_skipped, 1)
        self.assertEqual([file.filename for file in bundle.files], ["Campaign_Email.csv", "Campaign_Viber.csv"])
        self.assertEqual([file.count for file in bundle.files], [2, 2])
        expected_records = {record["primary_key"]: record for record in records}
        for file in bundle.files:
            rows = list(csv.reader(StringIO(file.data.decode("utf-8-sig"))))
            self.assertEqual(rows[0], headers)
            self.assertEqual(len(rows), 3)
            for row in rows[1:]:
                values = dict(zip(headers, row))
                expected = expected_records[values["primary_key"]]
                for source_header, target in FIELD_MAPPING.items():
                    actual_header = next(header for header in headers if header.casefold() == target.casefold())
                    value = None if target == ("phone" if file.channel == "email" else "email") else expected[target]
                    self.assertEqual(values[actual_header], "" if value is None else str(value))
                for header in ("Monthly", "Terms", "doclass", "AREA", "NewEndo"):
                    self.assertEqual(values[header], "")

    def test_each_nonempty_group_keeps_both_files_even_when_one_channel_has_no_rows(self):
        records = [self.records()[0]]
        bundle = build_sheet_exports(
            create_default_template(), records, "Template", 1, base_name="Email only", label="Test",
        )
        self.assertEqual([file.count for file in bundle.files], [1, 0])
        rows = list(csv.reader(StringIO(bundle.files[1].data.decode("utf-8-sig"))))
        self.assertEqual(rows, [list(DEFAULT_TEMPLATE_HEADERS)])

    def test_excel_template_headers_are_validated_before_csv_export(self):
        for headers, message in (
            (["email"], "Template headers are missing"),
            ([*TARGET_HEADERS, "Email"], "Duplicate or ambiguous header"),
        ):
            with self.subTest(headers=headers):
                template = Workbook()
                template.active.append(headers)
                with self.assertRaisesRegex(AutomationError, message):
                    build_sheet_exports(
                        save(template), self.records(), "Sheet", 1,
                        base_name="Campaign", label="Test",
                    )

    def test_every_group_including_review_files_gets_both_channels(self):
        cases = [
            ("QUALIFIED", "CARDS", "Level 1"), ("QUALIFIED", "CARDS", "Level 6"),
            ("QUALIFIED", "BEL", "Level 3"), ("QUALIFIED", "SLA", "Level 6"),
            (None, "CARDS", "Level 1"), (None, "CARDS", "Level 6"),
            (None, "BEL", "Level 5"), (None, "PL", "Level 6"), (None, "CARDS", "Level 2"),
        ]
        rows = [sample_row(product=product, level=level, **{"Christmas Amnesty": status})
                for status, product, level in cases]
        records = read_source(source_file(rows), "NEGATIVE", 1).records
        bundle = build_negative_exports(create_default_template(), records, "Template", 1)
        self.assertEqual(len(bundle.files), 18)
        self.assertEqual(sum(bundle.counts.values()), 9)
        self.assertEqual(sum(file.count for file in bundle.files), 18)
        for key in bundle.counts:
            self.assertEqual({file.channel for file in bundle.files if file.group_key == key}, {"email", "viber"})

    def test_channel_csv_routing_and_header_only_file(self):
        headers = list(reversed(TARGET_HEADERS))
        template = (";".join(headers) + "\n").encode()
        bundle = build_sheet_exports(
            template, self.records(), "", 1, base_name="Campaign.csv", label="Test", template_extension=".csv",
        )
        for file in bundle.files:
            rows = list(csv.reader(StringIO(file.data.decode("cp1252"))))
            self.assertEqual(rows[0], headers)
            expected_keys = ["0001", "0003"] if file.channel == "email" else ["0002", "0003"]
            self.assertEqual([row[headers.index("primary_key")] for row in rows[1:]], expected_keys)
            self.assertTrue(all(row[headers.index("phone" if file.channel == "email" else "email")] == "" for row in rows[1:]))
        bundle = build_sheet_exports(
            template, [self.records()[0]], "", 1, base_name="Only email", label="Test", template_extension=".csv",
        )
        self.assertEqual(list(csv.reader(StringIO(bundle.files[1].data.decode("cp1252")))), [headers])

    def test_channel_suffix_names_can_be_edited_and_must_be_unique(self):
        self.assertEqual(channel_filename("My campaign.xlsx", "email", ".csv"), "My campaign_Email.csv")
        with patch.dict(CHANNEL_SUFFIXES, {"viber": "_Phone"}):
            self.assertEqual(channel_filename("My campaign", "viber", ".csv"), "My campaign_Phone.csv")
        with patch.dict(CHANNEL_SUFFIXES, {"viber": "_Email"}):
            with self.assertRaisesRegex(AutomationError, "unique"):
                build_sheet_exports(create_default_template(), self.records(), "Template", 1, base_name="Test", label="Test")


class GeoTests(unittest.TestCase):
    def test_other_geo_sheet_keeps_area_grouping_and_doclass(self):
        records = read_source(geo_source_file([
            sample_row(Area="Mindanao", product="CARDS", level=1, **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(Area="Mindanao", product="CARDS", level=6, **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(Area="Region 3", product="BEL"),
        ]), "NEGATIVE", 1).records
        negative = build_geo_exports(create_default_csv_template(), records, "", 1, ".csv")
        other = build_geo_exports(create_default_csv_template(), records, "", 1, ".csv", source_sheet="OVERALL")
        self.assertEqual(other.counts, negative.counts)
        self.assertEqual(len(other.files), 4)
        self.assertEqual(other.zip_filename, "OVERALL - " + negative.zip_filename)
        for actual, original in zip(other.files, negative.files):
            self.assertEqual(actual.filename, "OVERALL - " + original.filename)
            self.assertEqual(actual.data, original.data)
            for row in actual.preview:
                self.assertEqual(row["doclass"], "SBC GEO Mindanao" if row["AREA"] == "Mindanao" else "SBC GEO R3")

    def test_doclass_matches_area_in_both_channels_template_formats_and_export_modes(self):
        cases = [
            ("Mindanao", "SBC GEO Mindanao"), ("  mInDaNaO  ", "SBC GEO Mindanao"),
            ("Region 3", "SBC GEO R3"), ("REGION-3", "SBC GEO R3"), (" r.3 ", "SBC GEO R3"),
            ("Other", ""), (None, ""),
        ]
        records = read_source(geo_source_file([
            sample_row(Area=area, product="CARDS", **{"Account no.": f"000{index}"})
            for index, (area, _) in enumerate(cases)
        ]), "NEGATIVE", 1).records
        original = [dict(record) for record in records]
        expected = {f"000{index}": value for index, (_, value) in enumerate(cases)}
        headers = ["Extra", "Do Class", *[header for header in DEFAULT_TEMPLATE_HEADERS if header != "doclass"]]
        workbook = Workbook()
        workbook.active.append(headers)
        templates = [(".csv", (",".join(headers) + "\n").encode(), ""), (".xlsx", save(workbook), "Sheet")]
        for extension, template, sheet_name in templates:
            for grouped in (True, False):
                with self.subTest(template=extension, grouped=grouped):
                    if grouped:
                        bundle = build_geo_exports(template, records, sheet_name, 1, extension)
                    else:
                        bundle = build_sheet_exports(template, records, sheet_name, 1,
                                                     template_extension=extension, base_name="GEO", label="GEO")
                    with ZipFile(BytesIO(bundle.zip_data)) as archive:
                        for file in bundle.files:
                            self.assertEqual(file.headers, headers)
                            rows = list(csv.DictReader(StringIO(archive.read(file.filename).decode("cp1252"))))
                            for row in rows:
                                self.assertEqual(row["Do Class"], expected[row["primary_key"]])
                                self.assertEqual(row["Extra"], "")
                            for row in file.preview:
                                self.assertEqual(row["doclass"] or "", expected[row["primary_key"]])
        self.assertEqual(records, original)

    def test_direct_template_writers_derive_doclass_and_reject_duplicate_class_headers(self):
        record = {**read_source(source_file(), "NEGATIVE", 1).records[0],
                  "AREA": "Region 3", "doclass": "Old template value"}
        csv_data = fill_csv_template(create_default_csv_template(), [record])
        self.assertEqual(list(csv.DictReader(StringIO(csv_data.decode("cp1252"))))[0]["doclass"], "SBC GEO R3")
        workbook = loaded(fill_excel_template(create_default_template(), [record], "Template", 1))
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.cell(2, DEFAULT_TEMPLATE_HEADERS.index("doclass") + 1).value, "SBC GEO R3")
        self.assertEqual(record["doclass"], "Old template value")
        duplicate = (",".join([*DEFAULT_TEMPLATE_HEADERS, "DO CLASS"]) + "\n").encode()
        with self.assertRaisesRegex(AutomationError, 'Duplicate or ambiguous header "doclass"'):
            fill_csv_template(duplicate, [record])

    def test_geo_headers_map_aliases_and_copy_area(self):
        data = geo_source_file([sample_row(Area="Region 3", level="Level 6")])
        info = next(sheet for sheet in inspect_workbook(data) if sheet.name == "NEGATIVE")
        self.assertEqual(info.matched_fields, len(FIELD_MAPPING))
        self.assertTrue(info.has_area)
        mapped = read_source(data, "NEGATIVE", 1)
        self.assertEqual(mapped.records[0], {
            **read_source(source_file([sample_row(level="Level 6")]), "NEGATIVE", 1).records[0],
            "AREA": "Region 3",
        })
        self.assertFalse(next(sheet for sheet in inspect_workbook(source_file()) if sheet.name == "NEGATIVE").has_area)

    def test_geo_aliases_and_area_cannot_be_ambiguous(self):
        for extra_header in ("debtor_email", "AREA"):
            with self.subTest(extra_header=extra_header):
                workbook = loaded(geo_source_file([sample_row(Area="Mindanao")]))
                sheet = workbook["NEGATIVE"]
                sheet.cell(1, sheet.max_column + 1, extra_header)
                with self.assertRaisesRegex(AutomationError, "Duplicate or ambiguous"):
                    read_source(save(workbook), "NEGATIVE", 1)

    def test_area_product_groups_combine_all_levels_and_retain_review_accounts(self):
        levels = [1, 2, 3, 4, 5, 6, None]
        cases = [(status, "CARDS", level) for status in ("QUALIFIED", None) for level in levels]
        cases.extend([
            ("QUALIFIED", "BEL", 3), ("QUALIFIED", "PL", 5),
            (None, "BEL", 5), (None, "PL", 3),
        ])
        rows = [
            sample_row(Area=area, product=product, level=level, **{"Christmas Amnesty": status})
            for area in (" Mindanao ", "REGION 3") for status, product, level in cases
        ]
        rows.extend([
            sample_row(Area="R3", product="BEL", **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(Area="Other", product="CARDS", level=1),
            sample_row(Area=None, product="BEL"),
            sample_row(Area="Mindanao", debtor_email=None, phone_1=None),
        ])
        for number, row in enumerate(rows):
            row["Account no."] = f"{number:018d}"
        expected_status = {row["Account no."]: row.get("Christmas Amnesty") or "" for row in rows}
        mapped = read_source(geo_source_file(rows), "NEGATIVE", 1)
        groups = separate_geo(mapped.records)
        self.assertEqual(mapped.no_contact_rows_skipped, 1)
        self.assertEqual(sum(map(len, groups.values())), len(cases) * 2 + 3)
        self.assertEqual([row["level"] for row in groups["mindanao_qualified_cards"]], levels)
        self.assertEqual([row["level"] for row in groups["mindanao_blank_cards"]], levels)
        self.assertEqual([row["level"] for row in groups["r3_qualified_cards"]], levels)
        self.assertEqual([row["level"] for row in groups["r3_blank_cards"]], levels)
        self.assertFalse(any("level" in key for key in groups))
        self.assertEqual(len(groups["r3_qualified_bel"]), 2)
        self.assertEqual(len(groups["mindanao_blank_other_products"]), 2)
        self.assertEqual(len(groups["r3_blank_other_products"]), 2)
        self.assertEqual(len(groups["mindanao_qualified_review"]), 0)
        self.assertEqual(len(groups["r3_qualified_review"]), 0)
        self.assertEqual(len(groups["area_review_blank_cards"]), 1)
        self.assertEqual(len(groups["area_review_blank_bel"]), 1)
        bundle = build_geo_exports(create_default_csv_template(), mapped.records, "", 1, ".csv")
        self.assertEqual(sum(bundle.counts.values()), len(cases) * 2 + 3)
        self.assertEqual(sum(file.count for file in bundle.files), (len(cases) * 2 + 3) * 2)
        with ZipFile(BytesIO(bundle.zip_data)) as archive:
            self.assertEqual(archive.namelist(), [file.filename for file in bundle.files])
            for file in bundle.files:
                records = list(csv.DictReader(StringIO(archive.read(file.filename).decode("cp1252"))))
                self.assertEqual(len(records), file.count)
                for record in records:
                    if file.group_key.startswith("mindanao_"):
                        self.assertEqual(record["AREA"].strip(), "Mindanao")
                        self.assertEqual(record["doclass"], "SBC GEO Mindanao")
                    elif file.group_key.startswith("r3_"):
                        self.assertIn(record["AREA"], ("REGION 3", "R3"))
                        self.assertEqual(record["doclass"], "SBC GEO R3")
                    else:
                        self.assertIn(file.group_key, GEO_REVIEW_GROUP_KEYS)
                        self.assertIn(record["AREA"], ("Other", ""))
                        self.assertEqual(record["doclass"], "")
                    self.assertEqual(record["ChristmasAmnesty"], expected_status[record["primary_key"]])
                    self.assertEqual(record["phone" if file.channel == "email" else "email"], "")

    def test_five_main_groups_per_area_produce_twenty_channel_files(self):
        rows = [
            sample_row(Area=area, product=product, level=level, **{"Christmas Amnesty": status})
            for area in ("Mindanao", "Region 3")
            for status, product in (("QUALIFIED", "CARDS"), ("QUALIFIED", "BEL"),
                                    (None, "CARDS"), (None, "BEL"), (None, "PL"))
            for level in (1, 2, 3, 4, 5, 6)
        ]
        records = read_source(geo_source_file(rows), "NEGATIVE", 1).records
        bundle = build_geo_exports(create_default_csv_template(), records, "", 1, ".csv")
        self.assertEqual(len(bundle.files), 20)
        self.assertEqual(sum(bundle.counts.values()), 60)
        self.assertTrue(all(file.count == 6 for file in bundle.files))
        self.assertTrue(all("LEVEL" not in file.filename and " - L" not in file.filename for file in bundle.files))
        self.assertFalse(any(bundle.counts[key] for key in GEO_REVIEW_GROUP_KEYS))

    def test_geo_channel_exclusions_header_only_files_and_configured_names(self):
        data = geo_source_file([
            sample_row(Area="Mindanao", product="BEL", debtor_email="bad..email@example.com", **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(Area="Region 3", product="CARDS", level=1, phone_1=None),
            sample_row(Area="Region 3", debtor_email="bad email@example.com", phone_1=None),
        ])
        mapped = read_source(data, "NEGATIVE", 1)
        self.assertEqual(mapped.no_contact_rows_skipped, 1)
        self.assertEqual(mapped.invalid_email_rows, 2)
        with patch.dict(GEO_OUTPUT_FILENAMES, {"mindanao_qualified_bel": "Custom.xlsx"}):
            bundle = build_geo_exports(create_default_template(), mapped.records, "Template", 1)
        self.assertEqual([file.count for file in bundle.files], [0, 1, 1, 0])
        self.assertEqual(bundle.files[0].filename, "Custom_Email.csv")
        for file in (bundle.files[0], bundle.files[-1]):
            self.assertEqual(list(csv.reader(StringIO(file.data.decode("cp1252")))), [list(DEFAULT_TEMPLATE_HEADERS)])
        with patch.dict(GEO_OUTPUT_FILENAMES, {"r3_blank_cards": GEO_OUTPUT_FILENAMES["mindanao_blank_cards"]}):
            with self.assertRaisesRegex(AutomationError, "unique"):
                build_geo_exports(create_default_template(), mapped.records, "Template", 1)
        with self.assertRaisesRegex(AutomationError, "AREA header"):
            build_geo_exports((",".join(TARGET_HEADERS) + "\n").encode(), mapped.records, "", 1, ".csv")


class FolderOutputTests(unittest.TestCase):
    def test_extracted_csv_can_be_edited_and_saved_in_place(self):
        records = read_source(source_file(), "NEGATIVE", 1).records
        bundle = build_sheet_exports(
            create_default_template(), records, "Template", 1, base_name="Campaign", label="Test",
        )
        with TemporaryDirectory(prefix="blasting-extracted-test-") as directory:
            with ZipFile(BytesIO(bundle.zip_data)) as archive:
                archive.extractall(directory)
            filename = Path(directory) / bundle.files[0].filename
            edited = filename.read_bytes().replace(b"Ana Cruz", b"Ana Updated")
            filename.write_bytes(edited)
            rows = list(csv.DictReader(StringIO(filename.read_text(encoding="cp1252"))))
            self.assertEqual(rows[0]["FullName"], "Ana Updated")

    def test_all_files_saved_without_zip_and_previous_exports_kept(self):
        with TemporaryDirectory(prefix="blasting-folder-test-") as directory:
            parent = Path(directory).resolve()
            files = [("first.csv", b"email\r\nana@example.com\r\n"), ("second.csv", b"email\r\n")]
            with patch("folder_output.datetime") as test_clock:
                test_clock.now.return_value = datetime(2026, 10, 26, 21, 0)
                first = save_files_to_folder(files, parent, batch_name="NEGATIVE")
                (first / "second.csv").write_bytes(b"previous export")
                second = save_files_to_folder(files, parent, batch_name="NEGATIVE")
            self.assertNotEqual(first, second)
            self.assertEqual(first.parent, parent)
            self.assertEqual(second.parent, parent)
            self.assertEqual((first / "second.csv").read_bytes(), b"previous export")
            self.assertEqual({file.name for file in second.iterdir()}, {name for name, _ in files})
            for name, data in files:
                self.assertEqual((second / name).read_bytes(), data)
            self.assertEqual((second / "first.csv").read_bytes(), files[0][1])

    def test_relative_output_folder_uses_project_directory(self):
        with TemporaryDirectory(prefix="blasting-folder-test-") as directory:
            project = Path(directory).resolve()
            with patch("folder_output.PROJECT_DIRECTORY", project):
                folder = save_files_to_folder(
                    [("data.csv", b"email\r\n")], "output", batch_name="../NEGATIVE",
                )
            self.assertEqual(folder.parent, project / "output")
            self.assertTrue(folder.is_relative_to(project))

    def test_invalid_and_duplicate_names_fail_before_writing(self):
        with TemporaryDirectory(prefix="blasting-folder-test-") as directory:
            parent = Path(directory).resolve() / "not_created"
            invalid_batches = [
                [], [("../outside.csv", b"data")], [("bad:name.csv", b"data")],
                [("same.csv", b"one"), ("SAME.csv", b"two")], [("not_supported.exe", b"data")],
            ]
            for files in invalid_batches:
                with self.subTest(files=[name for name, _ in files]):
                    with self.assertRaises(AutomationError):
                        save_files_to_folder(files, parent, batch_name="NEGATIVE")
                    self.assertFalse(parent.exists())

    def test_existing_file_as_destination_is_not_changed(self):
        with TemporaryDirectory(prefix="blasting-folder-test-") as directory:
            destination = Path(directory).resolve() / "already_a_file"
            destination.write_bytes(b"keep this")
            with self.assertRaisesRegex(AutomationError, "Could not save the files"):
                save_files_to_folder([("data.csv", b"email\r\n")], destination, batch_name="NEGATIVE")
            self.assertEqual(destination.read_bytes(), b"keep this")


class AppTests(unittest.TestCase):
    def test_partial_upload_maps_available_values_with_csv_and_excel_templates(self):
        headers = ['debtor_email', 'OLD I.C', 'Account no.', 'financing/Card no.', 'Name']
        upload = FakeUpload('partial.xlsx', source_file(headers=headers))
        for template in (None, FakeUpload('template.xlsx', create_default_template())):
            with self.subTest(excel_template=template is not None), patch(
                'streamlit.file_uploader', side_effect=lambda label, **kwargs: upload if kwargs.get('key') == 'source_upload' else template if kwargs.get('key') == 'template_upload' else None,
            ):
                app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
                self.assertEqual(len(app.exception), 0)
                self.assertEqual(app.radio[0].options, ['All accounts in one group'])
                app.button[0].click().run(timeout=20)
                self.assertEqual(len(app.exception), 0)
                self.assertEqual(len(app.error), 0)
                result = app.session_state['result']
                self.assertFalse(result['grouped'])
                self.assertEqual([file.count for file in result['files']], [1, 0])
                values = list(csv.DictReader(StringIO(result['files'][0].data.decode('cp1252'))))[0]
                self.assertEqual(values['FullName'], 'Ana Cruz')
                self.assertEqual(values['cardnumber'], sample_row()['financing/Card no.'])
                for field in ('EndoDate', 'Level', 'Product', 'OB', 'phone', 'ChristmasAmnesty'):
                    self.assertEqual(values[field], '')

    def test_partial_upload_keeps_grouping_when_group_fields_are_present(self):
        for geo in (False, True):
            with self.subTest(geo=geo):
                row = {'primary_key': '0001', 'email': 'first@example.com', 'PHONE': '09171234567',
                       'Product': 'CARDS', 'ChristmasAmnesty': 'QUALIFIED'}
                row.update({'Area': 'Mindanao'} if geo else {'Level': 1})
                upload = FakeUpload('partial-grouped.xlsx', source_file([row], headers=list(row)))
                with patch('streamlit.file_uploader', side_effect=lambda label, **kwargs: upload if kwargs.get('key') == 'source_upload' else None):
                    app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
                    app.button[0].click().run(timeout=20)
                    self.assertEqual(len(app.exception), 0)
                    self.assertEqual(len(app.error), 0)
                    result = app.session_state['result']
                    self.assertTrue(result['grouped'])
                    self.assertEqual(result['geo'], geo)
                    expected_group = 'mindanao_qualified_cards' if geo else 'qualified_cards_level_1'
                    self.assertEqual(result['group_counts'][expected_group], 1)
                    self.assertTrue(all(file.group_key == expected_group for file in result['files']))

    def test_manual_header_change_rechecks_grouping_fields(self):
        upload = FakeUpload('late-headers.xlsx', source_file(title_rows=59, other_sheet=False))
        with patch('streamlit.file_uploader', side_effect=lambda label, **kwargs: upload if kwargs.get('key') == 'source_upload' else None):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            self.assertEqual(app.radio[0].options, ['All accounts in one group'])
            field = next(widget for widget in app.number_input if widget.label == 'Source header row')
            field.set_value(60).run(timeout=20)
            self.assertEqual(app.radio[0].value, 'Separate Christmas Amnesty groups')
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(app.session_state['result']['count'], 1)

    def test_other_worksheets_default_to_negative_rules_and_remember_their_own_mode(self):
        rows = [
            sample_row(product="CARDS", level=1, **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(product="CARDS", level=6), sample_row(product="BEL", level=2),
            sample_row(product="PL", level=3),
        ]
        workbook = loaded(source_file(rows))
        for name in ("POSITIVE", "CONTACTED"):
            sheet = workbook.create_sheet(name)
            sheet.append(list(FIELD_MAPPING))
            for row in rows:
                sheet.append([name if header == "Name" else row.get(header) for header in FIELD_MAPPING])
        upload = FakeUpload("accounts.xlsx", save(workbook))
        with patch("streamlit.file_uploader", side_effect=lambda label, **kwargs: upload if kwargs.get("key") == "source_upload" else None), TemporaryDirectory(prefix="other-sheet-test-") as directory:
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            self.assertEqual(app.selectbox[0].value, "NEGATIVE")
            app.button[0].click().run(timeout=20)
            baseline = app.session_state["result"]
            app.selectbox[0].select("POSITIVE").run(timeout=20)
            self.assertNotIn("result", app.session_state)
            self.assertEqual(app.radio[0].value, "Separate Christmas Amnesty groups")
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state["result"]
            self.assertTrue(result["grouped"])
            self.assertFalse(result["geo"])
            self.assertEqual(result["group_counts"], baseline["group_counts"])
            self.assertEqual([file.filename for file in result["files"]], ["POSITIVE - " + file.filename for file in baseline["files"]])
            self.assertTrue(all(row["FullName"] == "POSITIVE" for file in result["files"] for row in file.preview))
            self.assertEqual(app.main.dataframe[-1].value.loc[0, "FullName"], "POSITIVE")
            app.text_input(key="output_folder").set_value(directory).run(timeout=20)
            app.button(key="save_to_folder").click().run(timeout=20)
            self.assertEqual(len(app.error), 0)
            saved = Path(app.session_state["result"]["saved_folder"])
            for file in result["files"]:
                self.assertEqual((saved / file.filename).read_bytes(), build_export_csv(file, protect_identifiers=True))
                self.assertEqual(list(csv.DictReader(StringIO(file.data.decode("cp1252"))))[0]["FullName"], "POSITIVE")
            app.radio[0].set_value("All accounts in one group").run(timeout=20)
            app.selectbox[0].select("CONTACTED").run(timeout=20)
            self.assertEqual(app.radio[0].value, "Separate Christmas Amnesty groups")
            app.button[0].click().run(timeout=20)
            self.assertTrue(all(file.filename.startswith("CONTACTED - ") for file in app.session_state["result"]["files"]))
            app.selectbox[0].select("POSITIVE").run(timeout=20)
            self.assertEqual(app.radio[0].value, "All accounts in one group")
            app.button[0].click().run(timeout=20)
            self.assertFalse(app.session_state["result"]["grouped"])
            self.assertEqual(len(app.session_state["result"]["files"]), 2)

    def test_grouping_defaults_on_workbooks_without_negative_and_on_other_geo_sheets(self):
        for geo in (False, True):
            with self.subTest(geo=geo):
                data = (geo_source_file([
                    sample_row(Area="Mindanao", product="CARDS", level=1, **{"Christmas Amnesty": "QUALIFIED"}),
                    sample_row(Area="Mindanao", product="CARDS", level=6, **{"Christmas Amnesty": "QUALIFIED"}),
                    sample_row(Area="Region 3", product="BEL"),
                ]) if geo else source_file([sample_row(product="CARDS", level=1, **{"Christmas Amnesty": "QUALIFIED"})], other_sheet=False))
                workbook = loaded(data)
                if "Summary" in workbook.sheetnames:
                    del workbook["Summary"]
                workbook["NEGATIVE"].title = "ACTIVE GEO" if geo else "PROMISE TO PAY"
                upload = FakeUpload("other.xlsx", save(workbook))
                with patch("streamlit.file_uploader", side_effect=lambda label, **kwargs: upload if kwargs.get("key") == "source_upload" else None):
                    app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
                    name = "ACTIVE GEO" if geo else "PROMISE TO PAY"
                    self.assertEqual(app.selectbox[0].value, name)
                    self.assertEqual(app.radio[0].value, "Separate Area and Christmas Amnesty groups" if geo else "Separate Christmas Amnesty groups")
                    app.button[0].click().run(timeout=20)
                    self.assertEqual(len(app.exception), 0)
                    self.assertEqual(len(app.error), 0)
                    result = app.session_state["result"]
                    self.assertTrue(result["grouped"])
                    self.assertEqual(result["geo"], geo)
                    self.assertTrue(all(file.filename.startswith(name + " - ") for file in result["files"]))
                    self.assertTrue(result["zip_filename"].startswith(name + " - "))
                    if geo:
                        self.assertEqual(result["group_counts"]["mindanao_qualified_cards"], 2)
                        self.assertEqual(len(result["files"]), 4)
                        for file in result["files"]:
                            for row in file.preview:
                                self.assertEqual(row["doclass"], "SBC GEO Mindanao" if row["AREA"] == "Mindanao" else "SBC GEO R3")

    def test_configured_geo_names_apply_to_downloads_preview_and_folder(self):
        upload = FakeUpload("GEO.xlsx", geo_source_file([
            sample_row(Area="Mindanao", product="CARDS", **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(Area="Region 3", product="BEL"),
        ]))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else None

        names = {"mindanao_qualified_cards": "My Mindanao Cards", "r3_blank_bel": "My R3 BEL.csv"}
        with TemporaryDirectory(prefix="blasting-config-test-") as directory, patch.dict(
            GEO_OUTPUT_FILENAMES, names,
        ), patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            app.selectbox(key="campaign_output_format").set_value("CSV (.csv)").run(timeout=20)
            app.checkbox(key="campaign_csv_protect_ids").set_value(False).run(timeout=20)
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state["result"]
            files = result["files"]
            expected = ["My Mindanao Cards_Email.csv", "My Mindanao Cards_Viber.csv",
                        "My R3 BEL_Email.csv", "My R3 BEL_Viber.csv"]
            self.assertEqual([file.filename for file in files], expected)
            self.assertEqual([button.label for button in app.get("download_button") if button.label.endswith(".csv")],
                             [f"Download {expected[0]}"])
            self.assertFalse(any(field.label == "ZIP filename" for field in app.text_input))
            with ZipFile(BytesIO(result["zip_data"])) as archive:
                self.assertEqual(archive.namelist(), expected)
                for file in files:
                    self.assertEqual(archive.read(file.filename), file.data)
            app.main.selectbox[-1].select("My R3 BEL_Email.csv").run(timeout=20)
            self.assertEqual(app.main.dataframe[-1].value.loc[0, "AREA"], "Region 3")
            app.text_input(key="output_folder").set_value(directory).run(timeout=20)
            app.button(key="save_to_folder").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            saved = Path(app.session_state["result"]["saved_folder"])
            self.assertEqual({path.name for path in saved.iterdir()}, set(expected))
            for file in files:
                self.assertEqual((saved / file.filename).read_bytes(), file.data)
            app.run(timeout=20)
            self.assertEqual(app.session_state["result"]["saved_folder"], str(saved))
            GEO_OUTPUT_FILENAMES["mindanao_qualified_cards"] = "Updated Mindanao"
            app.run(timeout=20)
            self.assertNotIn("result", app.session_state)
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.session_state["result"]["files"][0].filename, "Updated Mindanao_Email.csv")
            self.assertTrue((saved / expected[0]).exists())

    def test_geo_email_validation_counts_every_account_in_all_product_groups(self):
        rows = []
        for area in ("Mindanao", "Region 3"):
            for status, product in (("QUALIFIED", "CARDS"), ("QUALIFIED", "BEL"), (None, "CARDS"), (None, "BEL"), (None, "PL")):
                for email, phone in ((" Ana.Cruz+tag@EXAMPLE.COM ", "09000000000"),
                                     ("bad..email@example.com", "09000000000"),
                                     ("bad email@example.com", None), (None, "09000000000"), (None, None)):
                    rows.append(sample_row(Area=area, product=product, debtor_email=email, phone_1=phone, **{"Christmas Amnesty": status}))
        rows.insert(1, {})
        upload = FakeUpload("GEO.xlsx", geo_source_file(rows))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else None

        with patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state["result"]
            self.assertEqual(result["emails_checked"], 50)
            self.assertEqual(result["valid_emails"], 10)
            self.assertEqual(result["invalid_emails"], 20)
            self.assertEqual(result["missing_email_rows"], 20)
            self.assertEqual(result["blank_rows"], 1)
            self.assertEqual(result["count"], 30)
            self.assertEqual(result["no_contact_rows"], 20)
            self.assertEqual(len(result["files"]), 20)
            metrics = {metric.label: metric.value for metric in app.metric}
            self.assertEqual(metrics["Accounts checked for email"], "50")
            for file in result["files"]:
                exported = list(csv.DictReader(StringIO(file.data.decode("cp1252"))))
                if file.channel == "email":
                    self.assertEqual([row["email"] for row in exported], ["Ana.Cruz+tag@example.com"])
                else:
                    self.assertEqual(len(exported), 3)
                    self.assertTrue(all(row["email"] == "" for row in exported))

    def test_geo_area_grouping_preview_review_and_single_group_mode(self):
        upload = FakeUpload("GEO ACTIVE ACCOUNTS.xlsx", geo_source_file([
            sample_row(Area="Mindanao", product="CARDS", level=6, **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(Area="Mindanao", product="CARDS", level=2, **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(Area="Region 3", product="BEL", debtor_email=None),
            sample_row(Area="Mindanao", product="PL", phone_1=None),
            sample_row(Area="Other", product="CARDS", level=2, **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(Area="Mindanao", debtor_email=None, phone_1=None),
        ]))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else None

        with patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            self.assertEqual(app.radio[0].value, "Separate Area and Christmas Amnesty groups")
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state["result"]
            self.assertEqual(result["count"], 5)
            self.assertEqual(result["no_contact_rows"], 1)
            self.assertEqual(len(result["files"]), 8)
            self.assertEqual(result["group_counts"]["mindanao_qualified_cards"], 2)
            self.assertFalse(any("level" in key for key in result["group_counts"]))
            self.assertTrue(any("need review" in warning.value for warning in app.warning))
            file = next(file for file in result["files"] if file.group_key == "r3_blank_bel" and file.channel == "viber")
            app.main.selectbox[-1].select(file.filename).run(timeout=20)
            self.assertEqual(app.main.dataframe[-1].value.loc[0, "AREA"], "Region 3")
            self.assertEqual(app.main.dataframe[-1].value.loc[0, "doclass"], "SBC GEO R3")
            app.radio[0].set_value("All accounts in one group").run(timeout=20)
            self.assertEqual(len(app.get("download_button")), 5)
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(len(app.session_state["result"]["files"]), 2)
            self.assertEqual(app.session_state["result"]["count"], 5)
            for file in app.session_state["result"]["files"]:
                for row in file.preview:
                    expected = "SBC GEO Mindanao" if row["AREA"] == "Mindanao" else "SBC GEO R3" if row["AREA"] == "Region 3" else None
                    self.assertEqual(row["doclass"], expected)

    def test_file_type_defaults_to_standard_csv_and_change_refreshes_exports(self):
        upload = FakeUpload("source.xlsx", source_file([sample_row(Name="Pe\u00f1a, Ana")]))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else None

        with patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            self.assertEqual(app.selectbox(key="csv_file_type").value, "CSV (Comma delimited)")
            for output_format, encoding, has_bom in (
                ("CSV (Comma delimited)", "cp1252", False),
                ("CSV UTF-8 (Comma delimited)", "utf-8-sig", True),
            ):
                with self.subTest(output_format=output_format):
                    app.selectbox(key="csv_file_type").set_value(output_format).run(timeout=20)
                    self.assertEqual(len(app.get("download_button")), 5)
                    app.button[0].click().run(timeout=20)
                    self.assertEqual(len(app.exception), 0)
                    self.assertEqual(len(app.error), 0)
                    result = app.session_state["result"]
                    with ZipFile(BytesIO(result["zip_data"])) as archive:
                        for file in result["files"]:
                            self.assertEqual(file.data.startswith(b"\xef\xbb\xbf"), has_bom)
                            self.assertEqual(archive.read(file.filename), file.data)
                            rows = list(csv.DictReader(StringIO(file.data.decode(encoding))))
                            self.assertEqual(rows[0]["FullName"], "Pe\u00f1a, Ana")
                            self.assertEqual(list(rows[0]), list(DEFAULT_TEMPLATE_HEADERS))

    def test_invalid_emails_removed_and_reported_in_both_export_modes(self):
        invalid_only = sample_row(debtor_email="bad..email@example.com", phone_1=None)
        cases = [
            ([invalid_only], 0, 1),
            ([invalid_only, sample_row(debtor_email="bad email@example.com"),
              sample_row(debtor_email=" Ana.Cruz+tag@EXAMPLE.COM ", phone_1=None)], 2, 2),
        ]
        for rows, included, invalid in cases:
            for mode in ("Separate Christmas Amnesty groups", "All accounts in one group"):
                with self.subTest(included=included, mode=mode):
                    upload = FakeUpload("source.xlsx", source_file(rows))

                    def uploader(label, **kwargs):
                        return upload if kwargs.get("key") == "source_upload" else None

                    with patch("streamlit.file_uploader", side_effect=uploader):
                        app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
                        app.radio[0].set_value(mode).run(timeout=20)
                        app.button[0].click().run(timeout=20)
                        self.assertEqual(len(app.exception), 0)
                        self.assertEqual(len(app.error), 0)
                        result = app.session_state["result"]
                        self.assertEqual(result["count"], included)
                        self.assertEqual(result["no_contact_rows"], 1)
                        self.assertEqual(result["invalid_emails"], invalid)
                        metrics = {metric.label: metric.value for metric in app.metric}
                        self.assertEqual(metrics["Invalid emails removed"], str(invalid))
                        if included:
                            self.assertEqual([file.count for file in result["files"]], [1, 1])
                            email_file = next(file for file in result["files"] if file.channel == "email")
                            self.assertEqual(email_file.preview[0]["email"], "Ana.Cruz+tag@example.com")
                        else:
                            self.assertEqual(len(app.get("download_button")), 1)

    def test_initial_screen(self):
        app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
        self.assertEqual(len(app.exception), 0)
        self.assertIn("Upload a source workbook", app.info[0].value)

    def test_upload_preview_download_and_invalidation(self):
        workbook = loaded(source_file())
        sheet = workbook.create_sheet("NEXT")
        sheet.append(list(FIELD_MAPPING))
        sheet.append(list(sample_row(Name="Next person").values()))
        upload = FakeUpload("source.xlsx", save(workbook))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else None

        with patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.selectbox[0].value, "NEGATIVE")
            app.radio[0].set_value("All accounts in one group").run(timeout=20)
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(app.session_state["result"]["count"], 1)
            self.assertEqual(len(app.get("download_button")), 7)
            files = app.session_state["result"]["files"]
            self.assertTrue(all(file.filename.endswith(".csv") for file in files))
            rows = list(csv.DictReader(StringIO(files[0].data.decode("utf-8-sig"))))
            self.assertEqual(rows[0]["FullName"], "Ana Cruz")
            # A routine rerun must retain the prepared download.
            app.run(timeout=20)
            self.assertEqual(app.session_state["result"]["count"], 1)
            app.selectbox[0].select("NEXT").run(timeout=20)
            self.assertEqual(len(app.success), 0)
            self.assertEqual(len(app.get("download_button")), 5)

    def test_uploaded_csv_template_and_optional_missing_phone_header(self):
        upload = FakeUpload("source.xlsx", source_file())
        headers = ["Custom Field"] + list(reversed(TARGET_HEADERS))
        template = FakeUpload("template.csv", (";".join(headers) + "\n").encode())

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else template if kwargs.get("key") == "template_upload" else None

        with patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            app.radio[0].set_value("All accounts in one group").run(timeout=20)
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual([file.filename for file in app.session_state["result"]["files"]], [
                "source_NEGATIVE_filled_Email.csv", "source_NEGATIVE_filled_Viber.csv",
            ])
            preview = app.main.dataframe[-1].value
            self.assertEqual(list(preview.columns), [*headers, "Monthly", "Terms"])
            self.assertEqual(preview.loc[0, "Custom Field"], "")
            self.assertEqual(preview.loc[0, "FullName"], "Ana Cruz")
            upload = FakeUpload("bad.xlsx", source_file(headers=[header for header in FIELD_MAPPING if header != "phone_1"]))
            app.run(timeout=20)
            app.radio[0].set_value("All accounts in one group").run(timeout=20)
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual([file.count for file in app.session_state['result']['files']], [1, 0])
            self.assertEqual(len(app.get("download_button")), 7)

    def test_uploaded_excel_template_produces_csv_in_both_export_modes(self):
        upload = FakeUpload("source.xlsx", source_file([
            sample_row(product="CARDS", level=1, **{"Christmas Amnesty": "QUALIFIED"}),
        ]))
        headers = ["Extra"] + list(reversed(TARGET_HEADERS))
        workbook = Workbook()
        workbook.active.title = "Instructions"
        workbook.active.append(["Instructions are not exported"])
        sheet = workbook.create_sheet("Layout")
        sheet.append(["Template title is not exported"])
        sheet.append(headers)
        sheet.append(["Stale data is not exported"] * len(headers))
        template = FakeUpload("template.xlsx", save(workbook))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else template if kwargs.get("key") == "template_upload" else None

        with patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            self.assertEqual(app.selectbox[1].value, "Layout")
            for mode in ("Separate Christmas Amnesty groups", "All accounts in one group"):
                with self.subTest(mode=mode):
                    app.radio[0].set_value(mode).run(timeout=20)
                    app.button[0].click().run(timeout=20)
                    self.assertEqual(len(app.exception), 0)
                    self.assertEqual(len(app.error), 0)
                    result = app.session_state["result"]
                    self.assertEqual(len(result["files"]), 2)
                    preview = app.main.dataframe[-1].value
                    self.assertEqual(list(preview.columns), [*headers, "Monthly", "Terms"])
                    self.assertEqual(preview.loc[0, "Extra"], "")
                    self.assertEqual(preview.loc[0, "FullName"], "Ana Cruz")
                    self.assertEqual(preview.loc[0, "endodate"], "2026-10-26 00:00:00")
                    with ZipFile(BytesIO(result["zip_data"])) as archive:
                        self.assertEqual(archive.namelist(), [file.filename for file in result["files"]])
                        for file in result["files"]:
                            self.assertTrue(file.filename.endswith(".csv"))
                            self.assertEqual(archive.read(file.filename), file.data)
                            rows = list(csv.reader(StringIO(file.data.decode("utf-8-sig"))))
                            self.assertEqual(rows[0], [*headers, "Monthly", "Terms"])
                            self.assertEqual(len(rows), 2)
                            values = dict(zip(headers, rows[1]))
                            self.assertEqual(values["Extra"], "")
                            self.assertEqual(values["primary_key"], "0000000456")
                            self.assertEqual(values["FullName"], "Ana Cruz")
                            self.assertEqual(values["endodate"], "2026-10-26 00:00:00")
                            self.assertEqual(values["phone" if file.channel == "email" else "email"], "")

    def test_grouped_exports_default_preview_and_mode_change(self):
        rows = [sample_row(**{"product": product, "level": level, "Christmas Amnesty": status})
                for status, product, level in [
                    ("QUALIFIED", "CARDS", "Level 1"),
                    ("QUALIFIED", "CARDS", "Level 6"),
                    ("QUALIFIED", "BEL", "Level 3"),
                    (None, "CARDS", "Level 1"),
                    (None, "CARDS", "Level 6"),
                    (None, "BEL", "Level 5"),
                    (None, "PL", "Level 6"),
                ]]
        rows.extend([
            sample_row(debtor_email=None, phone_1=None, **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(debtor_email=" ", phone_1="\t"),
        ])
        upload = FakeUpload("source.xlsx", source_file(rows))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else None

        with patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            self.assertEqual(app.radio[0].value, "Separate Christmas Amnesty groups")
            self.assertTrue(any("export_config.py" in caption.value for caption in app.caption))
            app.button[0].click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(len(app.session_state["result"]["files"]), 14)
            for file in app.session_state["result"]["files"]:
                self.assertEqual(csv_template_headers(file.data), list(DEFAULT_TEMPLATE_HEADERS))
                rows = list(csv.DictReader(StringIO(file.data.decode("utf-8-sig"))))
                for row in rows:
                    for header in ("doclass", "AREA", "NewEndo"):
                        self.assertEqual(row[header], "")
                    self.assertEqual((row["Monthly"], row["Terms"]),
                                     tuple(app.table[0].value.loc[0, ["Monthly", "Terms"]]))
            preview = app.main.dataframe[-1].value
            self.assertEqual(list(preview.columns), list(DEFAULT_TEMPLATE_HEADERS))
            self.assertEqual(preview.loc[0, "Level"], "Level 1")
            self.assertEqual(preview.loc[0, "Product"], "CARDS")
            self.assertEqual(preview.loc[0, "Monthly"], app.table[0].value.loc[0, "Monthly"])
            self.assertEqual(app.session_state["result"]["no_contact_rows"], 2)
            self.assertEqual(app.metric[2].value, "2")
            self.assertEqual(len(app.get("download_button")), 7)
            app.main.selectbox[-1].select(OUTPUT_FILENAMES["blank_other_products"] + "_Email.csv").run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.session_state["result"]["count"], 7)
            app.radio[0].set_value("All accounts in one group").run(timeout=20)
            self.assertEqual(len(app.success), 0)
            self.assertEqual(len(app.get("download_button")), 5)

    def test_default_headers_refresh_cached_template_and_remain_in_empty_channel(self):
        upload = FakeUpload("source.xlsx", source_file([
            sample_row(phone_1=None, product="CARDS", level=1, **{"Christmas Amnesty": "QUALIFIED"}),
        ]))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else None

        with patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH))
            app.session_state["blank_template_headers"] = (TARGET_HEADERS, ".csv")
            app.session_state["blank_template"] = (",".join(TARGET_HEADERS) + "\n").encode()
            app.run(timeout=20)
            self.assertEqual(
                csv_template_headers(app.session_state["blank_template"]), list(DEFAULT_TEMPLATE_HEADERS),
            )
            for mode in ("Separate Christmas Amnesty groups", "All accounts in one group"):
                with self.subTest(mode=mode):
                    app.radio[0].set_value(mode).run(timeout=20)
                    app.button[0].click().run(timeout=20)
                    self.assertEqual(len(app.exception), 0)
                    self.assertEqual(len(app.error), 0)
                    files = app.session_state["result"]["files"]
                    self.assertEqual([file.count for file in files], [1, 0])
                    with ZipFile(BytesIO(app.session_state["result"]["zip_data"])) as archive:
                        for file in files:
                            self.assertEqual(csv_template_headers(archive.read(file.filename)), list(DEFAULT_TEMPLATE_HEADERS))
                    app.main.selectbox[-1].select(files[1].filename).run(timeout=20)
                    self.assertEqual(len(app.exception), 0)
                    preview = app.main.dataframe[-1].value
                    self.assertTrue(preview.empty)
                    self.assertEqual(list(preview.columns), list(DEFAULT_TEMPLATE_HEADERS))

    def test_all_accounts_without_contacts_show_exclusions_and_no_download(self):
        upload = FakeUpload("source.xlsx", source_file([sample_row(debtor_email=None, phone_1=None)]))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else None

        with patch("streamlit.file_uploader", side_effect=uploader):
            app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
            for mode in ("Separate Christmas Amnesty groups", "All accounts in one group"):
                app.radio[0].set_value(mode).run(timeout=20)
                app.button[0].click().run(timeout=20)
                self.assertEqual(len(app.exception), 0)
                self.assertEqual(len(app.error), 0)
                self.assertEqual(app.session_state["result"]["count"], 0)
                self.assertEqual(app.session_state["result"]["no_contact_rows"], 1)
                self.assertEqual(len(app.get("download_button")), 1)
                self.assertIn("No accounts remain", app.warning[0].value)

    def test_grouped_files_saved_to_folder_only_when_clicked(self):
        rows = [
            sample_row(product="CARDS", level="Level 1", **{"Christmas Amnesty": "QUALIFIED"}),
            sample_row(product="PL", level="Level 6"),
        ]
        upload = FakeUpload("source.xlsx", source_file(rows))

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else None

        with TemporaryDirectory(prefix="blasting-app-folder-test-") as directory:
            parent = Path(directory).resolve()
            with patch("streamlit.file_uploader", side_effect=uploader):
                app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
                app.button[0].click().run(timeout=20)
                app.text_input(key="output_folder").set_value(str(parent)).run(timeout=20)
                self.assertEqual(list(parent.iterdir()), [])
                app.button(key="save_to_folder").click().run(timeout=20)
                self.assertEqual(len(app.exception), 0)
                self.assertEqual(len(app.error), 0)
                saved = Path(app.session_state["result"]["saved_folder"])
                files = app.session_state["result"]["files"]
                self.assertFalse(app.session_state["result"]["excel_output"])
                self.assertEqual({path.name for path in saved.iterdir()}, {file.filename for file in files})
                for file in files:
                    saved_data = (saved / file.filename).read_bytes()
                    self.assertEqual(saved_data, build_export_csv(file, protect_identifiers=True))
                    self.assertEqual(csv_template_headers(saved_data), list(DEFAULT_TEMPLATE_HEADERS))
                    data = list(csv.DictReader(StringIO(saved_data.decode("cp1252"))))
                    self.assertEqual(len(data), file.count)
                    for row in data:
                        for header, expected in (("primary_key", "0000000456"), ("cardnumber", "000123456789012345"), ("indexoneuid", "000123")):
                            self.assertEqual(row[header], "\t" + expected)
                        self.assertEqual(row["OB"], "1250.75")
                app.run(timeout=20)
                self.assertEqual(list(parent.iterdir()), [saved])

    def test_single_csv_saved_to_folder(self):
        upload = FakeUpload("source.xlsx", source_file())
        template = FakeUpload("template.csv", (",".join(TARGET_HEADERS) + "\n").encode())

        def uploader(label, **kwargs):
            return upload if kwargs.get("key") == "source_upload" else template if kwargs.get("key") == "template_upload" else None

        with TemporaryDirectory(prefix="blasting-app-folder-test-") as directory:
            with patch("streamlit.file_uploader", side_effect=uploader):
                app = AppTest.from_file(str(APP_PATH)).run(timeout=20)
                app.selectbox(key="campaign_output_format").set_value("CSV (.csv)").run(timeout=20)
                app.checkbox(key="campaign_csv_protect_ids").set_value(False).run(timeout=20)
                app.radio[0].set_value("All accounts in one group").run(timeout=20)
                app.button[0].click().run(timeout=20)
                app.text_input(key="output_folder").set_value(directory).run(timeout=20)
                app.button(key="save_to_folder").click().run(timeout=20)
                self.assertEqual(len(app.exception), 0)
                self.assertEqual(len(app.error), 0)
                result = app.session_state["result"]
                saved = Path(result["saved_folder"])
                self.assertEqual({path.name for path in saved.iterdir()}, {file.filename for file in result["files"]})
                self.assertEqual(len(result["files"]), 2)
                for file in result["files"]:
                    self.assertEqual((saved / file.filename).read_bytes(), file.data)


if __name__ == "__main__":
    unittest.main()
