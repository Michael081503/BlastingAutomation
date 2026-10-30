from __future__ import annotations

import csv
import unittest
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from automation import create_default_csv_template
from exports import build_excel_zip, build_geo_exports, build_negative_exports


ROOT = Path(__file__).resolve().parents[1]


def records():
    result = []
    for area in ("Mindanao", "Region 3"):
        for product in ("CSL", "PL", "SLA", " csl ", "pl", "sla"):
            for level in (1, 3, 6, None):
                for status in (None, "QUALIFIED"):
                    number = len(result) + 1
                    result.append({
                        "primary_key": f"{number:018d}", "cardnumber": f"{number + 100:018d}",
                        "indexoneuid": f"UID-{number}", "product": product, "level": level, "AREA": area,
                        "email": f"account{number}@example.com", "phone": "09171234567", "OB": number,
                        "ChristmasAmnesty": status,
                    })
    for product in ("MC", "CARDS", "BEL"):
        result.append({**result[0], "primary_key": product, "product": product})
    return result


def csv_rows(file):
    return list(csv.DictReader(StringIO(file.data.decode("cp1252"))))


class Upload(BytesIO):
    def __init__(self, data):
        super().__init__(data)
        self.name = "GEO.xlsx"


class GeoLoanRoutingTests(unittest.TestCase):
    def test_geo_loan_codes_are_in_the_correct_area_files_in_csv_and_excel_zips(self):
        source = records()
        original = [dict(row) for row in source]
        bundle = build_geo_exports(create_default_csv_template(), source, "", 1, ".csv")
        expected_ids = {
            "MINDANAO": {row["primary_key"] for row in source[:48]},
            "R3": {row["primary_key"] for row in source[48:96]},
        }
        with ZipFile(BytesIO(bundle.zip_data)) as csv_zip, ZipFile(BytesIO(build_excel_zip(bundle.files))) as excel_zip:
            for area, ids in expected_ids.items():
                for channel in ("Email", "Viber"):
                    filename = f"GEO NEGA LOANS - {area}_{channel}.csv"
                    file = next(file for file in bundle.files if file.filename == filename)
                    self.assertEqual(csv_zip.read(filename), file.data)
                    self.assertEqual({row["primary_key"] for row in csv_rows(file)}, ids)
                    self.assertEqual({row["Product"].strip().upper() for row in csv_rows(file)}, {"CSL", "PL", "SLA"})
                    self.assertEqual({row["ChristmasAmnesty"] for row in csv_rows(file)}, {"", "QUALIFIED"})
                    self.assertIn("LOANS", file.label)
                    workbook = load_workbook(BytesIO(excel_zip.read(Path(filename).with_suffix(".xlsx").name)))
                    headers = [cell.value for cell in workbook.active[1]]
                    self.assertEqual({row[headers.index("primary_key")] for row in list(workbook.active.values)[1:]}, ids)
                    workbook.close()
            loan_ids = expected_ids["MINDANAO"] | expected_ids["R3"]
            for file in bundle.files:
                if "LOANS" not in file.filename:
                    self.assertTrue(loan_ids.isdisjoint(row["primary_key"] for row in csv_rows(file)))
        self.assertEqual(source, original)

    def test_geo_page_downloads_route_each_loan_product_to_its_area_loans_file(self):
        source = records()
        workbook = Workbook()
        workbook.active.title = "NEGATIVE"
        headers = ["primary_key", "cardnumber", "indexoneuid", "email", "phone", "product", "level", "AREA", "ChristmasAmnesty", "OB"]
        workbook.active.append(headers)
        for row in source:
            workbook.active.append([row.get(header) for header in headers])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        upload = Upload(output.getvalue())
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: upload if kw.get("key") == "source_upload" else None), \
                patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            next(button for button in app.button if button.label == "Preview and prepare files").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertTrue(app.session_state["result"]["geo"])
            calls = {call.kwargs.get("key"): call for call in download.call_args_list}
            for area, group in (("MINDANAO", "mindanao"), ("R3", "r3")):
                for channel in ("email", "viber"):
                    call = calls[f"download_{group}_blank_other_products_{channel}"]
                    self.assertEqual(call.args[2], f"GEO NEGA LOANS - {area}_{channel.title()}.csv")
                    data = list(csv.DictReader(StringIO(call.args[1].decode("cp1252"))))
                    self.assertEqual(len(data), 48)
                    self.assertEqual({row["Product"].strip().upper() for row in data}, {"CSL", "PL", "SLA"})
                    self.assertEqual({row["ChristmasAmnesty"] for row in data}, {"", "QUALIFIED"})

    def test_reco_keeps_its_existing_loan_routing_and_product_values(self):
        source = records()[:-3]
        bundle = build_negative_exports(create_default_csv_template(), source, "", 1, ".csv")
        self.assertEqual({file.filename for file in bundle.files}, {
            "RECO NEGA LOANS - PL_Email.csv", "RECO NEGA LOANS - PL_Viber.csv",
            "NEGATIVE_QUALIFIED_REVIEW_Email.csv", "NEGATIVE_QUALIFIED_REVIEW_Viber.csv",
        })
        for file in bundle.files:
            self.assertEqual({row["primary_key"]: row["Product"] for row in csv_rows(file)},
                             {row["primary_key"]: row["product"] for row in source
                              if bool(row["ChristmasAmnesty"]) == (file.group_key == "qualified_review")})


if __name__ == "__main__":
    unittest.main()
