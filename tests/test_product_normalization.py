from __future__ import annotations

import csv
import unittest
from io import BytesIO, StringIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from automation import create_default_csv_template, read_source
from export_config import CHANNEL_SUFFIXES, GEO_OUTPUT_FILENAMES, OUTPUT_FILENAMES
from exports import build_export_excel, build_geo_exports, build_negative_exports, build_sheet_exports


def account(number, product, *, level=1, status=None, area=None):
    return {
        "primary_key": f"{number:018d}", "cardnumber": f"{number + 100:018d}",
        "indexoneuid": f"UID-{number}", "email": f"account{number}@example.com", "phone": "09171234567",
        "product": product, "level": level, "ChristmasAmnesty": status, "AREA": area, "OB": number,
    }


class ProductNormalizationTests(unittest.TestCase):
    def assert_export_values(self, bundle, records):
        expected = {row["primary_key"]: row for row in records}
        with ZipFile(BytesIO(bundle.zip_data)) as archive:
            for file in bundle.files:
                self.assertEqual(archive.read(file.filename), file.data)
                exported = list(csv.DictReader(StringIO(file.data.decode("cp1252"))))
                self.assertEqual(len(exported), file.count)
                for row in exported:
                    original = expected[row["primary_key"]]
                    product = "CARDS" if original["product"].strip().casefold() == "mc" else original["product"]
                    self.assertEqual(row["Product"], product)
                    self.assertEqual(row["cardnumber"], original["cardnumber"])
                self.assertEqual([row["product"] for row in file.preview], [row["Product"] for row in exported])
                workbook = load_workbook(BytesIO(build_export_excel(file)))
                headers = [cell.value for cell in workbook.active[1]]
                self.assertEqual([row[headers.index("Product")] for row in list(workbook.active.values)[1:]],
                                 [row["Product"] for row in exported])
                workbook.close()

    def test_source_product_type_and_manual_product_mapping_normalize_before_export(self):
        products = ["MC", " mc ", "Mc", "mC", "CARDS", "BEL", "PL", "MCC", None]
        for manual in (False, True):
            with self.subTest(manual=manual):
                workbook = Workbook()
                workbook.active.append(["Account No.", "Email", "Custom category" if manual else "Product Type", "Phone"])
                for number, product in enumerate(products, 1):
                    workbook.active.append([f"{number:018d}", f"account{number}@example.com", product, "09171234567"])
                output = BytesIO()
                workbook.save(output)
                workbook.close()
                data = output.getvalue()
                mapping = {"primary_key": 0, "email": 1, "product": 2, "phone": 3} if manual else None
                result = read_source(data, "Sheet", 1, allow_missing=True, source_columns=mapping)
                self.assertEqual([row["product"] for row in result.records], ["CARDS"] * 4 + products[4:])
                self.assertEqual([row["primary_key"] for row in result.records], [f"{number:018d}" for number in range(1, 10)])
                original = load_workbook(BytesIO(data))
                self.assertEqual([cell.value for cell in original.active["C"][1:]], products)
                original.close()

    def test_reco_mc_routes_as_cards_before_status_and_level_groups(self):
        cases = [
            ("MC", "QUALIFIED", 1, "qualified_cards_level_1"),
            (" mc ", "QUALIFIED", 6, "qualified_cards_level_6"),
            ("Mc", None, 1, "blank_cards_level_1"),
            ("mC", None, 6, "blank_cards_level_6"),
            ("MC", "QUALIFIED", 2, "qualified_review"),
            ("MC", None, 2, "blank_review"),
            ("BEL", "QUALIFIED", 3, "qualified_bel"),
            ("BEL", None, 3, "blank_bel"),
            ("PL", None, 1, "blank_other_products"),
        ]
        records = [account(number, product, status=status, level=level)
                   for number, (product, status, level, _) in enumerate(cases, 1)]
        original = [dict(row) for row in records]
        bundle = build_negative_exports(create_default_csv_template(), records, "", 1, ".csv")
        self.assertEqual(sum(bundle.counts.values()), len(records))
        self.assertEqual(len(bundle.files), len(cases) * 2)
        for row, (_, _, _, group) in zip(records, cases):
            self.assertEqual(bundle.counts[group], 1)
            for channel in ("email", "viber"):
                file = next(file for file in bundle.files if file.group_key == group and file.channel == channel)
                self.assertEqual(file.filename, OUTPUT_FILENAMES[group] + CHANNEL_SUFFIXES[channel] + ".csv")
                self.assertEqual(file.records[0]["primary_key"], row["primary_key"])
        self.assert_export_values(bundle, records)
        self.assertEqual(records, original)

    def test_geo_mc_routes_as_cards_before_area_status_and_channel_groups(self):
        records, expected = [], {}
        for area, area_key in (("Mindanao", "mindanao"), ("Region 3", "r3"), ("Other", "area_review")):
            for status, status_key in (("QUALIFIED", "qualified"), (None, "blank")):
                group = f"{area_key}_{status_key}_cards"
                for product, level in zip(("MC", " mc ", "Mc", "mC"), (1, 2, 6, None)):
                    row = account(len(records) + 1, product, area=area, status=status, level=level)
                    records.append(row)
                    expected.setdefault(group, set()).add(row["primary_key"])
        records.extend([account(25, "BEL", area="Mindanao"), account(26, "PL", area="Region 3")])
        expected["mindanao_blank_bel"] = {records[-2]["primary_key"]}
        expected["r3_blank_other_products"] = {records[-1]["primary_key"]}
        original = [dict(row) for row in records]
        bundle = build_geo_exports(create_default_csv_template(), records, "", 1, ".csv")
        self.assertEqual(sum(bundle.counts.values()), len(records))
        self.assertEqual(len(bundle.files), len(expected) * 2)
        for file in bundle.files:
            self.assertEqual({row["primary_key"] for row in file.records}, expected[file.group_key])
            self.assertEqual(file.filename, GEO_OUTPUT_FILENAMES[file.group_key] + CHANNEL_SUFFIXES[file.channel] + ".csv")
        self.assert_export_values(bundle, records)
        self.assertEqual(records, original)

    def test_all_accounts_mode_converts_mc_before_separating_email_and_viber(self):
        records = [account(1, "MC"), account(2, " mc "), account(3, "BEL"), account(4, "PL")]
        original = [dict(row) for row in records]
        bundle = build_sheet_exports(create_default_csv_template(), records, "", 1, base_name="Accounts",
                                     label="All accounts", template_extension=".csv")
        self.assertEqual([file.count for file in bundle.files], [4, 4])
        self.assert_export_values(bundle, records)
        self.assertEqual(records, original)


if __name__ == "__main__":
    unittest.main()
