from __future__ import annotations

import csv
import unittest
from decimal import Decimal
from io import BytesIO, StringIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from automation import create_default_csv_template, create_default_template, ob_sort_key
from autostat import (
    EMAIL_HEADERS, VIBER_HEADERS, SourceSelection, build_autostat_excel,
    build_autostat_exports, inspect_contact_source,
)
from exports import build_geo_exports, build_negative_exports, build_sheet_exports


def source(name, headers, rows):
    if name.endswith(".csv"):
        output = StringIO(newline="")
        csv.writer(output).writerows([headers, *rows])
        data = output.getvalue().encode()
        extension = ".csv"
    else:
        workbook = Workbook()
        for row in [headers, *rows]:
            workbook.active.append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
        extension = ".xlsx"
    return SourceSelection(name, data, inspect_contact_source(data, extension)[0])


def csv_rows(file):
    return list(csv.DictReader(StringIO(file.data.decode("cp1252"))))


class OBOrderingTests(unittest.TestCase):
    def test_amount_formats_sort_numerically_with_invalid_values_after_negative_amounts(self):
        values = ["9", "1,000.50", None, "PHP 20,000", "₱ 2,000", "", "NaN",
                  "invalid", float("inf"), True, 0, "(50.25)", Decimal("10.5")]
        ordered = sorted(values, key=ob_sort_key, reverse=True)
        self.assertEqual(ordered, ["PHP 20,000", "₱ 2,000", "1,000.50", Decimal("10.5"),
                                   "9", 0, "(50.25)", None, "", "NaN", "invalid", float("inf"), True])

    def test_autostat_sorts_across_csv_and_excel_sources_after_channel_deduplication(self):
        first = source("RECO.csv", ["cardnumber", "indexoneuid", "email", "phone", "OB (MASTERLIST}"], [
            ["001", "A", "first@example.com", "09171234567", "9"],
            ["002", "B", "second@example.com", "09171234567", "1,000.00"],
            ["003", "C", "blank@example.com", "09171234567", None],
            ["004", "D", "fourth@example.com", None, "10"],
            ["007", "E", "invalid", "09171234567", "-2"],
            ["008", "F", "zero@example.com", "09171234567", 0],
            ["001", "A", "repeat@example.com", "09181234567", "999999"],
        ])
        second = source("GEO.xlsx", ["cardnumber", "indexoneuid", "email", "phone", "OB"], [
            ["005", "G", "fifth@example.com", "09171234567", 10],
            ["006", "H", "highest@example.com", "09171234567", 20000],
            ["002", "X", "shared@example.com", "09171234567", "2,000.00"],
            ["004", "D", "later@example.com", "09181234567", 500],
            ["009", "I", "invalidob@example.com", "09171234567", "bad"],
        ])
        unchanged = [first.data, second.data]
        bundle = build_autostat_exports([first, second], timestamp=None)
        expected = [
            ["006", "002", "002", "004", "005", "001", "008", "003", "009"],
            ["006", "002", "002", "004", "005", "001", "008", "007", "003", "009"],
        ]
        self.assertEqual(bundle.summary[0]["Duplicate Email rows removed"], 1)
        self.assertEqual(bundle.summary[1]["Duplicate Email rows removed"], 1)
        self.assertEqual(bundle.summary[1]["Duplicate Viber rows removed"], 0)
        self.assertEqual([row["EMAIL/ NUMBER"] for row in csv_rows(bundle.files[0])], [
            "highest@example.com", "shared@example.com", "second@example.com", "fourth@example.com",
            "fifth@example.com", "first@example.com", "zero@example.com", "blank@example.com", "invalidob@example.com",
        ])
        workbook = load_workbook(BytesIO(build_autostat_excel(bundle)))
        with ZipFile(BytesIO(bundle.zip_data)) as archive:
            for file, sheet, accounts, headers in zip(bundle.files, workbook, expected, (EMAIL_HEADERS, VIBER_HEADERS)):
                self.assertEqual(file.headers, headers)
                self.assertEqual([row["Account No."] for row in csv_rows(file)], accounts)
                self.assertEqual([row["Account No."] for row in file.preview], accounts)
                self.assertEqual([cell.value for cell in sheet["B"][1:]], accounts)
                self.assertTrue(all(cell.data_type == "s" for cell in sheet["B"][1:]))
                self.assertEqual(archive.read(file.filename), file.data)
        workbook.close()
        self.assertEqual([first.data, second.data], unchanged)

    def test_geo_and_reco_sort_each_channel_in_grouped_and_all_accounts_modes(self):
        amounts = [("low", "9"), ("high", "1,000"), ("blank", None), ("tie", 1000),
                   ("zero", 0), ("negative", -1), ("bad", "bad")]
        order = [("high", "tie"), ("low",), ("zero",), ("negative",), ("blank", "bad")]
        records = []
        for group, (area, product, level, status) in enumerate([
            ("Mindanao", "CARDS", 1, "QUALIFIED"), ("Region 3", "BEL", 3, None),
            ("Other", "PL", 2, "QUALIFIED"), ("Mindanao", "CARDS", 6, None),
        ]):
            for label, amount in amounts:
                records.append({
                    "primary_key": f"{group}-{label}", "cardnumber": f"00{group}-{label}",
                    "FullName": label, "OB": amount, "email": f"{label}@example.com",
                    "phone": "09171234567", "AREA": area, "product": product,
                    "level": level, "ChristmasAmnesty": status,
                })
        unchanged = [dict(record) for record in records]
        for extension, template, sheet in (
            (".csv", create_default_csv_template(), ""), (".xlsx", create_default_template(), "Template"),
        ):
            for mode in ("GEO", "RECO", "All accounts"):
                with self.subTest(template=extension, mode=mode):
                    if mode == "All accounts":
                        bundle = build_sheet_exports(template, records, sheet, 1,
                            template_extension=extension, base_name="All accounts", label=mode)
                    else:
                        builder = build_geo_exports if mode == "GEO" else build_negative_exports
                        bundle = builder(template, records, sheet, 1, extension)
                    self.assertEqual(sum(bundle.counts.values()), len(records))
                    self.assertEqual(sum(file.count for file in bundle.files), len(records) * 2)
                    with ZipFile(BytesIO(bundle.zip_data)) as archive:
                        for file in bundle.files:
                            rows = csv_rows(file)
                            ids = {row["primary_key"] for row in rows}
                            expected = [record["primary_key"] for labels in order for record in records
                                        if record["FullName"] in labels and record["primary_key"] in ids]
                            self.assertEqual([row["primary_key"] for row in rows], expected)
                            self.assertEqual([row["primary_key"] for row in file.preview], expected)
                            self.assertEqual(archive.read(file.filename), file.data)
                            for row in rows:
                                original = next(record for record in records if record["primary_key"] == row["primary_key"])
                                self.assertEqual(row["OB"], "" if original["OB"] is None else str(original["OB"]))
        self.assertEqual(records, unchanged)

    def test_full_export_is_sorted_before_the_preview_is_limited(self):
        data = source("accounts.csv", ["cardnumber", "email", "OB"], [
            [f"{index:04d}", f"account{index}@example.com", index] for index in range(110)
        ])
        file = build_autostat_exports([data], timestamp=None).files[0]
        self.assertEqual([row["Account No."] for row in csv_rows(file)],
                         [f"{index:04d}" for index in reversed(range(110))])
        self.assertEqual(len(file.preview), 100)
        self.assertEqual(file.preview[-1]["Account No."], "0010")


if __name__ == "__main__":
    unittest.main()
