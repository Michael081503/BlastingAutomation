from __future__ import annotations

import csv
import re
import unittest
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from automation import AutomationError, inspect_workbook, read_source
from autostat import SourceSelection, inspect_contact_source
from b2_autostat import B2ActivitySource, b2_metadata_columns, build_b2_autostat_exports
from b2_matching import B2_TEMPLATE_HEADERS, build_b2_exports, is_datagrid_sheet, matching_sheet_score


ROOT = Path(__file__).resolve().parents[1]
RAW_HEADERS = ['Account Key', 'Contract Number', 'Product Category Code', 'Bucket Code',
               'Phone (01)', 'Phone (02)', 'Email (01)', 'Email (02)']
GRID_HEADERS = ['No', 'Name', 'New I.C.', 'Account No.', 'Email', 'Financing/Card No.']


def workbook_data(headers, rows, *, sheet_name, wrong_dimensions=True):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if sheet_name == 'SBC':
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == 10:
                row[0].number_format = '00000'
                row[1].number_format = '000000'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    if not wrong_dimensions:
        return output.getvalue()
    rewritten = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(rewritten, 'w', ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename.startswith('xl/worksheets/sheet') and entry.filename.endswith('.xml'):
                data = re.sub(rb'(<dimension\b[^>]*\bref=")[^"]+(")', rb'\g<1>A1:C1\2', data, count=1)
            target.writestr(entry, data)
    return rewritten.getvalue()


def fixtures():
    raw = workbook_data(RAW_HEADERS, [
        ['00001', '0000000000000001', 'MC', 2, '09171234567', '0281234567', 'bad..email@example.com', 'first@example.com'],
        ['@BEL-1', '@BEL-1', 'BEL', 2, '0281234567', '+639181234567', 'bel@example.com', None],
        ['00002', '0002', 'MC', 4, '09191234567', None, 'b4@example.com', None],
        ['00003', '0003', 'MC', 2, '09191234567', None, 'database-only@example.com', None],
        ['00004', '0004', 'MC', 2, '09191234567', None, 'four@example.com', None],
        ['00005', '0005', 'MC', 2, '09191234567', None, 'five@example.com', None],
        ['00006', '0006', 'MC', 2, '09191234567', None, 'six@example.com', None],
        ['00006', '0007', 'MC', 2, '09191234567', None, 'seven@example.com', None],
        ['00008', '0008', 'PL', 2, '09191234567', None, 'loan@example.com', None],
        ['00009', '0009', 'MC', 2, '#VALUE!', '09191234567', '=1+1', 'recovered@example.com'],
        [10, 123, 'MC', 2, '09201234567', None, 'masked@example.com', None],
        ['00011', '000011', 'MC', 2, '123', None, 'bad..email@example.com', None],
    ], sheet_name='SBC')
    pairs = [('00001', '0000000000000001'), ('@BEL-1', '@BEL-1'), (' 00001 ', '0000000000000001'),
             ('00099', '00099'), (None, None), ('00002', '0002'), ('00004', '0005'), ('00006', None),
             ('00008', '0008'), ('00009', '0009'), ('00010', '000123'), ('00011', '000011')]
    grid = workbook_data(GRID_HEADERS, [[index, 'Account', None, account, None, contract]
                                       for index, (account, contract) in enumerate(pairs, 1)], sheet_name='Sheet')
    return raw, grid


def selection(name, data, database):
    sheet = max(inspect_contact_source(data, '.xlsx'), key=lambda s: matching_sheet_score(s, database=database))
    return SourceSelection(name, data, sheet)


def csv_rows(file):
    return list(csv.DictReader(StringIO(file.data.decode('cp1252'))))


class B2Tests(unittest.TestCase):
    def test_even_agent_assignment_survives_template_to_autostat_processing(self):
        raw, grid = fixtures()
        result = build_b2_exports(selection('database.xlsx', raw, True), selection('grid.xlsx', grid, False))
        expected = {row['primary_key']: row['Terms'] for file in result.bundle.files for row in csv_rows(file)}
        sources = []
        for file in result.bundle.files:
            sheet = inspect_contact_source(file.data, '.csv')[0]
            source = SourceSelection(file.filename, file.data, sheet)
            sources.append(B2ActivitySource(source, *b2_metadata_columns(sheet.headers)))
        activity = build_b2_autostat_exports(sources, timestamp=None)
        for row in csv_rows(activity.files[0]):
            self.assertEqual(row['Collector'], expected[row['Account No.']])
            self.assertTrue(row['REMARKS'].endswith(' | ' + expected[row['Account No.']]))
        self.assertEqual(result.bundle.summary[0]['AMRANCELI accounts'], 2)
        self.assertEqual(result.bundle.summary[0]['KPALIMA accounts'], 1)

    def test_email_01_and_mobile_1_take_priority_over_earlier_fields_with_valid_fallbacks(self):
        headers = ['Account Key', 'Contract Number', 'Product Category Code', 'Bucket Code',
                   'Phone (01)', 'Email (02)', 'Mobile 1', 'EMAIL (01)']
        raw = workbook_data(headers, [
            ['0001', '000001', 'MC', 2, '09171234567', 'old@example.com', '+63 918 123 4567', 'preferred@example.com'],
            ['BEL-2', 'BEL-002', 'BEL', 2, '09191234567', 'fallback@example.com', None, None],
            ['0003', '000003', 'MC', 2, '09201234567', 'other@example.com', '123', 'bad..email@example.com'],
        ], sheet_name='SBC')
        grid = workbook_data(GRID_HEADERS, [
            [1, 'Account', None, '0001', None, '000001'],
            [2, 'Account', None, 'BEL-2', None, 'BEL-002'],
            [3, 'Account', None, '0003', None, '000003'],
        ], sheet_name='Sheet')
        database = selection('database.xlsx', raw, True)
        result = build_b2_exports(database, selection('grid.xlsx', grid, False))
        self.assertEqual([file.count for file in result.bundle.files], [2, 1])
        rows = {row['cardnumber']: row for file in result.bundle.files for row in csv_rows(file)}
        self.assertEqual((rows['000001']['email'], rows['000001']['phone']),
                         ('preferred@example.com', '9181234567'))
        self.assertEqual((rows['BEL-002']['email'], rows['BEL-002']['phone']),
                         ('fallback@example.com', '9191234567'))
        self.assertEqual((rows['000003']['email'], rows['000003']['phone']),
                         ('other@example.com', '9201234567'))
        self.assertEqual([database.sheet.headers[index] for index in database.sheet.email_columns], ['Email (02)', 'EMAIL (01)'])
        self.assertEqual([database.sheet.headers[index] for index in database.sheet.phone_columns], ['Phone (01)', 'Mobile 1'])

    def test_datagrid_detection_uses_export_headers_without_routing_campaign_sources(self):
        self.assertTrue(is_datagrid_sheet(GRID_HEADERS))
        self.assertTrue(is_datagrid_sheet([' ACCOUNT NO. ', 'Financing/Card No.', 'Status Code', 'Email']))
        self.assertTrue(is_datagrid_sheet(['Account No.', 'Financing/Card No.', 'Client Name']))
        for headers in (RAW_HEADERS, ['primary_key', 'cardnumber', 'email', 'PHONE'],
                        ['Account no.', 'financing/Card no.', 'Name', 'debtor_email', 'phone_1'],
                        ['No', 'Name', 'New I.C.', 'Email']):
            with self.subTest(headers=headers):
                self.assertFalse(is_datagrid_sheet(headers))

    def test_matching_uses_only_grid_accounts_validates_keys_buckets_and_deduplicates(self):
        raw, grid = fixtures()
        result = build_b2_exports(selection('database.xlsx', raw, True), selection('grid.xlsx', grid, False))
        self.assertEqual(result.stats['Database rows'], 12)
        self.assertEqual(result.stats['DataGrid rows'], 12)
        self.assertEqual(result.stats['Matched rows'], 8)
        self.assertEqual(result.stats['Unique B2 accounts'], 5)
        self.assertEqual(result.stats['Unmatched rows'], 1)
        self.assertEqual(result.stats['Missing account numbers'], 1)
        self.assertEqual(result.stats['Conflicting or ambiguous keys'], 2)
        self.assertEqual(result.stats['Other or missing bucket'], 1)
        self.assertEqual(result.stats['Unsupported product'], 1)
        self.assertEqual(result.stats['Duplicate account rows removed'], 1)
        cards, bel = result.bundle.files
        self.assertEqual([file.filename for file in result.bundle.files], ['NEGATIVE B2 - CARDS.csv', 'NEGATIVE B2 - BEL.csv'])
        self.assertEqual([file.count for file in result.bundle.files], [3, 1])
        self.assertEqual({row['cardnumber'] for row in csv_rows(cards)}, {'0000000000000001', '0009', '000123'})
        self.assertEqual({row['primary_key'] for row in csv_rows(cards)}, {'00001', '00009', '00010'})
        self.assertEqual({row['cardnumber'] for row in csv_rows(bel)}, {'@BEL-1'})
        self.assertEqual(result.bundle.summary[0]['No valid contact'], 1)
        self.assertEqual(result.bundle.summary[0]['Accounts exported'], 3)
        for file in result.bundle.files:
            rows = csv_rows(file)
            self.assertEqual(tuple(rows[0]), B2_TEMPLATE_HEADERS)
            identities = [row['cardnumber'] for row in rows]
            self.assertEqual(len(identities), len(set(identities)))
            for index, row in enumerate(rows):
                self.assertEqual((row['Monthly'], row['Terms']),
                                 [('Almera Ranceli', 'AMRANCELI'), ('Kyarill Palima', 'KPALIMA')][index % 2])
                self.assertTrue(row['email'] or row['phone'])
                self.assertTrue(not row['phone'] or re.fullmatch(r'9[0-9]{9}', row['phone']))
        self.assertIn('recovered@example.com', {row['email'] for row in csv_rows(cards)})
        self.assertIn('9191234567', {row['phone'] for row in csv_rows(cards)})
        self.assertNotIn('database-only@example.com', {row['email'] for row in csv_rows(cards)})
        with ZipFile(BytesIO(result.bundle.zip_data)) as archive:
            self.assertEqual(archive.namelist(), [file.filename for file in result.bundle.files])
            for file in result.bundle.files:
                self.assertEqual(archive.read(file.filename), file.data)

    def test_exact_blasting_headers_metadata_and_one_row_with_both_contacts(self):
        self.assertEqual(B2_TEMPLATE_HEADERS, (
            'email', 'primary_key', 'indexoneuid', 'cardnumber', 'FullName', 'phone', 'OB',
            'Monthly', 'Terms', 'doclass', 'Level', 'Product', 'EndoDate', 'AREA', 'ChristmasAmnesty', 'NewEndo',
        ))
        raw = workbook_data(RAW_HEADERS + ['Name', 'Total Outstanding', 'Assignment Date', 'Area', 'Monthly', 'Terms'], [
            ['0001', '000001', 'MC', 2, '09171234567', None, 'first@example.com', None,
             'Database name', 1234.5, datetime(2026, 10, 2), 'Mindanao', 'Sample collector', 'SAMPLE'],
            ['0002', '000002', 'MC', 2, 'invalid', None, 'email-only@example.com', None,
             None, 0, None, None, None, None],
            ['@BEL-3', '@BEL-3', 'BEL', 2, '+63 918 123 4567', None, 'invalid', None,
             'BEL name', 321, datetime(2026, 10, 3), 'R3', None, None],
        ], sheet_name='SBC')
        grid = workbook_data(GRID_HEADERS + ['Old I.C.', 'Amount Referred', 'EndoDate', 'Area'], [
            [1, 'Old name', 'WRONG-NEW-IC', '0001', 'ignore-grid@example.com', '000001', 'SBCC03-001', 999, None, None],
            [2, 'Grid fallback', 'WRONG-NEW-IC', '0002', None, '000002', None, 999, datetime(2026, 10, 5), 'R3'],
            [3, 'Old BEL name', None, '@BEL-3', None, '@BEL-3', 'SBCC03-003', 999, None, None],
        ], sheet_name='Sheet')
        result = build_b2_exports(selection('database.xlsx', raw, True), selection('grid.xlsx', grid, False))
        cards, bel = [csv_rows(file) for file in result.bundle.files]
        self.assertEqual([len(cards), len(bel)], [2, 1])
        self.assertEqual(cards[0]['primary_key'], '0001')
        self.assertEqual(cards[0]['cardnumber'], '000001')
        self.assertEqual(cards[0]['indexoneuid'], 'SBCC03-001')
        self.assertEqual(cards[0]['FullName'], 'Database name')
        self.assertEqual(cards[0]['OB'], '1234.5')
        self.assertEqual(cards[0]['EndoDate'], '2026-10-02 00:00:00')
        self.assertEqual(cards[0]['doclass'], 'SBC CARDS B2')
        self.assertEqual(cards[0]['AREA'], 'Mindanao')
        self.assertEqual(cards[0]['email'], 'first@example.com')
        self.assertEqual(cards[0]['phone'], '9171234567')
        self.assertEqual(cards[1]['FullName'], 'Grid fallback')
        self.assertEqual(cards[1]['OB'], '0')
        self.assertEqual(cards[1]['indexoneuid'], '')
        self.assertEqual(cards[1]['EndoDate'], '2026-10-05 00:00:00')
        self.assertEqual(cards[1]['email'], 'email-only@example.com')
        self.assertEqual(cards[1]['phone'], '')
        self.assertEqual(bel[0]['email'], '')
        self.assertEqual(bel[0]['phone'], '9181234567')
        self.assertEqual(bel[0]['doclass'], 'SBC BEL B2')
        self.assertEqual(bel[0]['Product'], 'BEL')
        for row in cards + bel:
            for column in ('Level', 'ChristmasAmnesty', 'NewEndo'):
                self.assertEqual(row[column], '')
            self.assertNotIn('Time', row)
            self.assertNotIn('Remark', row)
        self.assertEqual([(row['Monthly'], row['Terms']) for row in cards],
                         [('Almera Ranceli', 'AMRANCELI'), ('Kyarill Palima', 'KPALIMA')])
        self.assertEqual((bel[0]['Monthly'], bel[0]['Terms']), ('Almera Ranceli', 'AMRANCELI'))

    def test_incorrect_excel_dimensions_do_not_hide_headers_or_rows_in_either_reader(self):
        raw, grid = fixtures()
        workbook = load_workbook(BytesIO(grid), read_only=True)
        self.assertEqual((workbook.active.max_row, workbook.active.max_column), (1, 3))
        workbook.close()
        info = inspect_workbook(grid)[0]
        self.assertEqual(len(info.headers), len(GRID_HEADERS))
        self.assertIn('email', info.available_fields)
        self.assertIn('cardnumber', info.available_fields)
        mapped = read_source(grid, info.name, info.header_row, allow_missing=True)
        self.assertEqual(mapped.no_contact_rows_skipped, 12)
        source = selection('database.xlsx', raw, True)
        self.assertEqual(len(source.sheet.headers), len(RAW_HEADERS))
        self.assertTrue(source.sheet.ready)

    def test_missing_database_key_headers_produce_an_actionable_error(self):
        raw = workbook_data(['Account Key', 'Contract Number', 'Email (01)'], [['0001', '001', 'first@example.com']], sheet_name='SBC')
        _, grid = fixtures()
        with self.assertRaisesRegex(AutomationError, 'missing Bucket Code'):
            build_b2_exports(selection('database.xlsx', raw, True), selection('grid.xlsx', grid, False))

    def test_datagrid_used_as_database_explains_which_workbook_is_needed(self):
        raw, grid = fixtures()
        with self.assertRaisesRegex(AutomationError, 'this is the DataGrid account list.*full SBC database'):
            build_b2_exports(selection('NEGATIVE B2 - BP.xlsx', grid, True), selection('database.xlsx', raw, False))

    def test_full_database_cannot_replace_the_datagrid_account_population(self):
        raw, _ = fixtures()
        with self.assertRaisesRegex(AutomationError, 'this is the SBC database.*DataGrid account list'):
            build_b2_exports(selection('database.xlsx', raw, True), selection('another database.xlsx', raw, False))


class Upload(BytesIO):
    def __init__(self, name, data):
        super().__init__(data)
        self.name = name


class B2AppTests(unittest.TestCase):
    def test_reversed_uploads_use_the_correct_roles_and_keep_the_same_exports(self):
        raw, grid = fixtures()
        expected = build_b2_exports(selection('database.xlsx', raw, True), selection('grid.xlsx', grid, False))
        uploads = {'b2_database_upload': Upload('NEGATIVE B2 - BP.xlsx', grid)}
        with patch('streamlit.file_uploader', side_effect=lambda label, **kw: uploads.get(kw['key'])):
            app = AppTest.from_file(str(ROOT / 'pages' / 'NEGATIVE_B2.py')).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertTrue(any('DataGrid account list' in item.value and 'empty box' in item.value for item in app.info))
            self.assertFalse(any(widget.key == 'b2_prepare' for widget in app.button))
            uploads['b2_grid_upload'] = Upload('sbc cards-bel database as of October 29, 2026.xlsx', raw)
            app.run(timeout=20)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(next(widget.value for widget in app.selectbox if widget.label == 'SBC database sheet'), 'SBC')
            self.assertEqual(next(widget.value for widget in app.selectbox if widget.label == 'DataGrid accounts sheet'), 'Sheet')
            self.assertTrue(any('Detected the SBC database' in item.value for item in app.info))
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            actual = app.session_state['b2_result']['matched']
            self.assertEqual(actual.stats, expected.stats)
            self.assertEqual([file.data for file in actual.bundle.files], [file.data for file in expected.bundle.files])
            # Replacing the daily database must automatically replace the downloads.
            uploads['b2_grid_upload'] = Upload('sbc cards-bel database as of October 29, 2026.xlsx', workbook_data(
                RAW_HEADERS, [['00001', '0000000000000001', 'MC', 2, '09181234567', None, 'updated@example.com', None]],
                sheet_name='SBC',
            ))
            app.run(timeout=20)
            self.assertIn('b2_result', app.session_state)
            self.assertEqual(len(app.get('download_button')), 1)
            self.assertEqual(csv_rows(app.session_state['b2_result']['matched'].bundle.files[0])[0]['email'], 'updated@example.com')

    def test_two_account_lists_clear_old_downloads_and_report_missing_database_before_matching(self):
        raw, grid = fixtures()
        uploads = {'b2_database_upload': Upload('database.xlsx', raw), 'b2_grid_upload': Upload('NEGATIVE B2 - BP.xlsx', grid)}
        with patch('streamlit.file_uploader', side_effect=lambda label, **kw: uploads.get(kw['key'])):
            app = AppTest.from_file(str(ROOT / 'pages' / 'NEGATIVE_B2.py')).run(timeout=20)
            self.assertIn('b2_result', app.session_state)
            uploads['b2_database_upload'] = Upload('NEGATIVE B2 - CARDS AND BEL.xlsx', grid)
            app.run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 1)
            self.assertIn('both DataGrid account lists', app.error[0].value)
            self.assertIn('NEGATIVE B2 - CARDS AND BEL.xlsx', app.error[0].value)
            self.assertIn('NEGATIVE B2 - BP.xlsx', app.error[0].value)
            self.assertIn('full SBC database', app.error[0].value)
            self.assertNotIn('b2_result', app.session_state)
            self.assertEqual(len(app.get('download_button')), 0)
            self.assertFalse(any(widget.key == 'b2_prepare' for widget in app.button))

    def test_database_uploaded_alone_in_second_box_requests_the_account_list(self):
        raw, grid = fixtures()
        uploads = {'b2_grid_upload': Upload('database.xlsx', raw)}
        with patch('streamlit.file_uploader', side_effect=lambda label, **kw: uploads.get(kw['key'])):
            app = AppTest.from_file(str(ROOT / 'pages' / 'NEGATIVE_B2.py')).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertTrue(any('Upload the DataGrid account list in the empty box' in item.value for item in app.info))
            uploads['b2_database_upload'] = Upload('NEGATIVE B2 - CARDS AND BEL.xlsx', grid)
            app.run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual([file.count for file in app.session_state['b2_result']['matched'].bundle.files], [3, 1])

    def test_selected_lite_sheet_is_validated_and_full_sheet_recovers_matching(self):
        workbook = Workbook()
        workbook.active.title = 'SBC'
        workbook.active.append(RAW_HEADERS)
        workbook.active.append(['00001', '0000000000000001', 'MC', 2, '09171234567', None, 'first@example.com', None])
        lite = workbook.create_sheet('SBC Lite')
        lite.append(['Account Key', 'Contract Number', 'Product Category Code'])
        lite.append(['00001', '0000000000000001', 'MC'])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        _, grid = fixtures()
        uploads = {'b2_database_upload': Upload('database.xlsx', output.getvalue()), 'b2_grid_upload': Upload('list.xlsx', grid)}
        with patch('streamlit.file_uploader', side_effect=lambda label, **kw: uploads.get(kw['key'])):
            app = AppTest.from_file(str(ROOT / 'pages' / 'NEGATIVE_B2.py')).run(timeout=20)
            self.assertIn('b2_result', app.session_state)
            database_sheet = next(widget for widget in app.selectbox if widget.label == 'SBC database sheet')
            database_sheet.set_value('SBC Lite').run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 1)
            self.assertIn('missing Bucket Code', app.error[0].value)
            self.assertNotIn('b2_result', app.session_state)
            self.assertFalse(any(widget.key == 'b2_prepare' for widget in app.button))
            database_sheet = next(widget for widget in app.selectbox if widget.label == 'SBC database sheet')
            database_sheet.set_value('SBC').run(timeout=20)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(app.session_state['b2_result']['matched'].bundle.files[0].count, 1)

    def test_dedicated_b2_page_refreshes_contacts_from_daily_database(self):
        raw, grid = fixtures()
        uploads = {'b2_grid_upload': Upload('Daily account list.xlsx', grid)}
        with patch('streamlit.file_uploader', side_effect=lambda label, **kw: uploads.get(kw['key'])):
            app = AppTest.from_file(str(ROOT / 'app.py')).run(timeout=20)
            app.switch_page('pages/NEGATIVE_B2.py').run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any('latest full SBC database' in message.value for message in app.info))
            self.assertNotIn('b2_result', app.session_state)
            self.assertFalse(any(widget.key == 'b2_prepare' for widget in app.button))

            uploads['b2_database_upload'] = Upload('daily database.xlsx', raw)
            app.run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertNotIn('result', app.session_state)
            files = app.session_state['b2_result']['matched'].bundle.files
            self.assertEqual([file.filename for file in files], ['NEGATIVE B2 - CARDS.csv', 'NEGATIVE B2 - BEL.csv'])
            self.assertIn('first@example.com', {row['email'] for row in csv_rows(files[0])})
            self.assertFalse(any(widget.key == 'b2_remark_by' for widget in app.text_input))
            downloads = [widget.proto.label for widget in app.get('download_button')]
            self.assertIn('Download NEGATIVE B2 - CARDS.csv', downloads)
            self.assertIn('Download NEGATIVE B2 - BEL.csv', downloads)
            self.assertFalse(any('_filled_' in label for label in downloads))

            updated = workbook_data(RAW_HEADERS, [
                ['00001', '0000000000000001', 'MC', 2, '09221234567', None, 'next-day@example.com', None],
                ['@BEL-1', '@BEL-1', 'BEL', 2, '09181234567', None, 'bel@example.com', None],
            ], sheet_name='SBC')
            uploads['b2_database_upload'] = Upload('daily database.xlsx', updated)
            app.run(timeout=20)
            self.assertIn('b2_result', app.session_state)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            files = app.session_state['b2_result']['matched'].bundle.files
            self.assertEqual(len(csv_rows(files[0])), 1)
            self.assertEqual(csv_rows(files[0])[0]['email'], 'next-day@example.com')
            self.assertEqual(csv_rows(files[0])[0]['phone'], '9221234567')
            self.assertEqual([file.filename for file in files], ['NEGATIVE B2 - CARDS.csv', 'NEGATIVE B2 - BEL.csv'])

            uploads.pop('b2_database_upload')
            app.run(timeout=20)
            self.assertNotIn('b2_result', app.session_state)

    def test_matching_workflow_downloads_two_files_saves_and_invalidates(self):
        raw, grid = fixtures()
        database, accounts = Upload('database.xlsx', raw), Upload('DataGrid.xlsx', grid)
        def uploader(label, **kwargs):
            return database if kwargs.get('key') == 'b2_database_upload' else accounts if kwargs.get('key') == 'b2_grid_upload' else None
        with patch('streamlit.file_uploader', side_effect=uploader), TemporaryDirectory(prefix='b2-matching-test-') as directory:
            app = AppTest.from_file(str(ROOT / 'app.py')).run(timeout=20)
            app.switch_page('pages/NEGATIVE_B2.py').run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.title[0].value, 'NEGATIVE B2')
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state['b2_result']['matched']
            self.assertEqual(len(result.bundle.files), 2)
            self.assertEqual(len(app.get('download_button')), 3)
            app.text_input(key='b2_output_folder').set_value(directory).run(timeout=20)
            app.button(key='b2_save').click().run(timeout=20)
            self.assertEqual(len(app.error), 0)
            saved = Path(app.session_state['b2_result']['saved_folder'])
            self.assertEqual({path.name for path in saved.iterdir()}, {file.filename for file in result.bundle.files})
            for file in result.bundle.files:
                self.assertEqual((saved / file.filename).read_bytes(), file.data)
            app.selectbox(key='b2_csv_type').set_value('CSV UTF-8 (Comma delimited)').run(timeout=20)
            self.assertTrue(app.session_state['b2_result']['matched'].bundle.files[0].data.startswith(b'\xef\xbb\xbf'))
            self.assertNotIn('saved_folder', app.session_state['b2_result'])


if __name__ == '__main__':
    unittest.main()
