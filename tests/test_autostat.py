from __future__ import annotations

import csv
import unittest
from dataclasses import replace
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from automation import AutomationError, DEFAULT_TEMPLATE_HEADERS, create_default_csv_template
from autostat import (
    EMAIL_HEADERS, REMARK_PREFIX, VIBER_HEADERS, SourceSelection,
    autostat_source_excluded, build_autostat_excel, build_autostat_exports, build_bel_exports, extract_contacts, has_bel_primary_keys,
    inspect_contact_source, split_bel_contacts, unique_account_contacts, validated_phone,
)
from export_config import AUTOSTAT_BEL_FILENAMES
from email_validation import validated_email
from folder_output import MANILA_TIMEZONE
from exports import build_export_csv, build_export_excel, build_geo_exports, build_negative_exports
from geo_autostat import mapped_contact_source


ROOT = Path(__file__).resolve().parents[1]
NOW = datetime(2026, 10, 28, 15, 4, 5, tzinfo=MANILA_TIMEZONE)


class Upload(BytesIO):
    def __init__(self, name, data):
        super().__init__(data)
        self.name = name


def workbook_bytes():
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Account totals"])
    sheet = workbook.create_sheet("Contacts")
    sheet.append(["Exported contacts"])
    sheet.append(["Account no.", "debtor_email", "phone_1", "Phone 2", "Other data"])
    for row in (
        ["00001", " Ana.Cruz+tag@EXAMPLE.COM ", "+63 (917) 123-4567", "09171234567", None],
        ["00002", "bad..email@example.com;other@example.com", "08171234567", "09181234567", None],
        ["00003", "bad email@example.com", "123", None, None],
        [None, "orphan@example.com", "09191234567", None, None],
        [None, None, None, None, None],
        ["00004", "repeat@example.com;repeat@example.com", 9171234567.0, None, None],
    ):
        sheet.append(row)
    next_sheet = workbook.create_sheet("Phone only")
    next_sheet.append(["primary_key", "phone"])
    next_sheet.append([123, 639201234567])
    next_sheet["A2"].number_format = "000000"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def exported_rows(file):
    return list(csv.DictReader(StringIO(file.data.decode("utf-8-sig" if file.data.startswith(b"\xef\xbb\xbf") else "cp1252"))))


class PhoneTests(unittest.TestCase):
    def test_rightmost_ten_digits_and_numeric_inputs(self):
        for value in ("09171234567", "+63 (917) 123-4567", "917 123 4567", "917.123.4567",
                      "000639171234567", 9171234567, 639171234567.0, "9171234567.0", "9.171234567E9"):
            with self.subTest(value=value):
                self.assertEqual(validated_phone(value), "9171234567")

    def test_invalid_values_are_not_repaired_into_mobile_numbers(self):
        for value in (None, "", "  ", True, "8171234567", "0917123456", "123", "call 09171234567",
                      "09171234567 ext 9", -9171234567, 9171234567.5, float("nan"), float("inf"), "9e999"):
            with self.subTest(value=value):
                self.assertIsNone(validated_phone(value))


class AutostatTests(unittest.TestCase):
    def test_whole_number_accounts_lose_only_zero_decimals_and_remain_exact_text_in_all_outputs(self):
        cases = [
            ("5156041234567897.0", "5156041234567897"),
            ("000123.000", "000123"),
            ("123456789012345678901234.00", "123456789012345678901234"),
            ("5146000000000001", "5146000000000001"),
            ("BEL-123.0", "BEL-123.0"),
            ("123.45", "123.45"),
            ("001.0", "001"),
            ("+1234567890123456.0", "+1234567890123456"),
            ("  002.00  ", "002"),
            ("000123", "000123"),  # The cleaned duplicate keeps the first contact.
        ]
        rows = [["cardnumber", "email", "phone"], *[
            [account, f"contact{index}@example.com", "09171234567"]
            for index, (account, _) in enumerate(cases)
        ]]
        csv_output = StringIO(newline="")
        csv.writer(csv_output).writerows(rows)
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        excel_output = BytesIO()
        workbook.save(excel_output)
        workbook.close()
        expected = [account for _, account in cases[:-1]]
        for extension, data in ((".csv", csv_output.getvalue().encode()), (".xlsx", excel_output.getvalue())):
            with self.subTest(source=extension):
                source = SourceSelection("accounts" + extension, data, inspect_contact_source(data, extension)[0])
                bundle = build_autostat_exports([source], timestamp=None)
                self.assertEqual(bundle.summary[0]["Account precision issues"], 0)
                self.assertEqual(bundle.summary[0]["Duplicate Email rows removed"], 1)
                self.assertEqual(bundle.summary[0]["Duplicate Viber rows removed"], 1)
                with ZipFile(BytesIO(bundle.zip_data)) as archive:
                    for file in bundle.files:
                        self.assertEqual([row["Account No."] for row in exported_rows(file)], expected)
                        self.assertEqual([row["Account No."] for row in file.preview], expected)
                        self.assertEqual(archive.read(file.filename), file.data)
                self.assertEqual(exported_rows(bundle.files[0])[1]["EMAIL/ NUMBER"], "contact1@example.com")
                saved = load_workbook(BytesIO(build_autostat_excel(bundle)))
                for sheet in saved:
                    self.assertEqual([cell.value for cell in sheet["B"][1:]], expected)
                    self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in sheet["B"][1:]))
                saved.close()
                self.assertEqual(source.data, data)

    def test_excel_download_preserves_full_identifiers_and_all_rows_as_literal_text(self):
        accounts = [
            "5156041234567897", "001234567890123456", "123456789012345678901234",
            "=2+2", "+1234567890123456", "@BEL-1",
            *[f"000000{number:012d}" for number in range(110)],
        ]
        rows = [["financing/Card no.", "email", "phone"], *[
            [account, f"contact{index}@example.com", "09171234567"]
            for index, account in enumerate(accounts)
        ]]
        csv_output = StringIO(newline="")
        csv.writer(csv_output).writerows(rows)
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        # A literal formula-like identifier must not become a source formula.
        workbook.active["A5"].data_type = "s"
        workbook_output = BytesIO()
        workbook.save(workbook_output)
        workbook.close()
        for extension, source_data in ((".csv", csv_output.getvalue().encode()), (".xlsx", workbook_output.getvalue())):
            for encoding, author in (("cp1252", "Muñoz"), ("utf-8-sig", "李")):
                with self.subTest(extension=extension, encoding=encoding):
                    source = SourceSelection(
                        "contacts" + extension, source_data, inspect_contact_source(source_data, extension)[0],
                    )
                    bundle = build_autostat_exports([source], timestamp=None, remark_by=author, encoding=encoding)
                    self.assertEqual(bundle.summary[0]["Account precision issues"], 0)
                    saved_csvs = [file.data for file in bundle.files]
                    result = load_workbook(BytesIO(build_autostat_excel(bundle)))
                    self.assertEqual(result.sheetnames, ["Email", "Viber"])
                    for sheet, file in zip(result, bundle.files):
                        self.assertEqual([cell.value for cell in sheet[1]], list(file.headers))
                        self.assertEqual([cell.value for cell in sheet["B"][1:]], accounts)
                        self.assertEqual(sheet.max_row, len(accounts) + 1)
                        for row in sheet.iter_rows(min_row=2):
                            self.assertEqual(row[1].data_type, "s")
                            self.assertEqual(row[1].number_format, "@")
                            self.assertEqual(row[4].value, "MBURGOS")
                            self.assertIsNone(row[0].value)
                        self.assertEqual([row["Account No."] for row in exported_rows(file)], accounts)
                    result.close()
                    self.assertEqual([file.data for file in bundle.files], saved_csvs)

    def test_excel_download_keeps_headers_for_an_empty_channel(self):
        data = b"cardnumber,email\n0001,first@example.com\n"
        source = SourceSelection("source.csv", data, inspect_contact_source(data, ".csv")[0])
        bundle = build_autostat_exports([source], timestamp=NOW)
        workbook = load_workbook(BytesIO(build_autostat_excel(bundle)))
        self.assertEqual(workbook["Viber"].max_row, 1)
        self.assertEqual([cell.value for cell in workbook["Viber"][1]], list(VIBER_HEADERS))
        workbook.close()

    def test_abbreviated_scientific_accounts_are_skipped_and_numeric_sources_retain_supplied_digits(self):
        data = (
            b"cardnumber,indexoneuid,email,phone\n"
            b"000001,UID-0,before@example.com,09171234567\n"
            b"5.15604E+15,UID-1,invalid@example.com,09181234567\n"
            b"5156041234567897,UID-2,after@example.com,09191234567\n"
        )
        source = SourceSelection("source.csv", data, inspect_contact_source(data, ".csv")[0])
        contacts = extract_contacts(source)
        self.assertEqual(contacts.source_error_rows, [3])
        bundle = build_autostat_exports([source], timestamp=NOW)
        self.assertEqual(bundle.summary[0]["Source error rows skipped"], 1)
        self.assertEqual(bundle.summary[0]["Account precision issues"], 0)
        for file in bundle.files:
            self.assertEqual([row["Account No."] for row in exported_rows(file)], ["000001", "5156041234567897"])
        workbook = Workbook()
        for row in csv.reader(StringIO(data.decode())):
            workbook.active.append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
        source = SourceSelection("source.xlsx", data, inspect_contact_source(data, ".xlsx")[0])
        self.assertEqual(extract_contacts(source).source_error_rows, [3])
        bundle = build_autostat_exports([source], timestamp=NOW)
        self.assertEqual(bundle.summary[0]["Source error rows skipped"], 1)
        for file in bundle.files:
            self.assertEqual([row["Account No."] for row in exported_rows(file)], ["000001", "5156041234567897"])
        workbook = Workbook()
        workbook.active.append(["cardnumber", "email"])
        workbook.active.append([5156040000000000, "first@example.com"])
        workbook.active.append(["001234567890123456", "second@example.com"])
        workbook.active.append([123, "third@example.com"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        source = SourceSelection("source.xlsx", output.getvalue(), inspect_contact_source(output.getvalue(), ".xlsx")[0])
        bundle = build_autostat_exports([source], timestamp=NOW)
        self.assertEqual(bundle.summary[0]["Account precision issues"], 1)
        self.assertEqual([row["Account No."] for row in exported_rows(bundle.files[0])],
                         ["5156040000000000", "001234567890123456", "123"])

    def test_excel_errors_and_missing_formula_values_skip_each_bad_row_once(self):
        workbook = Workbook()
        workbook.active.append(["cardnumber", "email", "phone", "Unused"])
        for row in (
            ["000001", "before@example.com", "09171234567", None],
            ["#N/A", "bad-account@example.com", "09181234567", None],
            ["000002", "#VALUE!", "09181234567", None],
            ["000003", "bad-phone@example.com", "#DIV/0!", None],
            ["#REF!", "#N/A", "09181234567", None],
            ["=1+1", "formula-account@example.com", "09181234567", None],
            ["000004", "formula-phone@example.com", "=1+1", None],
            ["000009", "unused-error@example.com", "09191234567", "#N/A"],
            ["000002", "after@example.com", "09201234567", None],
        ):
            workbook.active.append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
        source = SourceSelection("errors.xlsx", data, inspect_contact_source(data, ".xlsx")[0])
        self.assertEqual(extract_contacts(source).source_error_rows, [3, 4, 5, 6, 7, 8])
        bundle = build_autostat_exports([source], timestamp=NOW)
        self.assertEqual(bundle.summary[0]["Source error rows skipped"], 6)
        self.assertEqual(bundle.summary[0]["Rows checked"], 3)
        for file in bundle.files:
            self.assertEqual([row["Account No."] for row in exported_rows(file)], ["000001", "000009", "000002"])

    def test_each_single_email_is_validated_once_including_addresses_with_a_pipe(self):
        data = (
            'cardnumber,email\n'
            '0001,first+tag@example.com\n'
            '0002,pipe|tag@example.com\n'
            '0003,bad..email@example.com\n'
        ).encode()
        source = SourceSelection("contacts.csv", data, inspect_contact_source(data, ".csv")[0])
        with patch("autostat.validated_email", wraps=validated_email) as email_checks:
            contacts = extract_contacts(source)
        self.assertEqual(email_checks.call_count, 3)
        self.assertEqual(contacts.emails, [
            ("0001", "first+tag@example.com"), ("0002", "pipe|tag@example.com"),
        ])
        self.assertEqual(contacts.invalid_emails, 1)

    def test_excluded_campaign_files_do_not_enter_counts_exports_or_duplicate_tracking(self):
        data = b"cardnumber,email,phone\n0001,first@example.com,09171234567\n"
        sheet = inspect_contact_source(data, ".csv")[0]
        excluded = [SourceSelection(name, data, sheet) for name in (
            "SBC INSURANCE.csv", "sbc_insurance_2026.csv", "NEGATIVE B2 - BEL.csv",
            "NEGATIVE_B2_CARDS.csv", "negative-b4-bel.csv", "B4.csv",
        )]
        included = SourceSelection("GEO_Viber.csv", data, sheet)
        bundle = build_autostat_exports([*excluded, included], timestamp=None)
        self.assertEqual([file.count for file in bundle.files], [1, 1])
        self.assertEqual([row["File"] for row in bundle.summary], [included.name])
        self.assertEqual(bundle.summary[0]["Duplicate Viber rows removed"], 0)
        self.assertEqual(exported_rows(bundle.files[0])[0]["EMAIL/ NUMBER"], "first@example.com")
        with self.assertRaisesRegex(AutomationError, "All selected files are excluded"):
            build_autostat_exports(excluded, timestamp=None)
        for name in ("NEGATIVE B20.csv", "B40.csv", "club2.csv", "GEO_Email.csv"):
            self.assertFalse(autostat_source_excluded(name), name)

    def test_distinct_source_ids_survive_shared_card_numbers_and_true_repeats_are_removed(self):
        for identity_header in ("indexoneuid", "Old I.C."):
            with self.subTest(identity_header=identity_header):
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(["cardnumber", "primary_key", identity_header, "email", "phone"])
                for values in (
                    [9999000000000000, 1000000000000, "UID-1", "first@example.com", "09171234567"],
                    [9999000000000000, 1000000000000, "UID-2", "second@example.com", "09171234567"],
                    [9999000000000000, 1000000000000, "UID-3", "invalid", "09181234567"],
                    [9999000000000000, 1000000000000, "UID-3", "third@example.com", "09191234567"],
                ):
                    sheet.append(values)
                output = BytesIO()
                workbook.save(output)
                workbook.close()
                data = output.getvalue()
                source = SourceSelection("shared-cards.xlsx", data, inspect_contact_source(data, ".xlsx")[0])
                repeated = (
                    f'cardnumber,{identity_header},email,phone\n'
                    '9999000000000000, uid-1 ,later@example.com,09201234567\n'
                ).encode()
                duplicate = SourceSelection("repeat.csv", repeated, inspect_contact_source(repeated, ".csv")[0])
                bundle = build_autostat_exports([source, duplicate], timestamp=NOW)
                self.assertEqual([file.count for file in bundle.files], [3, 3])
                self.assertEqual([row["EMAIL/ NUMBER"] for row in exported_rows(bundle.files[0])],
                                 ["first@example.com", "second@example.com", "third@example.com"])
                self.assertTrue(all(row["Account No."] == "9999000000000000"
                                    for file in bundle.files for row in exported_rows(file)))
                self.assertEqual(bundle.summary[0]["Duplicate Viber rows removed"], 1)
                self.assertEqual(bundle.summary[1]["Duplicate Email rows removed"], 1)
                self.assertEqual(bundle.summary[1]["Duplicate Viber rows removed"], 1)

    def test_accounts_are_unique_per_channel_across_rows_cells_and_files(self):
        data = (
            'cardnumber,email,phone\n'
            '0001,bad..email@example.com,09171234567\n'
            '0001,first@example.com,09181234567\n'
            ' 0001 ,later@example.com,09191234567\n'
            '0002,second@example.com,09201234567\n'
            '0003,third@example.com;alternate@example.com,09211234567;09221234567\n'
        ).encode()
        next_data = b'cardnumber,email,phone\n0002,repeated@example.com,09231234567\n0004,new@example.com,09241234567\n'
        sources = [SourceSelection(name, content, inspect_contact_source(content, '.csv')[0])
                   for name, content in [('first.csv', data), ('second.csv', next_data)]]
        bundle = build_autostat_exports(sources, timestamp=NOW)
        self.assertEqual([file.filename for file in bundle.files], ['AUTOSTAT_Email.csv', 'AUTOSTAT_Viber.csv'])
        self.assertEqual([file.count for file in bundle.files], [4, 4])
        email_rows = exported_rows(bundle.files[0])
        self.assertEqual([row['Account No.'] for row in email_rows], ['0001', '0002', '0003', '0004'])
        self.assertEqual([row['EMAIL/ NUMBER'] for row in email_rows], ['first@example.com', 'second@example.com', 'third@example.com', 'new@example.com'])
        self.assertIn('9171234567', exported_rows(bundle.files[1])[0]['Remark'])
        self.assertEqual(bundle.summary[0]['Duplicate Email rows removed'], 2)
        self.assertEqual(bundle.summary[0]['Duplicate Viber rows removed'], 3)
        self.assertEqual(bundle.summary[1]['Duplicate Email rows removed'], 1)
        self.assertEqual(bundle.summary[1]['Duplicate Viber rows removed'], 1)
        for file in bundle.files:
            accounts = [row['Account No.'] for row in exported_rows(file)]
            self.assertEqual(len(accounts), len(set(accounts)))
        with ZipFile(BytesIO(bundle.zip_data)) as archive:
            self.assertEqual(archive.namelist(), ['AUTOSTAT_Email.csv', 'AUTOSTAT_Viber.csv'])
            for file in bundle.files:
                self.assertEqual(archive.read(file.filename), file.data)

    def test_later_upload_supplies_a_missing_channel_without_dropping_existing_contacts(self):
        first = (
            'cardnumber,email,phone\n'
            'AbC-1,first@example.com,invalid\n'
            '0002,bad..email@example.com,09181234567\n'
            '0003,third@example.com,09191234567\n'
        ).encode()
        second = (
            'cardnumber,email,phone\n'
            ' abc-1 ,later@example.com,09171234567\n'
            '0002,second@example.com,09201234567\n'
            '0004,fourth@example.com,09211234567\n'
        ).encode()
        sources = [SourceSelection('same.csv', data, inspect_contact_source(data, '.csv')[0])
                   for data in (first, second)]
        bundle = build_autostat_exports(sources, timestamp=NOW)
        self.assertEqual([file.count for file in bundle.files], [4, 4])
        emails, phones = [exported_rows(file) for file in bundle.files]
        self.assertEqual({row['Account No.'].strip().casefold(): row['EMAIL/ NUMBER'] for row in emails}, {
            'abc-1': 'first@example.com', '0002': 'second@example.com',
            '0003': 'third@example.com', '0004': 'fourth@example.com',
        })
        phone_remarks = {row['Account No.'].strip().casefold(): row['Remark'] for row in phones}
        self.assertIn('9171234567', phone_remarks['abc-1'])
        self.assertIn('9181234567', phone_remarks['0002'])
        self.assertEqual([row['Account No.'] for row in emails], ['AbC-1', '0003', '0002', '0004'])
        self.assertEqual([row['Account No.'] for row in phones], ['0002', '0003', 'abc-1', '0004'])
        self.assertEqual(len(bundle.summary), 2)
        self.assertEqual(bundle.summary[1]['Duplicate Email rows removed'], 1)
        self.assertEqual(bundle.summary[1]['Duplicate Viber rows removed'], 1)

    def test_compiled_outputs_keep_an_empty_channel_with_its_own_template_headers(self):
        for contact_header, contacts, empty_index, empty_headers in (
            ('email', ('first@example.com', 'second@example.com'), 1, VIBER_HEADERS),
            ('phone', ('09171234567', '09181234567'), 0, EMAIL_HEADERS),
        ):
            with self.subTest(channel=contact_header):
                sources = []
                for number, contact in enumerate(contacts, 1):
                    data = f'cardnumber,{contact_header}\n000{number},{contact}\n'.encode()
                    sources.append(SourceSelection(f'{number}.csv', data, inspect_contact_source(data, '.csv')[0]))
                bundle = build_autostat_exports(sources, timestamp=NOW)
                self.assertEqual([file.filename for file in bundle.files], ['AUTOSTAT_Email.csv', 'AUTOSTAT_Viber.csv'])
                empty = bundle.files[empty_index]
                self.assertEqual(empty.count, 0)
                self.assertEqual(empty.preview, [])
                self.assertEqual(empty.headers, empty_headers)
                self.assertEqual(empty.data.decode('cp1252').splitlines(), [','.join(empty_headers)])
                self.assertEqual(bundle.files[1 - empty_index].count, 2)

    def test_account_matching_preserves_leading_zeros_and_does_not_mutate_contacts(self):
        data = b'cardnumber,email,phone\nAbC-1,first@example.com,09171234567\nabc-1,later@example.com,09181234567\n0001,zero@example.com,09191234567\n1,short@example.com,09201234567\n'
        source = SourceSelection('source.csv', data, inspect_contact_source(data, '.csv')[0])
        contacts = extract_contacts(source)
        original_emails, original_phones = list(contacts.emails), list(contacts.phones)
        result = unique_account_contacts(contacts)
        self.assertEqual([account for account, _ in result.emails], ['AbC-1', '0001', '1'])
        self.assertEqual(result.duplicate_email_rows, 1)
        self.assertEqual(result.duplicate_phone_rows, 1)
        self.assertEqual(contacts.emails, original_emails)
        self.assertEqual(contacts.phones, original_phones)

    def test_two_named_group_logs_keep_both_channel_templates_and_contacts(self):
        data = (
            'primary_key,cardnumber,email,phone\n'
            '@BELLTLSUL01R00001,00001,first@example.com,09171234567\n'
            '1234567890123456,00002,other@example.com,09181234567\n'
            '@bellhusur01r00002,00003,bad..email@example.com,09191234567\n'
            '1234567890123457,00004,bad..email@example.com,08171234567\n'
        ).encode()
        source = SourceSelection('mixed.csv', data, inspect_contact_source(data, '.csv')[0])
        bundle = build_bel_exports([source], timestamp=NOW)
        self.assertEqual([file.filename for file in bundle.files], ['NEGATIVE B4 - BEL.csv', 'NEGATIVE B4 - CARDS.csv'])
        bel, cards = [exported_rows(file) for file in bundle.files]
        self.assertEqual([row['Account No.'] for row in bel], ['00001', '00001', '00003'])
        self.assertEqual([row['EMAIL/ NUMBER'] for row in bel], ['first@example.com', '9171234567', '9191234567'])
        self.assertEqual([row['EMAIL/ NUMBER'] for row in cards], ['other@example.com', '9181234567'])
        for file in bundle.files:
            self.assertEqual(file.headers, EMAIL_HEADERS)
            self.assertEqual(file.count, len(exported_rows(file)))
            for row in exported_rows(file):
                email = '@' in row['EMAIL/ NUMBER']
                ending = 'FOLLOW UP EMAIL SENT' if email else 'VIBER SENT - BLASTING'
                self.assertEqual(row['Remark'], f"{REMARK_PREFIX} - {row['EMAIL/ NUMBER']} - {ending}")
                self.assertEqual(row['Status'], ('EMAIL' if email else 'SMS') + ' BLAST SENT - WAITING FOR REPLY')
                self.assertEqual(row['Time'], '2026-10-28 03:04:05 PM')
                self.assertEqual(row['Remark By'], 'MBURGOS')
                self.assertTrue(all(row[header] == '' for header in EMAIL_HEADERS[6:]))
        self.assertEqual(sum(row['Rows checked'] for row in bundle.summary), 4)
        with ZipFile(BytesIO(bundle.zip_data)) as archive:
            self.assertEqual(archive.namelist(), [file.filename for file in bundle.files])
            for file in bundle.files:
                self.assertEqual(archive.read(file.filename), file.data)

    def test_group_outputs_combine_sources_keep_empty_headers_and_validate_filenames(self):
        data = b'primary_key,email\n@BEL-1,first@example.com\n'
        source = SourceSelection('mixed.csv', data, inspect_contact_source(data, '.csv')[0])
        bundle = build_bel_exports([source, source], timestamp=NOW)
        self.assertEqual([file.count for file in bundle.files], [1, 0])
        self.assertEqual(sum(row['Duplicate Email rows removed'] for row in bundle.summary), 1)
        self.assertEqual(exported_rows(bundle.files[1]), [])
        self.assertEqual(bundle.files[1].data.decode('cp1252').splitlines()[0], ','.join(EMAIL_HEADERS))
        with patch.dict(AUTOSTAT_BEL_FILENAMES, {'Other': AUTOSTAT_BEL_FILENAMES['BEL']}):
            with self.assertRaisesRegex(AutomationError, 'unique'):
                build_bel_exports([source], timestamp=NOW)
        with patch.dict(AUTOSTAT_BEL_FILENAMES, {'BEL': '../outside.csv'}):
            with self.assertRaisesRegex(AutomationError, 'Invalid filename'):
                build_bel_exports([source], timestamp=NOW)

    def test_bel_detection_scans_beyond_first_fifty_rows_in_csv_and_workbook(self):
        headers = ['primary_key', 'email']
        rows = [[str(index), 'first@example.com'] for index in range(60)]
        rows.append(['@BELLTLSUL01R00001', 'last@example.com'])
        workbook = Workbook()
        workbook.active.append(headers)
        for row in rows:
            workbook.active.append(row)
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        csv_data = (','.join(headers) + '\n' + '\n'.join(','.join(row) for row in rows)).encode()
        for name, data in (('late.xlsx', stream.getvalue()), ('late.csv', csv_data)):
            with self.subTest(name=name):
                source = SourceSelection(name, data, inspect_contact_source(data, Path(name).suffix)[0])
                self.assertTrue(has_bel_primary_keys(source))

    def test_bel_grouping_uses_primary_key_while_accounts_use_cardnumber(self):
        data = (
            'primary_key,cardnumber,email,phone\n'
            '@BELLTLSUL01R00001,00001,first@example.com,09171234567\n'
            '1234567890123456,00002,other@example.com,09181234567\n'
            '  @bellhusur02r00002  ,00003,second@example.com,09191234567\n'
            '1234567890123457,00004,bad..email@example.com,08171234567\n'
        ).encode()
        source = SourceSelection('mixed.csv', data, inspect_contact_source(data, '.csv')[0])
        groups = split_bel_contacts(source)
        self.assertEqual(groups['BEL'].emails, [('00001', 'first@example.com'), ('00003', 'second@example.com')])
        self.assertEqual(groups['Other'].phones, [('00002', '9181234567')])
        self.assertEqual(groups['Other'].no_contacts, 1)
        self.assertEqual(sum(group.checked_rows for group in groups.values()), 4)
        self.assertEqual(groups['Other'].invalid_emails, 1)
        self.assertEqual(groups['Other'].invalid_phones, 1)

    def test_bel_grouping_requires_one_unambiguous_primary_key_header(self):
        for headers in ('cardnumber,email', 'primary_key,primary_key,email'):
            with self.subTest(headers=headers):
                data = (headers + '\n0001,first@example.com\n').encode()
                source = SourceSelection('source.csv', data, inspect_contact_source(data, '.csv')[0])
                with self.assertRaisesRegex(AutomationError, 'one primary_key column'):
                    split_bel_contacts(source)

    def test_cardnumber_is_used_for_both_logs_ahead_of_other_account_headers(self):
        for header in ("cardnumber", "Card Number", " CARDNUMBER "):
            with self.subTest(header=header):
                data = (
                    f"primary_key,Account No.,email,{header},phone\n"
                    "INTERNAL-1,OTHER-1,first@example.com,000123456789012345,09171234567\n"
                    "INTERNAL-2,OTHER-2,second@example.com,000000000000000002,09181234567\n"
                    "INTERNAL-3,OTHER-3,missing@example.com,,09191234567\n"
                ).encode()
                sheet = inspect_contact_source(data, ".csv")[0]
                bundle = build_autostat_exports([SourceSelection("contacts.csv", data, sheet)], timestamp=NOW)
                self.assertEqual(bundle.summary[0]["Missing account"], 1)
                for file in bundle.files:
                    rows = exported_rows(file)
                    self.assertEqual([row["Account No."] for row in rows], ["000123456789012345", "000000000000000002"])
                    self.assertEqual([row["Account No."] for row in file.preview], ["000123456789012345", "000000000000000002"])
                    for row, contact in zip(rows, ("first@example.com", "second@example.com") if "_Email" in file.filename else ("9171234567", "9181234567")):
                        self.assertIn(contact, row["Remark"])

    def test_workbook_with_cardnumber_and_contacts_is_detected_and_preserves_zero_mask(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Contact list"])
        sheet.append(["email", "phone", "cardnumber"])
        sheet.append(["first@example.com", "09171234567", 123])
        sheet["C3"].number_format = "000000"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
        detected = inspect_contact_source(data, ".xlsx")[0]
        self.assertTrue(detected.ready)
        self.assertEqual(detected.header_row, 2)
        bundle = build_autostat_exports([SourceSelection("contacts.xlsx", data, detected)], timestamp=NOW)
        for file in bundle.files:
            self.assertEqual(exported_rows(file)[0]["Account No."], "000123")

    def test_detection_multiple_contact_columns_and_validation_counts(self):
        data = workbook_bytes()
        sheets = inspect_contact_source(data, ".xlsx")
        self.assertFalse(sheets[0].ready)
        self.assertEqual(sheets[1].header_row, 2)
        self.assertEqual(sheets[1].phone_columns, (2, 3))
        rows = extract_contacts(SourceSelection("source.xlsx", data, sheets[1]))
        self.assertEqual(rows.emails, [("00001", "Ana.Cruz+tag@example.com"), ("00002", "other@example.com"), ("00004", "repeat@example.com")])
        self.assertEqual(rows.phones, [("00001", "9171234567"), ("00002", "9181234567"), ("00004", "9171234567")])
        self.assertEqual((rows.checked_rows, rows.blank_rows, rows.missing_accounts, rows.no_contacts), (5, 1, 1, 1))
        self.assertEqual((rows.invalid_emails, rows.invalid_phones), (2, 2))

    def test_exact_template_values_numeric_ids_multiple_sheets_and_zip(self):
        data = workbook_bytes()
        sheets = inspect_contact_source(data, ".xlsx")
        bundle = build_autostat_exports([SourceSelection("source.xlsx", data, sheet) for sheet in sheets if sheet.ready], timestamp=NOW)
        self.assertEqual(len(bundle.files), 2)
        email, viber = bundle.files
        self.assertEqual([file.count for file in bundle.files], [3, 4])
        self.assertEqual(len(bundle.summary), 2)
        self.assertEqual(tuple(exported_rows(email)[0]), EMAIL_HEADERS)
        self.assertEqual(tuple(exported_rows(viber)[0]), VIBER_HEADERS)
        first = exported_rows(email)[0]
        self.assertEqual(first["Time"], "2026-10-28 03:04:05 PM")
        self.assertEqual(first["Account No."], "00001")
        self.assertEqual(first["Status"], "EMAIL BLAST SENT - WAITING FOR REPLY")
        self.assertEqual(first["Remark By"], "MBURGOS")
        self.assertEqual(first["Remark"], f"{REMARK_PREFIX} - Ana.Cruz+tag@example.com - FOLLOW UP EMAIL SENT")
        self.assertEqual(first["EMAIL/ NUMBER"], "Ana.Cruz+tag@example.com")
        self.assertTrue(all(first[header] == "" for header in EMAIL_HEADERS[6:]))
        first_viber = exported_rows(viber)[0]
        self.assertEqual(first_viber["Status"], "SMS BLAST SENT - WAITING FOR REPLY")
        self.assertEqual(first_viber["Remark"], f"{REMARK_PREFIX} - 9171234567 - VIBER SENT - BLASTING")
        self.assertEqual(exported_rows(viber)[-1]["Account No."], "000123")
        self.assertEqual(exported_rows(viber)[-1]["Remark"], f"{REMARK_PREFIX} - 9201234567 - VIBER SENT - BLASTING")
        with ZipFile(BytesIO(bundle.zip_data)) as archive:
            self.assertEqual(archive.namelist(), [file.filename for file in bundle.files])
            for file in bundle.files:
                self.assertEqual(archive.read(file.filename), file.data)

    def test_csv_mixed_contacts_manual_mapping_and_duplicate_source_names(self):
        data = 'Report\r\nAccount;EMAIL/ NUMBER\r\n0001;good@example.com|09171234567\r\n'.encode()
        sheet = inspect_contact_source(data, ".csv")[0]
        self.assertEqual(sheet.header_row, 2)
        self.assertEqual(sheet.email_columns, sheet.phone_columns)
        selection = SourceSelection("same.csv", data, sheet)
        contacts = extract_contacts(selection)
        self.assertEqual(contacts.emails, [("0001", "good@example.com")])
        self.assertEqual(contacts.phones, [("0001", "9171234567")])
        self.assertEqual((contacts.invalid_emails, contacts.invalid_phones), (0, 0))
        bundle = build_autostat_exports([selection, selection], timestamp=NOW)
        self.assertEqual([file.filename for file in bundle.files], ['AUTOSTAT_Email.csv', 'AUTOSTAT_Viber.csv'])
        self.assertEqual([file.count for file in bundle.files], [1, 1])
        self.assertEqual(len(bundle.summary), 2)
        custom = b'Notes\nCustomer,Mail contact,Cell detail\n00002,next@example.com,09181234567\n'
        detected = inspect_contact_source(custom, ".csv", sheet_name="CSV", header_row=2)[0]
        mapped = replace(detected, account_column=0, email_columns=(1,), phone_columns=(2,))
        rows = extract_contacts(SourceSelection("custom.csv", custom, mapped))
        self.assertEqual(rows.phones, [("00002", "9181234567")])

    def test_actionable_errors_and_formula_cache(self):
        with self.assertRaisesRegex(AutomationError, "empty"):
            inspect_contact_source(b"", ".xlsx")
        with self.assertRaisesRegex(AutomationError, "Cannot read"):
            inspect_contact_source(b"broken", ".xlsx")
        with self.assertRaisesRegex(AutomationError, "Select at least"):
            build_autostat_exports([], timestamp=NOW)
        workbook = Workbook()
        workbook.active.append(["Account", "email"])
        workbook.active.append(["0001", '=CONCAT("a", "@example.com")'])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
        source = SourceSelection("formula.xlsx", data, inspect_contact_source(data, ".xlsx")[0])
        contacts = extract_contacts(source)
        self.assertEqual(contacts.source_error_rows, [2])
        self.assertEqual(contacts.emails, [])
        data = 'Account,email\n0001,ana\u6f22@example.com\n'.encode()
        source = SourceSelection("unicode.csv", data, inspect_contact_source(data, ".csv")[0])
        with self.assertRaisesRegex(AutomationError, "CSV UTF-8"):
            build_autostat_exports([source], timestamp=NOW)
        self.assertEqual(build_autostat_exports([source], timestamp=NOW, encoding="utf-8-sig").files[0].count, 1)


class AutostatAppTests(unittest.TestCase):
    def test_generated_campaign_files_match_geo_reco_logs_without_the_original_upload(self):
        records = [
            dict(primary_key=f"KEY-{index}", cardnumber=card, indexoneuid=f"UID-{index}",
                 email=email, phone=phone, OB=balance, product="CARDS", level=1,
                 AREA="Mindanao", ChristmasAmnesty="QUALIFIED")
            for index, (card, email, phone, balance) in enumerate((
                ("001234567890123456", "one@example.com", "09171234567", 9),
                ("001234567890123456", "two@example.com", "09181234567", 1000),
                ("5156041234567897.0", "bad..email@example.com", "09191234567", 100),
                ("", "missing-card@example.com", "09201234567", 2000),
            ))
        ]
        source_csv = StringIO(newline="")
        csv.writer(source_csv).writerows([
            DEFAULT_TEMPLATE_HEADERS,
            *[[record.get(header, record.get(header.lower(), "")) for header in DEFAULT_TEMPLATE_HEADERS]
              for record in records],
        ])
        positions = {header: DEFAULT_TEMPLATE_HEADERS.index(header)
                     for header in ("primary_key", "cardnumber", "indexoneuid", "email", "phone", "OB")}
        original = mapped_contact_source(
            "original.csv", source_csv.getvalue().encode(), "CSV", 1, len(DEFAULT_TEMPLATE_HEADERS), positions,
        )
        expected = build_autostat_exports([original], timestamp=None)
        self.assertEqual([file.count for file in expected.files], [2, 3])
        for builder in (build_geo_exports, build_negative_exports):
            campaign = builder(create_default_csv_template(), records, "", 1, ".csv")
            for excel in (False, True):
                with self.subTest(workflow=builder.__name__, excel=excel):
                    uploads = [Upload(
                        Path(file.filename).with_suffix(".xlsx").name if excel else file.filename,
                        build_export_excel(file) if excel else build_export_csv(file, protect_identifiers=True),
                    ) for file in campaign.files]
                    with patch("streamlit.file_uploader", return_value=uploads), patch("streamlit.download_button") as download:
                        app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
                        app.checkbox(key="autostat_blank_time").set_value(True).run(timeout=20)
                        self.assertEqual(len(app.exception), 0)
                        self.assertEqual(len(app.error), 0)
                        self.assertFalse(any(button.key == "autostat_prepare" for button in app.button))
                        actual = app.session_state["autostat_result"]["bundle"]
                        self.assertEqual([file.data for file in actual.files], [file.data for file in expected.files])
                        calls = {call.kwargs.get("key"): call for call in download.call_args_list}
                        self.assertEqual(calls["autostat_download_0"].kwargs["type"], "primary")
                        self.assertEqual(calls["autostat_download_1"].kwargs["type"], "primary")

    def test_automatic_logs_reuse_the_timestamp_and_hide_downloads_after_invalid_inputs(self):
        uploads = [Upload("campaign.csv", b"cardnumber,email,phone\n0001,one@example.com,09171234567\n")]
        with patch("streamlit.file_uploader", side_effect=lambda *args, **kwargs: uploads), \
                patch("autostat_page.build_autostat_exports", wraps=build_autostat_exports) as prepare:
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            result = app.session_state["autostat_result"]
            self.assertEqual(prepare.call_count, 1)
            app.run(timeout=20)
            self.assertEqual(prepare.call_count, 1)
            self.assertEqual(app.session_state["autostat_result"]["timestamp"], result["timestamp"])
            account = next(widget for widget in app.selectbox if widget.label == "Account number column")
            account.select(1).run(timeout=20)
            self.assertNotIn("autostat_result", app.session_state)
            self.assertEqual(len(app.get("download_button")), 0)
            self.assertEqual(prepare.call_count, 1)
            uploads[0] = Upload("campaign.csv", b"cardnumber,email,phone\n0002,bad-email,123\n")
            app.run(timeout=20)
            self.assertEqual([file.count for file in app.session_state["autostat_result"]["bundle"].files], [0, 0])
            self.assertEqual(len(app.get("download_button")), 0)
            self.assertTrue(any("No valid email or mobile" in item.value for item in app.info))
            self.assertEqual(len(app.exception), 0)

    def test_standard_excludes_bucket_and_insurance_worksheets_like_geo_reco(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for index, name in enumerate(("GEO", "B2", "B4", "SBC INSURANCE")):
            sheet = workbook.create_sheet(name)
            sheet.append(["cardnumber", "email", "phone"])
            sheet.append([f"000{index}", f"person{index}@example.com", "09171234567"])
        data = BytesIO()
        workbook.save(data)
        workbook.close()
        with patch("streamlit.file_uploader", return_value=[Upload("campaign.xlsx", data.getvalue())]):
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            result = app.session_state["autostat_result"]
            self.assertEqual([row["Sheet"] for row in result["bundle"].summary], ["GEO"])
            self.assertEqual([file.count for file in result["bundle"].files], [1, 1])
            self.assertEqual(app.multiselect[0].options, ["GEO"])

    def test_excel_download_is_generated_only_when_requested_and_precision_issues_are_visible(self):
        source = Workbook()
        source.active.append(["cardnumber", "email"])
        source.active.append([5156041234567897, "first@example.com"])
        output = BytesIO()
        source.save(output)
        source.close()
        upload = Upload("contacts.xlsx", output.getvalue())
        with patch("streamlit.file_uploader", return_value=[upload]), \
                patch("autostat_downloads.build_autostat_excel", wraps=build_autostat_excel) as build_excel, \
                patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertTrue(any("missing digits cannot be recovered" in warning.value for warning in app.warning))
            build_excel.assert_not_called()
            excel_download = next(call for call in download.call_args_list if call.kwargs.get("key") == "autostat_excel")
            self.assertEqual(excel_download.args[2], "AUTOSTAT.xlsx")
            self.assertEqual(excel_download.kwargs["type"], "secondary")
            self.assertEqual(app.selectbox(key="autostat_output_format").value, "Excel (.xlsx)")
            self.assertTrue(callable(excel_download.args[1]))
            workbook = load_workbook(BytesIO(excel_download.args[1]()))
            self.assertEqual(workbook["Email"]["B2"].value, "5156041234567897")
            self.assertEqual(workbook["Email"]["B2"].data_type, "s")
            workbook.close()
            build_excel.assert_called_once()

    def test_default_individual_downloads_and_zip_contain_literal_text_excel_accounts(self):
        upload = Upload("contacts.csv", b"cardnumber,email,phone\n5156041234567897.0,first@example.com,09171234567\n001234567890123456,second@example.com,09181234567\n")
        with patch("streamlit.file_uploader", return_value=[upload]), patch("streamlit.download_button") as download:
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            expected = ["5156041234567897", "001234567890123456"]
            calls = {call.kwargs.get("key"): call for call in download.call_args_list}
            for index in (0, 1):
                call = calls[f"autostat_download_{index}"]
                self.assertTrue(call.args[2].endswith(".xlsx"))
                self.assertTrue(callable(call.args[1]))
                workbook = load_workbook(BytesIO(call.args[1]()))
                self.assertEqual([cell.value for cell in workbook.active["B"][1:]], expected)
                self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in workbook.active["B"][1:]))
                workbook.close()
            with ZipFile(BytesIO(calls["autostat_zip"].args[1]())) as archive:
                self.assertEqual(archive.namelist(), ["AUTOSTAT_Email.xlsx", "AUTOSTAT_Viber.xlsx"])
                for name in archive.namelist():
                    workbook = load_workbook(BytesIO(archive.read(name)))
                    self.assertEqual([cell.value for cell in workbook.active["B"][1:]], expected)
                    self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in workbook.active["B"][1:]))
                    workbook.close()
            app.selectbox(key="autostat_output_format").select("CSV (.csv)").run(timeout=20)
            self.assertIn("autostat_result", app.session_state)
            calls = {call.kwargs.get("key"): call for call in download.call_args_list}
            self.assertTrue(calls["autostat_download_0"].args[2].endswith(".csv"))
            with ZipFile(BytesIO(calls["autostat_zip"].args[1])) as archive:
                self.assertEqual(archive.namelist(), ["AUTOSTAT_Email.csv", "AUTOSTAT_Viber.csv"])

    def test_abbreviated_source_accounts_are_skipped_and_new_valid_rows_replace_previous_downloads(self):
        uploads = [Upload("contacts.csv", b"cardnumber,email\n5156041234567897,first@example.com\n")]
        with patch("streamlit.file_uploader", side_effect=lambda *args, **kwargs: uploads):
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            self.assertIn("autostat_result", app.session_state)
            uploads[0] = Upload("contacts.csv", b"cardnumber,email\n5.15604E+15,first@example.com\n000002,second@example.com\n")
            app.run(timeout=20)
            self.assertIn("autostat_result", app.session_state)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(exported_rows(app.session_state["autostat_result"]["bundle"].files[0])[0]["Account No."], "000002")
            self.assertGreater(len(app.get("download_button")), 0)
            self.assertTrue(any("Skipped 1 database row with" in item.value for item in app.warning))

    def test_excluded_uploads_are_skipped_before_reading_and_only_included_accounts_are_counted(self):
        included = Upload("GEO.csv", b"cardnumber,email,phone\n0001,first@example.com,09171234567\n")
        uploads = [included, *[Upload(name, b"not a valid source") for name in (
            "SBC INSURANCE.csv", "NEGATIVE B2 - BEL.csv", "NEGATIVE B4 - CARDS.csv",
        )]]
        with patch("streamlit.file_uploader", side_effect=lambda *args, **kwargs: uploads):
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.multiselect), 3)
            self.assertEqual([(metric.label, metric.value) for metric in app.metric], [
                ("Valid email accounts", "1"), ("Valid phone accounts", "1"),
            ])
            self.assertEqual([row["File"] for row in app.session_state["autostat_result"]["bundle"].summary], ["GEO.csv"])
            uploads.pop(0)
            app.run(timeout=20)
            self.assertNotIn("autostat_result", app.session_state)
            self.assertFalse(any(button.key == "autostat_prepare" for button in app.button))
            self.assertEqual(len(app.get("download_button")), 0)

    def test_blank_time_exports_empty_cells_and_changing_time_settings_rebuilds_results(self):
        upload = Upload("contacts.csv", b"cardnumber,email,phone\n0001,first@example.com,09171234567\n")
        with patch("streamlit.file_uploader", return_value=[upload]):
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            timestamp = app.session_state["autostat_result"]["timestamp"]
            self.assertEqual(timestamp.tzinfo, MANILA_TIMEZONE)
            for file in app.session_state["autostat_result"]["bundle"].files:
                self.assertEqual(exported_rows(file)[0]["Time"], timestamp.strftime("%Y-%m-%d %I:%M:%S %p"))

            app.checkbox(key="autostat_blank_time").set_value(True).run(timeout=20)
            self.assertIsNone(app.session_state["autostat_result"]["timestamp"])
            self.assertTrue(app.checkbox(key="autostat_current_time").disabled)
            self.assertEqual(len(app.date_input), 0)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state["autostat_result"]
            self.assertIsNone(result["timestamp"])
            self.assertEqual([file.count for file in result["bundle"].files], [1, 1])
            for file in result["bundle"].files:
                self.assertIn("Time", file.headers)
                self.assertEqual(exported_rows(file)[0]["Time"], "")
                self.assertEqual(file.preview[0]["Time"], "")

            app.checkbox(key="autostat_blank_time").set_value(False).run(timeout=20)
            self.assertIsNotNone(app.session_state["autostat_result"]["timestamp"])
            app.checkbox(key="autostat_current_time").set_value(False).run(timeout=20)
            manual_timestamp = NOW.replace(second=0, microsecond=0)
            app.date_input(key="autostat_date").set_value(manual_timestamp.date()).run(timeout=20)
            app.time_input(key="autostat_time").set_value(manual_timestamp.time()).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            for file in app.session_state["autostat_result"]["bundle"].files:
                self.assertEqual(exported_rows(file)[0]["Time"], manual_timestamp.strftime("%Y-%m-%d %I:%M:%S %p"))

    def test_bel_cards_and_cardnumber_only_uploads_compile_into_two_channel_downloads(self):
        data = b'primary_key,email,PHONE\n@BEL-1,bel@example.com,09171234567\n00001234,cards@example.com,09181234567\n'
        uploads = [
            Upload('BEL CARDS BLASTING.xlsx.csv', data),
            Upload('cardnumber-only.csv', b'cardnumber,email,phone\n00009999,extra@example.com,09191234567\n'),
        ]
        with patch('streamlit.file_uploader', return_value=uploads), TemporaryDirectory(prefix='autostat-bel-test-') as directory:
            app = AppTest.from_file(str(ROOT / 'pages' / 'AUTOSTAT.py')).run(timeout=20)
            app.selectbox(key="autostat_output_format").select("CSV (.csv)").run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertFalse(any(widget.label == 'Separate BEL and CARDS using primary_key' for widget in app.checkbox))
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state['autostat_result']
            self.assertEqual([file.filename for file in result['bundle'].files], ['AUTOSTAT_Email.csv', 'AUTOSTAT_Viber.csv'])
            self.assertEqual([file.count for file in result['bundle'].files], [3, 3])
            self.assertEqual({row['Account No.'] for row in exported_rows(result['bundle'].files[0])},
                             {'@BEL-1', '00001234', '00009999'})
            self.assertEqual([file.headers for file in result['bundle'].files], [EMAIL_HEADERS, VIBER_HEADERS])
            self.assertEqual(len(app.get('download_button')), 4)
            self.assertEqual(app.dataframe[-1].value.loc[0, 'Account No.'], '@BEL-1')
            app.text_input(key='autostat_output_folder').set_value(directory).run(timeout=20)
            app.button(key='autostat_save').click().run(timeout=20)
            self.assertEqual(len(app.error), 0)
            saved = Path(app.session_state['autostat_result']['saved_folder'])
            self.assertEqual({path.name for path in saved.iterdir()}, {'AUTOSTAT_Email.csv', 'AUTOSTAT_Viber.csv'})
            for file in result['bundle'].files:
                self.assertEqual((saved / file.filename).read_bytes(), file.data)
            with ZipFile(BytesIO(result['bundle'].zip_data)) as archive:
                self.assertEqual(set(archive.namelist()), {'AUTOSTAT_Email.csv', 'AUTOSTAT_Viber.csv'})

    def test_cardnumber_is_selected_automatically_for_export(self):
        upload = Upload("contacts.csv", b"primary_key,email,cardnumber,phone\nINTERNAL-1,first@example.com,00001234,09171234567\n")
        with patch("streamlit.file_uploader", return_value=[upload]):
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            account = next(widget for widget in app.selectbox if widget.label == "Account number column")
            self.assertEqual(account.value, 2)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            for file in app.session_state["autostat_result"]["bundle"].files:
                self.assertEqual(exported_rows(file)[0]["Account No."], "00001234")

    def test_navigation_registers_autostat(self):
        app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
        self.assertEqual(len(app.exception), 0)
        app.switch_page("pages/AUTOSTAT.py").run(timeout=20)
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.title[0].value, "AUTOSTAT")

    def test_multiple_uploads_preview_download_folder_and_input_invalidation(self):
        uploads = [Upload("source.xlsx", workbook_bytes()), Upload("extra.csv", b"Account,email,phone\n00999,extra@example.com,09191234567\n")]
        def uploader(label, **kwargs):
            return uploads if kwargs.get("key") == "autostat_uploads" else None

        with patch("streamlit.file_uploader", side_effect=uploader), TemporaryDirectory(prefix="autostat-test-") as directory:
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual([panel.label for panel in app.expander], ["Log settings", "Edit source columns", "Save to a folder", "Preview"])
            source_panel = app.expander[1]
            self.assertFalse(source_panel.proto.expanded)
            self.assertEqual(len(source_panel.multiselect), 8)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state["autostat_result"]
            self.assertEqual([file.filename for file in result['bundle'].files], ['AUTOSTAT_Email.csv', 'AUTOSTAT_Viber.csv'])
            self.assertEqual([file.count for file in result['bundle'].files], [4, 5])
            self.assertEqual([(metric.label, metric.value) for metric in app.metric], [
                ("Valid email accounts", "4"), ("Valid phone accounts", "5"),
            ])
            self.assertEqual(len(app.dataframe), 1)
            self.assertEqual(len(result['bundle'].summary), 3)
            self.assertEqual(len(app.get("download_button")), 4)
            self.assertEqual(sum(row["Viber rows"] for row in result["bundle"].summary), 5)
            self.assertFalse(Path(directory).joinpath("AUTOSTAT").exists())
            app.run(timeout=20)
            self.assertEqual(app.session_state["autostat_result"]["timestamp"], result["timestamp"])
            app.text_input(key="autostat_output_folder").set_value(directory).run(timeout=20)
            app.button(key="autostat_save").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            saved = Path(app.session_state["autostat_result"]["saved_folder"])
            self.assertEqual({path.name for path in saved.iterdir()}, {Path(file.filename).with_suffix(".xlsx").name for file in result["bundle"].files})
            for file in result["bundle"].files:
                workbook = load_workbook(saved / Path(file.filename).with_suffix(".xlsx").name)
                self.assertEqual([cell.value for cell in workbook.active["B"][1:]],
                                 [row["Account No."] for row in exported_rows(file)])
                self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in workbook.active["B"][1:]))
                workbook.close()
            app.session_state["autostat_remark_by"] = "USER2"
            app.run(timeout=20)
            self.assertEqual(app.text_input(key="autostat_remark_by").value, "MBURGOS")
            self.assertTrue(app.text_input(key="autostat_remark_by").proto.disabled)
            self.assertTrue(all(file.preview[0]["Remark By"] == "MBURGOS" for file in app.session_state["autostat_result"]["bundle"].files if file.count))
            uploads.pop()
            app.run(timeout=20)
            self.assertEqual([file.count for file in app.session_state["autostat_result"]["bundle"].files], [3, 4])


if __name__ == "__main__":
    unittest.main()
