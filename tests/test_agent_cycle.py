from __future__ import annotations

import csv
import unittest
from collections import Counter
from datetime import date, timedelta
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from agent_cycle import AgentCycle, assign_agents_to_accounts, cycle_agents, default_cycle, load_cycle, save_cycle
from automation import AutomationError, DEFAULT_TEMPLATE_HEADERS, create_default_csv_template
from export_config import GEO_RECO_AGENTS
from exports import build_export_excel, build_geo_exports, build_negative_exports, build_sheet_exports


ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 30)


def records():
    return [{"primary_key": f"{number:018d}", "cardnumber": f"{number + 100:018d}", "indexoneuid": f"UID-{number}",
             "email": f"account{number}@example.com" if number != 1 else None,
             "phone": "09171234567" if number != 2 else None, "OB": number, "level": 1,
             "product": "MC" if number < 8 else "PL", "AREA": "Mindanao", "ChristmasAmnesty": None}
            for number in range(1, 12)]


def rows(file):
    return list(csv.DictReader(StringIO(file.data.decode("cp1252"))))


class AgentCycleTests(unittest.TestCase):
    def test_daily_first_agent_moves_to_last_and_repeats_with_matching_names(self):
        cycle = default_cycle()
        for offset in range(-1, 8):
            agents = cycle_agents(cycle, TODAY + timedelta(days=offset))
            first = offset % 3
            self.assertEqual(agents, GEO_RECO_AGENTS[first:] + GEO_RECO_AGENTS[:first])

    def test_changed_cycle_is_saved_and_loaded_without_losing_the_start_date(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "settings" / "cycle.json"
            self.assertEqual(load_cycle(path), default_cycle())
            cycle = AgentCycle(TODAY, ("ROSVINOYA", "JBATALLER", "APMAMARIL"))
            save_cycle(cycle, path)
            self.assertEqual(load_cycle(path), cycle)
            self.assertEqual(cycle_agents(load_cycle(path), TODAY + timedelta(days=1))[0], GEO_RECO_AGENTS[2])
            saved = path.read_bytes()
            with patch("agent_cycle.os.replace", side_effect=PermissionError):
                with self.assertRaisesRegex(AutomationError, "Could not save"):
                    save_cycle(default_cycle(), path)
            self.assertEqual(path.read_bytes(), saved)
            self.assertEqual(list(path.parent.iterdir()), [path])
            path.write_text("broken json", encoding="utf-8")
            with self.assertRaisesRegex(AutomationError, "cannot be read"):
                load_cycle(path)

    def test_duplicate_agents_are_rejected_and_duplicate_account_rows_keep_one_agent(self):
        with self.assertRaisesRegex(AutomationError, "exactly once"):
            cycle_agents(AgentCycle(TODAY, ("APMAMARIL",) * 3), TODAY)
        source = records()
        source.insert(1, dict(source[0]))
        assigned = assign_agents_to_accounts(source, GEO_RECO_AGENTS)
        self.assertEqual(assigned[0]["Terms"], assigned[1]["Terms"])
        self.assertEqual(assigned[2]["Terms"], "ROSVINOYA")
        self.assertNotIn("Terms", source[0])

    def test_geo_reco_and_all_accounts_use_balanced_paired_assignments_in_all_formats(self):
        source = records()
        original = [dict(row) for row in source]
        headers = [header for header in DEFAULT_TEMPLATE_HEADERS if header not in ("Monthly", "Terms")]
        for mode in ("GEO", "RECO", "All accounts"):
            for include_headers in (True, False):
                with self.subTest(mode=mode, include_headers=include_headers):
                    template = create_default_csv_template() if include_headers else (",".join(headers) + "\n").encode()
                    if mode == "All accounts":
                        bundle = build_sheet_exports(template, source, "", 1, base_name="Accounts", label="Accounts",
                                                     template_extension=".csv", agents=GEO_RECO_AGENTS)
                    else:
                        builder = build_geo_exports if mode == "GEO" else build_negative_exports
                        bundle = builder(template, source, "", 1, ".csv", agents=GEO_RECO_AGENTS)
                    assignments, group_accounts = {}, {}
                    for file in bundle.files:
                        exported = rows(file)
                        for row in exported:
                            pair = row["Monthly"], row["Terms"]
                            self.assertIn(pair, GEO_RECO_AGENTS)
                            self.assertEqual(assignments.setdefault(row["primary_key"], pair), pair)
                            group_accounts.setdefault(file.group_key, {})[row["primary_key"]] = row["Terms"]
                        self.assertEqual([row["Monthly"] for row in exported], [row["Monthly"] for row in file.preview])
                        workbook = load_workbook(BytesIO(build_export_excel(file)))
                        saved_headers = [cell.value for cell in workbook.active[1]]
                        saved_rows = list(workbook.active.values)[1:]
                        self.assertEqual([row[saved_headers.index("Terms")] for row in saved_rows], [row["Terms"] for row in exported])
                        self.assertEqual([row[saved_headers.index("Monthly")] for row in saved_rows], [row["Monthly"] for row in exported])
                        workbook.close()
                    self.assertEqual(len(assignments), len(source))
                    for accounts in group_accounts.values():
                        counts = Counter(accounts.values())
                        self.assertLessEqual(max(counts.get(code, 0) for _, code in GEO_RECO_AGENTS)
                                             - min(counts.get(code, 0) for _, code in GEO_RECO_AGENTS), 1)
                    self.assertEqual(assignments[f"{11:018d}"][1], "APMAMARIL")
                    self.assertEqual(source, original)


class Upload(BytesIO):
    def __init__(self):
        workbook = Workbook()
        workbook.active.title = "NEGATIVE"
        headers = ["primary_key", "cardnumber", "indexoneuid", "email", "phone", "OB", "product", "level", "AREA", "ChristmasAmnesty"]
        workbook.active.append(headers)
        for row in records():
            workbook.active.append([row.get(header) for header in headers])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        super().__init__(output.getvalue())
        self.name = "GEO.xlsx"


class AgentCyclePageTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "cycle.json"
        self.addCleanup(patch.stopall)
        patch("agent_cycle.SETTINGS_PATH", self.path).start()
        patch("agent_cycle_page.manila_today", return_value=TODAY).start()

    def test_cycle_visible_without_uploads_and_next_day_rotates_to_rosalie(self):
        with patch("streamlit.file_uploader", return_value=None):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.table[0].value["Terms"].tolist(), ["APMAMARIL", "ROSVINOYA", "JBATALLER"])
            app.date_input(key=f"agent_cycle_date_{TODAY.isoformat()}").set_value(TODAY + timedelta(days=1)).run(timeout=20)
            self.assertEqual(app.table[0].value["Terms"].tolist(), ["ROSVINOYA", "JBATALLER", "APMAMARIL"])
            self.assertFalse(self.path.exists())

    def test_saving_a_new_cycle_invalidates_exports_and_survives_clear_and_new_sessions(self):
        upload = Upload()
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: upload if kw.get("key") == "source_upload" else None):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            next(button for button in app.button if button.label == "Preview and prepare files").click().run(timeout=20)
            self.assertIn("result", app.session_state)
            choices = [widget for widget in app.selectbox if widget.key.startswith("agent_cycle_order_")]
            choices[0].select("ROSVINOYA").run(timeout=20)
            self.assertTrue(app.button(key="agent_cycle_save").disabled)
            choices = [widget for widget in app.selectbox if widget.key.startswith("agent_cycle_order_")]
            choices[1].select("JBATALLER")
            choices[2].select("APMAMARIL")
            app.run(timeout=20)
            app.button(key="agent_cycle_save").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertNotIn("result", app.session_state)
            self.assertEqual(load_cycle().agent_codes, ("ROSVINOYA", "JBATALLER", "APMAMARIL"))
            next(button for button in app.button if button.label == "Preview and prepare files").click().run(timeout=20)
            for file in app.session_state["result"]["files"]:
                self.assertEqual(file.records[0]["Terms"], "ROSVINOYA")
            self.assertTrue(all(value in {name for name, _ in GEO_RECO_AGENTS}
                                for value in app.main.dataframe[-1].value["Monthly"]))
            app.button(key="clear_all_files").click().run(timeout=20)
            self.assertEqual(load_cycle().agent_codes[0], "ROSVINOYA")
            fresh = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            self.assertEqual(fresh.table[0].value["Terms"].tolist(), ["ROSVINOYA", "JBATALLER", "APMAMARIL"])

    def test_new_manila_day_rotates_automatically_and_clears_previous_exports(self):
        upload = Upload()
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: upload if kw.get("key") == "source_upload" else None):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            next(button for button in app.button if button.label == "Preview and prepare files").click().run(timeout=20)
            self.assertIn("result", app.session_state)
            with patch("agent_cycle_page.manila_today", return_value=TODAY + timedelta(days=1)):
                app.run(timeout=20)
                self.assertEqual(len(app.exception), 0)
                self.assertNotIn("result", app.session_state)
                self.assertEqual(app.table[0].value["Terms"].tolist(), ["ROSVINOYA", "JBATALLER", "APMAMARIL"])
                self.assertEqual(next(widget for widget in app.date_input if widget.label == "Assignment date (Asia/Manila)").value,
                                 TODAY + timedelta(days=1))

    def test_corrupt_saved_cycle_blocks_assignments_and_can_be_reset(self):
        self.path.write_text("bad json", encoding="utf-8")
        with patch("streamlit.file_uploader", return_value=None):
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any("cannot be read" in item.value for item in app.error))
            self.assertNotIn("result", app.session_state)
            app.button(key="agent_cycle_reset").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(load_cycle(), default_cycle())


if __name__ == "__main__":
    unittest.main()
