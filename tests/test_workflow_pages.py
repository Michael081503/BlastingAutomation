from __future__ import annotations

import csv
import unittest
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from automation import DEFAULT_TEMPLATE_HEADERS, FIELD_MAPPING
from export_config import CHANNEL_SUFFIXES, GEO_OUTPUT_FILENAMES, OUTPUT_FILENAMES


ROOT = Path(__file__).resolve().parents[1]


class Upload(BytesIO):
    def __init__(self, data: bytes):
        super().__init__(data)
        self.name = "mixed-workflows.xlsx"


def mixed_workbook() -> bytes:
    workbook = Workbook()
    grid = workbook.active
    grid.title = "DataGrid"
    grid.append(["No", "Name", "New I.C.", "Account No.", "Email", "Financing/Card No.", "Status Code"])
    grid.append([1, "DataGrid account", None, "GRID-001", "grid@example.com", "GRID-CARD-001", "ACTIVE"])

    def add_accounts(name: str, cases: list[tuple], *, geo: bool = False) -> None:
        sheet = workbook.create_sheet(name)
        headers = [*FIELD_MAPPING, *(["Area"] if geo else [])]
        sheet.append(headers)
        for account, product, level, status, area in cases:
            record = {
                "debtor_email": f"{account.lower()}@example.com",
                "OLD I.C": f"IC-{account}",
                "Account no.": account,
                "financing/Card no.": f"CARD-{account}",
                "Name": f"Account {account}",
                "endo date": datetime(2026, 10, 28),
                "level": level,
                "product": product,
                "OB (Masterlist)": 1234.5,
                "phone_1": "09171234567",
                "Christmas Amnesty": status,
                "Area": area,
            }
            sheet.append([record.get(header) for header in headers])

    add_accounts("NEGATIVE", [
        ("RECO001", "CARDS", 1, None, None),
        ("RECO006", "CARDS", 6, None, None),
        ("RECOBEL", "BEL", 4, None, None),
        ("RECOLOAN", "PL", 3, None, None),
    ])
    add_accounts("GEO", [
        ("GEOM001", "CARDS", 1, "QUALIFIED", "Mindanao"),
        ("GEOM006", "CARDS", 6, "QUALIFIED", "Mindanao"),
        ("GEORBEL", "BEL", 3, None, "Region 3"),
    ], geo=True)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def csv_rows(file) -> list[dict[str, str]]:
    return list(csv.DictReader(StringIO(file.data.decode("cp1252"))))


class WorkflowPageTests(unittest.TestCase):
    def assert_no_embedded_b2(self, app) -> None:
        self.assertFalse(any(button.key == "b2_prepare" for button in app.button))
        self.assertFalse(any(widget.label == "SBC database sheet" for widget in app.selectbox))
        self.assertFalse(any("match accounts and contacts" in heading.value for heading in app.subheader))
        self.assertFalse(any(heading.value in ("B2 files", "B2 activity logs") for heading in app.subheader))

    def test_negative_b2_activity_uploads_prepare_all_channels_independently_of_matching(self):
        matching = {}
        csv_upload = Upload(
            b"cardnumber,primary_key,FullName,email,phone,Collector\n"
            b"CARD1,000000000000000001,CSV Person,csv@example.com,+639171234567,COL1\n"
            b"CARD2,000000000000000002,Email only,only@example.com,123,COL2\n"
        )
        csv_upload.name = "NEGATIVE B2 - CARDS.csv"
        workbook = Workbook()
        workbook.active.append(["primary_key", "cardnumber", "FullName", "email", "phone", "Collector"])
        workbook.active.append(["000000000000000003", "CARD3", "Excel Person", None, "09181234567", "COL3"])
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        excel_upload = Upload(output.getvalue())
        excel_upload.name = "NEGATIVE B2 - BEL.xlsx"
        b2 = [csv_upload, excel_upload]

        def uploader(label, **kwargs):
            return b2 if kwargs.get("key") == "geo_b2_autostat_uploads" else matching.get(kwargs.get("key"))

        with patch("streamlit.file_uploader", side_effect=uploader), TemporaryDirectory(prefix="negative-b2-test-") as directory:
            app = AppTest.from_file(str(ROOT / "pages" / "NEGATIVE_B2.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            self.assertNotIn("result", app.session_state)
            self.assertNotIn("geo_autostat_result", app.session_state)
            result = app.session_state["geo_b2_autostat_result"]
            self.assertEqual([file.count for file in result["bundle"].files], [2, 2, 2])
            self.assertEqual([row["Account No."] for row in csv_rows(result["bundle"].files[0])],
                             ["000000000000000001", "000000000000000002"])
            for file in result["bundle"].files[1:]:
                self.assertEqual([row["Account number"] for row in csv_rows(file)],
                                 ["000000000000000001", "000000000000000003"])
                self.assertEqual([row["Mobile Number"] for row in csv_rows(file)], ["09171234567", "09181234567"])
            app.run(timeout=20)
            self.assertEqual(app.session_state["geo_b2_autostat_result"]["timestamp"], result["timestamp"])
            app.text_input(key="geo_b2_autostat_output_folder").set_value(directory).run(timeout=20)
            app.button(key="geo_b2_autostat_save").click().run(timeout=20)
            self.assertEqual(len(app.error), 0)
            saved = Path(app.session_state["geo_b2_autostat_result"]["saved_folder"])
            self.assertEqual(len(list(saved.iterdir())), 3)
            for path in saved.iterdir():
                book = load_workbook(path)
                account_column = "B" if path.stem.endswith("Email") else "A"
                self.assertTrue(all(cell.value.startswith("00000000000000000") and cell.data_type == "s"
                                    and cell.number_format == "@" for cell in book.active[account_column][1:]))
                book.close()
            for key, headers, row in (
                ("b2_database_upload", ["Account Key", "Contract Number", "Bucket Code", "Product Category Code", "EMAIL (01)", "Mobile 1"],
                 ["MATCH-1", "MATCH-CARD", 2, "MC", "matched@example.com", "09191234567"]),
                ("b2_grid_upload", ["Account No.", "Financing/Card No."], ["MATCH-1", "MATCH-CARD"]),
            ):
                workbook = Workbook()
                workbook.active.append(headers)
                workbook.active.append(row)
                output = BytesIO()
                workbook.save(output)
                workbook.close()
                matching[key] = Upload(output.getvalue())
            app.run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.session_state["geo_b2_autostat_result"]["timestamp"], result["timestamp"])
            app.button(key="b2_prepare").click().run(timeout=20)
            self.assertEqual(len(app.error), 0)
            self.assertIn("b2_result", app.session_state)
            self.assertEqual(app.session_state["geo_b2_autostat_result"]["timestamp"], result["timestamp"])
            replacement = Upload(b"cardnumber,primary_key,email,phone\nCARD9,000009,updated@example.com,09191234567\n")
            replacement.name = csv_upload.name
            b2[0] = replacement
            app.run(timeout=20)
            self.assertEqual(len(app.error), 0)
            self.assertEqual(csv_rows(app.session_state["geo_b2_autostat_result"]["bundle"].files[0])[0]["Account No."], "000009")
            b2.clear()
            app.run(timeout=20)
            self.assertNotIn("geo_b2_autostat_result", app.session_state)
            self.assertIn("b2_result", app.session_state)

    def test_mixed_workbook_keeps_geo_and_reco_grouped_exports_on_main_page(self):
        upload = Upload(mixed_workbook())
        with patch("streamlit.file_uploader", side_effect=lambda label, **kw: upload if kw.get("key") == "source_upload" else None) as uploader:
            app = AppTest.from_file(str(ROOT / "app.py")).run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.title[0].value, "GEO & RECO")
            self.assertEqual({call.kwargs["key"] for call in uploader.call_args_list}, {"source_upload", "template_upload"})
            source = next(widget for widget in app.selectbox if widget.label == "Source sheet")
            self.assertEqual(set(source.options), {"DataGrid", "NEGATIVE", "GEO"})
            self.assertEqual(source.value, "NEGATIVE")
            self.assert_no_embedded_b2(app)

            source.select("GEO").run(timeout=20)
            self.assertEqual(app.radio[0].value, "Separate Area and Christmas Amnesty groups")
            next(button for button in app.button if button.label == "Preview and prepare files").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state["result"]
            self.assertTrue(result["geo"])
            self.assertTrue(result["grouped"])
            self.assertEqual(result["group_counts"]["mindanao_qualified_cards"], 2)
            self.assertEqual(result["group_counts"]["r3_blank_bel"], 1)
            self.assertEqual({file.filename for file in result["files"]}, {
                f"GEO - {GEO_OUTPUT_FILENAMES[group]}{CHANNEL_SUFFIXES[channel]}.csv"
                for group in ("mindanao_qualified_cards", "r3_blank_bel")
                for channel in ("email", "viber")
            })
            for file in result["files"]:
                rows = csv_rows(file)
                self.assertEqual(tuple(rows[0]), DEFAULT_TEMPLATE_HEADERS)
                self.assertTrue(all(row["primary_key"] in {"GEOM001", "GEOM006", "GEORBEL"} for row in rows))
                self.assertTrue(all(row["cardnumber"] == "CARD-" + row["primary_key"] for row in rows))
                expected_class = "SBC GEO Mindanao" if file.group_key.startswith("mindanao_") else "SBC GEO R3"
                self.assertTrue(all(row["doclass"] == expected_class for row in rows))
            self.assert_no_embedded_b2(app)

            next(widget for widget in app.selectbox if widget.label == "Source sheet").select("NEGATIVE").run(timeout=20)
            self.assertNotIn("result", app.session_state)
            self.assertEqual(app.radio[0].value, "Separate Christmas Amnesty groups")
            next(button for button in app.button if button.label == "Preview and prepare files").click().run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(app.error), 0)
            result = app.session_state["result"]
            self.assertFalse(result["geo"])
            self.assertTrue(result["grouped"])
            groups = ("blank_cards_level_1", "blank_cards_level_6", "blank_bel", "blank_other_products")
            self.assertTrue(all(result["group_counts"][group] == 1 for group in groups))
            self.assertEqual({file.filename for file in result["files"]}, {
                f"{OUTPUT_FILENAMES[group]}{CHANNEL_SUFFIXES[channel]}.csv"
                for group in groups for channel in ("email", "viber")
            })
            self.assertEqual({row["primary_key"] for file in result["files"] for row in csv_rows(file)},
                             {"RECO001", "RECO006", "RECOBEL", "RECOLOAN"})
            self.assertTrue(all(row["cardnumber"] == "CARD-" + row["primary_key"]
                                for file in result["files"] for row in csv_rows(file)))
            with ZipFile(BytesIO(result["zip_data"])) as archive:
                self.assertEqual(set(archive.namelist()), {file.filename for file in result["files"]})
                for file in result["files"]:
                    self.assertEqual(archive.read(file.filename), file.data)
            self.assert_no_embedded_b2(app)

    def test_remembered_b2_workflow_does_not_replace_autostat(self):
        with patch("streamlit.file_uploader", return_value=None):
            app = AppTest.from_file(str(ROOT / "pages" / "AUTOSTAT.py"))
            app.session_state["autostat_workflow"] = "B2: Match DataGrid with SBC database"
            app.run(timeout=20)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(app.title[0].value, "AUTOSTAT")
            self.assertTrue(any(widget.key == "autostat_remark_by" for widget in app.text_input))
            self.assertNotIn("autostat_result", app.session_state)
            self.assertFalse(any(widget.key == "autostat_workflow" for widget in app.radio))
            self.assert_no_embedded_b2(app)


if __name__ == "__main__":
    unittest.main()
